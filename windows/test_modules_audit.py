"""Ошибки модулей (свод, объединение, входящие, Word→Excel, формулы, орфография) — чтобы не вернулись."""
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

import iskra_compose as compose
import iskra_index as ix
import iskra_inbox as ib
import iskra_svod as sv
import iskra_telegram as tg
from iskra_crossformat import as_number
from iskra_formula import build_formula
from iskra_spell import SpellChecker, fix_typical


class ModulesAudit(unittest.TestCase):
    def setUp(self):
        self.d = Path(tempfile.mkdtemp())

    def test_excel_with_empty_first_column_is_read(self):
        wb = Workbook(); ws = wb.active
        ws['B1'] = 'Наименование'; ws['C1'] = 'по списку'; ws['B2'] = 'Госпиталь'; ws['C2'] = 5
        wb.save(self.d / 'e.xlsx')
        self.assertIn('Госпиталь | 5', ix.extract(self.d / 'e.xlsx'))

    def test_svod_title_row_hidden_after_total(self):
        wb = Workbook(); ws = wb.active
        ws['A1'] = 'Сведения о численности по организациям'; ws['D1'] = 'на 01.10.2026'
        ws.append(['№', 'Наименование', 'по списку', 'на лицо'])
        ws.append([1, 'Госпиталь', 120, 112]); ws.append([2, 'Полигон', 50, 48])
        ws.append([None, 'Итого', 170, 160]); ws.append([None, 'старые данные', 100, 100])
        ws.append([3, 'Скрыто', 7, 7]); ws.row_dimensions[ws.max_row].hidden = True
        wb.save(self.d / 'a.xlsx')
        t = sv.read_tables(self.d / 'a.xlsx')[0]
        self.assertEqual(t.columns, ['по списку', 'на лицо'])
        self.assertEqual([r[0] for r in t.rows], ['Госпиталь', 'Полигон'])
        self.assertEqual(len(t.notes), 2)

    def test_svod_module_names_merge(self):
        for n, (org, a) in enumerate([('Модуль1', 10), ('Модуль №1', 3)]):
            wb = Workbook(); ws = wb.active; ws.append(['Наименование', 'по списку']); ws.append([org, a])
            wb.save(self.d / f'm{n}.xlsx')
        s = sv.combine(sv.read_tables(self.d / 'm0.xlsx') + sv.read_tables(self.d / 'm1.xlsx'))
        self.assertEqual([(o.name, s.value(o, 'по списку')[0]) for o in s.orgs], [('Модуль1', 13)])

    def test_merge_keeps_each_organisation_and_deadline(self):
        m = tg.merge([('a', 'Текст:\nГоспиталю представить донесение о численности до 15.11.2026.'),
                      ('b', 'Текст:\nПолигону представить донесение о численности до 15.11.2026.')])
        text = m.text()
        self.assertIn('Госпиталю', text); self.assertIn('Полигону', text)
        m = tg.merge([('a', 'Текст:\nСрочно представить подробное донесение о численности личного состава в штаб округа.'),
                      ('b', 'Текст:\nПредставить донесение о численности личного состава в штаб к 20.11.2026.')])
        self.assertIn('20.11.2026', m.text())

    def test_sentences_keep_addresses(self):
        self.assertEqual(len(tg._sentences('Прибыть по адресу: г. Москва, ул. Ленина, д. 5. Иметь документы.')), 2)

    def test_fix_typical_keeps_correct_text(self):
        self.assertEqual(fix_typical('Направить в в/ч 12345. Оплата за счёт средств.')[0], 'Направить в в/ч 12345. Оплата за счёт средств.')

    def test_inbox_own_number_and_sender(self):
        t = ('ТЕЛЕГРАММА\nВо исполнение приказа от 01.09.2026 № 345 представить сведения до 30.10.2026.\n'
             'Начальник штаба\nполковник И.И. Иванов\nИсх. № 12/345 от 20.10.2026')
        self.assertEqual(compose.reference(t), ('12/345', '20.10.2026'))
        self.assertIn('Иванов', ib.sender_of(t))
        self.assertEqual(ib.kind_of('x.docx', 'П Р И К А З\nтекст'), 'Приказ')

    def test_old_formats_shown_not_read(self):
        (self.d / 'тлг.doc').write_bytes(b'x')
        item = ib.analyze(self.d / 'тлг.doc')
        self.assertIn('формат не читается', item['kind'])

    def test_word_numbers_become_numbers(self):
        self.assertEqual(as_number('1 234'), 1234)
        self.assertEqual(as_number('12,5'), 12.5)
        self.assertEqual(as_number('007'), '007')

    def test_formulas(self):
        self.assertEqual(build_formula('если в столбце C больше 100 то да иначе нет')[0], '=ЕСЛИ(C2>100;"Да";"Нет")')
        self.assertEqual(build_formula('сумма столбца D где в столбце B "госпиталь" и в столбце C больше 10')[0],
                         '=СУММЕСЛИМН(D:D;B:B;"госпиталь";C:C;">10")')

    def test_office_words_not_typos(self):
        bad = [i.text for i in SpellChecker().check('ВрИО начальника медслужбы направил шифротелеграмму по спецсвязи. Превет.',
                                                   with_style=False)]
        self.assertEqual(bad, ['Превет'])


if __name__ == '__main__':
    unittest.main()
