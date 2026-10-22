import os, re, threading, queue, time
from dataclasses import replace
from datetime import date
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from iskra_services import AutoIndexer
from iskra_tools_ui import ToolsMixin
from iskra_work_ui import WorkMixin
from iskra_ai_ui import AIMixin
from iskra_plan_ui import PlanMixin
from iskra_hotkey_ui import HotkeyMixin
from iskra_people_ui import PeopleMixin
from iskra_journal_ui import JournalMixin
from iskra_voice_ui import VoiceMixin
from iskra_learn_ui import LearnMixin
import iskra_ai
import iskra_bugreport
from iskra_formula import build_formula, explain_formula
from iskra_macros import MACROS, export_macro
from iskra_workflows import compare_texts, save_history, read_history
from iskra_consolidate import consolidate
from iskra_index import ROOT, extract, index_folder, search_docs, forget_folder, stats, iter_documents
import iskra_assistant as assistant

ROOT.mkdir(parents=True, exist_ok=True)
VERSION = '1.9.7'
ACCENT = '#24518A'


# Команды, после которых фраза может быть уточнением прошлого ответа нейросети.
DOC_CMDS = {'deadlines', 'dates', 'people', 'numbers', 'money', 'contacts', 'summary'}
FOLLOW_CMDS = {'search', 'write', 'ask', 'telegram'} | DOC_CMDS


class App(LearnMixin, VoiceMixin, JournalMixin, PeopleMixin, HotkeyMixin, PlanMixin, AIMixin, WorkMixin, ToolsMixin):
    def __init__(self, root):
        self.root = root
        import iskra_clipboard
        iskra_clipboard.install(root)                    # Ctrl+C/V и при русской раскладке, меню правой кнопкой
        iskra_bugreport.install(root, lambda text: self.events.put(('call', lambda: self.say(
            'Произошла ошибка, она записана в журнал. Если мешает работать — меню Файл → «Сообщить об ошибке…».'))))
        root.title(f'Искра {VERSION} — помощник по документам')
        root.geometry('1180x740'); root.minsize(860, 540)
        self.events = queue.Queue()
        self._ai_queue, self._bg_jobs, self._bg_yield = [], [], threading.Event()
        self.results = []          # [(path, snippet)] — так же, как раньше, для инструментов
        self.doc_cache = {}        # path -> text выбранных документов
        self.indexer = AutoIndexer(ROOT, index_folder, self.events.put)
        self.indexer.start()
        self.start_inbox()
        self.start_planning()
        self.start_people()
        self.start_learning()                          # в фоне: люди и слова из документов
        self.start_journal()
        self.telegram_template_path()                    # шаблон телеграмм — и для нейросети
        self.ai_ready = iskra_ai.available()
        root.bind('<Escape>', self.ai_stop)
        root.protocol('WM_DELETE_WINDOW', self.close)
        self._style()
        self._menu()

        top = ttk.Frame(root, padding=(12, 10, 12, 4)); top.pack(fill='x')
        ttk.Label(top, text='Искра', style='Title.TLabel').pack(side='left')
        ttk.Label(top, text='  поиск и помощник по документам · всё на этом компьютере, без интернета',
                  style='Muted.TLabel').pack(side='left')
        ttk.Button(top, text='Папки…', command=self.folders_window).pack(side='right')
        ttk.Button(top, text='Обновить индекс', command=self.indexer.scan_now).pack(side='right', padx=6)

        bar = ttk.Frame(root, padding=(12, 4)); bar.pack(fill='x')
        self.query = tk.StringVar()
        entry = ttk.Entry(bar, textvariable=self.query, font=('Segoe UI', 12))
        entry.pack(side='left', fill='x', expand=True); entry.bind('<Return>', lambda e: self.find())
        ttk.Button(bar, text='Найти', style='Accent.TButton', command=self.find).pack(side='left', padx=(8, 0))
        entry.focus_set()

        pane = ttk.Panedwindow(root, orient='horizontal'); pane.pack(fill='both', expand=True, padx=12, pady=6)
        left = ttk.Frame(pane); pane.add(left, weight=2)
        self.found_label = tk.StringVar(value='Документы: введите запрос и нажмите «Найти»')
        ttk.Label(left, textvariable=self.found_label, style='Muted.TLabel').pack(anchor='w')
        lf = ttk.Frame(left); lf.pack(fill='both', expand=True, pady=(4, 0))
        self.listbox = tk.Listbox(lf, font=('Segoe UI', 10), activestyle='none', selectbackground=ACCENT, exportselection=False)
        self.listbox.pack(side='left', fill='both', expand=True)
        sc = ttk.Scrollbar(lf, command=self.listbox.yview); sc.pack(side='right', fill='y'); self.listbox.config(yscrollcommand=sc.set)
        self.listbox.bind('<<ListboxSelect>>', self.show)
        self.listbox.bind('<Double-Button-1>', lambda e: self.open_selected())
        self.listbox.bind('<Return>', lambda e: self.open_selected())

        right = ttk.Frame(pane); pane.add(right, weight=3)
        self.tabs = ttk.Notebook(right); self.tabs.pack(fill='both', expand=True)
        self._assistant_tab()
        self._document_tab()
        self._deadlines_tab()
        self._inbox_tab()
        self._control_tab()

        self.status = tk.StringVar(value='Готов к работе')
        ttk.Label(root, textvariable=self.status, style='Muted.TLabel', padding=(12, 2, 12, 8)).pack(fill='x')
        self.say(assistant.HELP)
        self.say(f'Сейчас в поиске документов: {stats()}. Папки для поиска добавляются кнопкой «Папки…».')
        self.say(self.ai_status_line(), 'proof')
        self.say(self.addons_status_line(), 'proof')
        self.ai_gather()
        if self.ai_ready:
            # Модель грузится в память заранее, пока человек читает сводку: первый ответ — без ожидания загрузки.
            self.root.after(8000, lambda: threading.Thread(target=iskra_ai.warm, daemon=True).start())
        self.start_hotkey()
        self.root.after(600, self.maybe_first_run)
        self.root.after(2500, self.offer_restore)            # после сбоя — вернуть несохранённые черновики
        # Сводку при запуске не показываем в первый запуск (сначала мастер) и не выдёргиваем фокус, если человек уже печатает.
        self.root.after(4000, self._startup_morning)
        self.root.after(100, self.poll)

    # ---------- оформление ----------

    def _style(self):
        st = ttk.Style()
        if 'vista' in st.theme_names():
            st.theme_use('vista')
        st.configure('.', font=('Segoe UI', 10))
        st.configure('Title.TLabel', font=('Segoe UI', 18, 'bold'), foreground=ACCENT)
        st.configure('Muted.TLabel', foreground='#555555')
        st.configure('Accent.TButton', font=('Segoe UI', 10, 'bold'))
        st.configure('TNotebook.Tab', padding=(14, 6))

    def _menu(self):
        m = tk.Menu(self.root)
        f = tk.Menu(m, tearoff=False)
        f.add_command(label='Папки для поиска…', command=self.folders_window)
        f.add_command(label='Папки входящих…', command=self.inbox_folders_window)
        f.add_command(label='Мастер первого запуска…', command=self.first_run_window)
        f.add_command(label='Обновить индекс', command=self.indexer.scan_now)
        f.add_command(label='Открыть папку данных', command=lambda: os.startfile(str(ROOT)))
        f.add_command(label='Проверить нейросеть', command=self.ai_check)
        f.add_command(label='Сообщить об ошибке…', command=lambda: iskra_bugreport.window(self, VERSION))
        f.add_command(label='Чему научилась Искра…', command=self.learning_window)
        f.add_command(label='Справочник людей…', command=self.people_window)
        f.add_command(label='Журнал исходящих…', command=self.journal_window)
        f.add_command(label='Образцы документов отдела (папка)', command=self.samples_folder)
        f.add_separator(); f.add_command(label='Выход', command=self.close)
        m.add_cascade(label='Файл', menu=f)
        t = tk.Menu(m, tearoff=False)
        t.add_command(label='Утренняя сводка', command=self.morning_window)
        t.add_command(label='Поручения (контроль исполнения)', command=lambda: self.tabs.select(self.control_tab))
        t.add_command(label='Автосвод таблиц за день…', command=self.autosvod_window)
        t.add_command(label='Обзор папок', command=self.overview_window)
        t.add_command(label='Контроль сроков', command=lambda: (self.tabs.select(2), self.refresh_deadlines()))
        t.add_command(label='Word → Excel: собрать таблицу из документов…', command=self.word_to_excel)
        t.add_command(label='Excel → Word: документы по образцу…', command=self.excel_to_word)
        t.add_command(label='Свод таблиц Excel по организациям…', command=self.svod_window)
        t.add_command(label='Сводная Excel по шаблону столбцов…', command=self.consolidation)
        t.add_separator()
        t.add_command(label='Составить телеграмму…', command=self.compose_window)
        t.add_command(label='Редактор телеграмм', command=self.telegrams)
        t.add_command(label='Черновики и версии…', command=self.drafts_window)
        t.add_command(label='Мои реквизиты для телеграмм…', command=self.my_details_window)
        t.add_command(label='Объединить телеграммы…', command=self.merge_telegrams_window)
        t.add_command(label='Перенести изменения (зелёная подсветка) в образец…', command=self.highlight_window)
        t.add_command(label='Пересчитать сутки в документе…', command=self.recount_days_window)
        t.add_command(label='История по объектам (ПУ)…', command=self.pu_history_window)
        t.add_command(label='Проверка орфографии', command=self.spell_window)
        t.add_command(label='Сравнить две редакции…', command=self.compare_files)
        t.add_separator()
        t.add_command(label='Формулы Excel', command=self.formulas)
        t.add_command(label='Конструктор макросов VBA', command=self.macro_builder)
        t.add_command(label='Библиотека макросов VBA', command=self.macros)
        t.add_separator(); t.add_command(label='История действий', command=self.history)
        m.add_cascade(label='Инструменты', menu=t)
        h = tk.Menu(m, tearoff=False)
        h.add_command(label='Что умеет помощник', command=lambda: (self.tabs.select(0), self.say(assistant.HELP)))
        h.add_command(label='О программе', command=lambda: messagebox.showinfo(
            'Искра', f'Искра {VERSION}\nЛокальный поиск и помощник по документам.\nДанные: {ROOT}\nСеть не используется.'))
        m.add_cascade(label='Справка', menu=h)
        self.root.config(menu=m)

    # ---------- вкладка «Помощник» ----------

    def _assistant_tab(self):
        tab = ttk.Frame(self.tabs, padding=8); self.tabs.add(tab, text='Помощник')
        self.chat = tk.Text(tab, wrap='word', font=('Segoe UI', 10), state='disabled', relief='flat',
                            background='#FAFAFA', padx=10, pady=8, spacing1=2, spacing3=4)
        self.chat.tag_configure('me', foreground=ACCENT, font=('Segoe UI', 10, 'bold'), spacing1=10)
        self.chat.tag_configure('iskra', foreground='#1A1A1A')
        self.chat.tag_configure('value', foreground='#1A1A1A', font=('Segoe UI', 10, 'bold'))
        self.chat.tag_configure('proof', foreground='#666666', lmargin1=18, lmargin2=18)
        sc = ttk.Scrollbar(tab, command=self.chat.yview); self.chat.config(yscrollcommand=sc.set)
        tasks = ttk.Frame(tab); tasks.pack(side='bottom', fill='x', pady=(6, 0))
        ttk.Label(tasks, text='Задачи:').pack(side='left', padx=(0, 6))
        for label, cmd in (('Сводка', self.morning_window), ('Автосвод', self.autosvod_window),
                           ('Обзор папок', self.overview_window), ('Свод Excel', self.svod_window), ('Составить телеграмму', self.compose_window),
                           ('Объединить телеграммы', self.merge_telegrams_window), ('Word → Excel', self.word_to_excel),
                           ('Excel → Word', self.excel_to_word), ('Орфография', self.spell_window),
                           ('Формула Excel', self.formulas), ('Макрос', self.macro_builder)):
            ttk.Button(tasks, text=label, command=cmd).pack(side='left', padx=(0, 4))
        quick = ttk.Frame(tab); quick.pack(side='bottom', fill='x', pady=(6, 0))
        ttk.Label(quick, text='По документу:').pack(side='left', padx=(0, 6))
        for label, cmd in (('Кратко', 'кратко'), ('Сроки', 'сроки'), ('Кто', 'кто'), ('Номера', 'номера'),
                           ('Суммы', 'суммы'), ('Все даты', 'даты'), ('Ближайшие сроки', 'ближайшие сроки')):
            ttk.Button(quick, text=label, command=lambda c=cmd: self.ask(c)).pack(side='left', padx=(0, 4))
        row = ttk.Frame(tab); row.pack(side='bottom', fill='x', pady=(6, 0))
        self.question = tk.StringVar()
        q = ttk.Entry(row, textvariable=self.question, font=('Segoe UI', 11))
        q.pack(side='left', fill='x', expand=True); q.bind('<Return>', lambda e: self.ask())
        ttk.Button(row, text='Спросить', style='Accent.TButton', command=self.ask).pack(side='left', padx=(6, 0))
        self.voice_button(row).pack(side='left', padx=(6, 0))
        self.stop_btn = ttk.Button(row, text='Стоп', command=self.ai_stop)
        self.stop_btn.pack(side='left', padx=(6, 0)); self.stop_btn.state(['disabled'])
        sc.pack(side='right', fill='y'); self.chat.pack(fill='both', expand=True)

    def say(self, text, tag='iskra'):
        self.chat.configure(state='normal')
        self.chat.insert('end', text + '\n', tag)
        self.chat.configure(state='disabled'); self.chat.see('end')

    def say_found(self, title, items, empty):
        if not items:
            self.say(empty); return
        self.say(title)
        self.chat.configure(state='normal')
        for f in items[:40]:
            self.chat.insert('end', '• ' + f.value + '\n', 'value')
            self.chat.insert('end', f.context + '\n', 'proof')
        if len(items) > 40:
            self.chat.insert('end', f'…и ещё {len(items) - 40}\n', 'proof')
        self.chat.configure(state='disabled'); self.chat.see('end')

    def ask(self, text=None, routed=None):
        text = (text if text is not None else self.question.get()).strip()
        if not text:
            return
        self.question.set('')
        if routed:
            cmd, arg = routed, text
            if cmd in ('formula', 'macro'):
                cmd, arg = assistant.intent(('формула ' if cmd == 'formula' else 'макрос ') + text)
        else:
            self.say('Вы: ' + text, 'me')
            fresh = assistant.new_topic_text(text)
            if fresh:
                text, self.last_ai = fresh, None          # «новое: …» — забыть прошлый ответ
            cmd, arg = assistant.intent(text)
            # Уточнение к прошлому ответу нейросети («сделай короче», «поменяй срок», «а кто подписал?») —
            # правим тот же ответ, а не ищем заново.
            last = getattr(self, 'last_ai', None)
            if last and time.monotonic() - last['time'] > 1800:
                last = self.last_ai = None              # прошло больше 30 минут — это уже не уточнение
            if (last and last['kind'] in ('ask', 'chat') and assistant.is_question(text)
                    and not re.match(r'^(а|и|но)\b', text.lower())
                    and not assistant.related(text, last['request'] + ' ' + last['answer'])
                    and time.monotonic() - last['time'] > 120):
                last = self.last_ai = None              # новый вопрос на другую тему — искать заново, а не в старых документах
            if (self.ai_ready and cmd in FOLLOW_CMDS and assistant.is_followup(text, (last or {}).get('kind', 'write'))
                    and (last or self.ai_busy())):
                if last is None or self.ai_busy():
                    # Ответ ещё пишется — уточнение применится к нему, как только он будет готов.
                    self.ai_background(lambda cancel: None, lambda _: self.ai_followup(text), label=f'«{text}»')
                else:
                    self.ai_followup(text)
                return
            self.last_ai = None                          # новая просьба — прошлый ответ больше не уточняем
            # Вопрос обычной фразой («кто подписал телеграмму 154?», «когда сдавать сведения?») — ответ нейросети
            # по документам, а не список ФИО/сроков выбранного документа. Короткие команды кнопок — как раньше.
            if (self.ai_ready and cmd in DOC_CMDS and len(text.split()) >= 3
                    and (text.rstrip().endswith('?') or self.selected() is None)):
                cmd, arg = 'ask', text
            # Не похоже ни на одну команду и это не явный поиск — пусть нейросеть поймёт, что нужно.
            if (cmd == 'search' and self.ai_ready and len(text.split()) >= 3
                    and not re.match(r'^(найди|найти|поиск|ищи|покажи|где)\b', text.lower())):
                self.ai_route(text)
                return
        if cmd == 'help':
            self.say(assistant.HELP)
        elif cmd == 'stats':
            self.say(f'В поиске документов: {stats()}. Папки: ' + ('; '.join(self.indexer.folders()) or 'не добавлены'))
        elif cmd == 'formula':
            formula, note = build_formula(arg or text)
            if not formula and self.ai_ready:
                self.ai_formula(arg or text)
                return
            self.say((formula + '\n' if formula else '') + note)
            if formula:
                self.say('Из чего состоит:\n' + explain_formula(formula))
        elif cmd == 'write':
            self.ai_write(arg)
        elif cmd == 'ask':
            if self.ai_ready:
                self.ai_ask(arg)
            else:
                self.query.set(arg); self.find(report=True)
        elif cmd == 'overview':
            self.chat_overview()
        elif cmd == 'morning':
            self.morning_window()
        elif cmd == 'control':
            self.tabs.select(self.control_tab)
            late, soon = self.control.overdue(), self.control.due(3)
            self.say(f'На контроле: {len(self.control.open_items())}, просрочено: {len(late)}, срок в ближайшие 3 дня: {len(soon)}.')
            for t in (late + soon)[:10]:
                self.say(f'• {t["deadline"]} — {t["text"]}' + (f' [{t["responsible"]}]' if t['responsible'] else ''), 'proof')
        elif cmd == 'autosvod':
            lines = self.autosvod.summary_lines()
            if lines:
                for line in lines:
                    self.say(line)
            elif not self.autosvod.settings.get('enabled'):
                self.say('Автосвод не настроен — укажите папку, куда присылают таблицы.', 'proof')
            else:
                self.say('Сегодня таблиц в папке автосвода ещё нет — не прислал никто.', 'proof')
            self.autosvod_window()
        elif cmd == 'journal':
            waiting = self.journal.waiting()
            self.say(f'Исходящих в журнале: {len(self.journal.items)}, ждём ответа: {len(waiting)}, '
                     f'из них просрочено: {sum(1 for _, late in waiting if late)}. Следующий номер: {self.journal.peek()}.')
            for line in self.journal.lines_for_morning()[:10]:
                self.say('• ' + line, 'proof')
            self.journal_window()
        elif cmd == 'directory':
            self.people_window()
        elif cmd == 'report':
            self.morning_window()
            self.say('Сводка открыта — нажмите «Доклад по образцу» (с нейросетью сначала «Главное — нейросеть»).', 'proof')
        elif cmd == 'explain':
            self.say('Что делает формула:\n' + explain_formula(arg))
        elif cmd == 'spell':
            self.spell_window()
        elif cmd == 'highlight':
            self.highlight_window()
        elif cmd == 'pu_history':
            import iskra_pu
            self.pu_history_window(iskra_pu.query_from_phrase(text))
        elif cmd == 'recount_days':
            self.recount_days_window()
        elif cmd == 'word2excel':
            self.word_to_excel()
        elif cmd == 'excel2word':
            self.excel_to_word()
        elif cmd == 'merge':
            self.merge_telegrams_window()
        elif cmd == 'telegram':
            if self.ai_ready and len(text.split()) > 3:
                self.ai_write(text)
            else:
                low = text.lower()
                kind = next((k for w, k in (('запрос', 'Запрос сведений'), ('указани', 'Указание'), ('напомин', 'Напоминание о сроке'),
                                            ('ответ', 'Ответ на входящую'), ('доклад', 'Доклад')) if w in low), None)
                self.compose_window(kind)
        elif cmd == 'reply':
            chosen = self.selected() if re.search(r'\b(эт[уоа]|этот|этой|выбранн\w*|на не[её]|на него|по нему)\b', text.lower()) else None
            if chosen is not None:
                # «ответь на эту телеграмму» — на выбранную в списке слева
                if self.ai_ready:
                    self.ai_write(text + ' (по этому документу)')
                else:
                    self.compose_window('Ответ на входящую', source=self.results[chosen][0])
            elif self.inbox_tree.selection():
                self.tabs.select(self.inbox_tab)
                self.inbox_reply()
            elif self.ai_ready:
                self.ai_write(text)
            else:
                self.compose_window('Ответ на входящую')
        elif cmd == 'svod':
            self.svod_window()
        elif cmd == 'inbox':
            items = [i for i in self.inbox.items() if i['status'] != 'выполнен']
            if not self.inbox.folders():
                self.say('Папка входящих не выбрана: «Файл» → «Папки входящих…». Я буду проверять её каждые 30 секунд.')
            elif not items:
                self.say('Новых и невыполненных входящих нет.')
            else:
                self.say(f'Входящих в работе: {len(items)}.')
                for i in items[:15]:
                    line = f'• {i["kind"]} «{i["name"]}» {i["number"]}'.rstrip()
                    if i['deadline']:
                        line += f' — срок {i["deadline"]}'
                    self.say(line + f' ({i["status"]})')
        elif cmd == 'macro':
            self.macro_builder(arg)
        elif cmd == 'search':
            self.query.set(arg); self.find(report=True)
        elif cmd == 'upcoming':
            self.tabs.select(2); self.refresh_deadlines(report=True)
        else:
            self.about_selected(cmd)

    def about_selected(self, cmd):
        i = self.selected()
        if i is None:
            self.say('Сначала выберите документ в списке слева — найдите его поиском.')
            return
        path = self.results[i][0]
        name = Path(path).name

        def answer(text):
            if cmd == 'summary':
                lines = assistant.summary(text)
                self.say(f'Кратко о «{name}» (главные предложения из текста):' if lines else f'В «{name}» не нашла текста.')
                for line in lines:
                    self.say('• ' + line)
            elif cmd == 'deadlines':
                items = [replace(f, value=f'{f.value} — {assistant.deadline_status(f.sort_key)}') for f in assistant.deadlines(text)]
                self.say_found(f'Сроки в «{name}»:', items, f'В «{name}» сроков не нашла (ищу даты после слов «до», «не позднее», «срок»).')
            elif cmd == 'dates':
                self.say_found(f'Все даты в «{name}»:', assistant.all_dates(text), f'В «{name}» дат нет.')
            elif cmd == 'people':
                self.say_found(f'ФИО в «{name}»:', assistant.people(text), f'В «{name}» не нашла ФИО вида «Иванов И.И.».')
            elif cmd == 'numbers':
                self.say_found(f'Номера в «{name}»:', assistant.numbers(text), f'В «{name}» номеров со знаком № нет.')
            elif cmd == 'money':
                self.say_found(f'Суммы в «{name}»:', assistant.money(text), f'В «{name}» сумм в рублях нет.')
            elif cmd == 'contacts':
                self.say_found(f'Контакты в «{name}»:', assistant.contacts(text), f'В «{name}» телефонов и почты нет.')
            save_history(ROOT, 'Помощник', f'{cmd}: {name}')

        self.with_text(path, answer)

    def with_text(self, path, callback):
        """Read a document in the background (big PDFs take time) and call back in the window thread."""
        if path in self.doc_cache:
            callback(self.doc_cache[path]); return
        self.status.set('Читаю документ…')

        def work():
            try:
                text = extract(Path(path))
                self.events.put(('call', lambda: (self.doc_cache.__setitem__(path, text), self.status.set('Готов к работе'), callback(text))))
            except Exception as exc:
                msg = str(exc)
                self.events.put(('call', lambda: (self.status.set('Готов к работе'), self.say(f'Не удалось прочитать документ: {msg}'))))
        threading.Thread(target=work, daemon=True).start()

    # ---------- вкладка «Документ» ----------

    def _document_tab(self):
        tab = ttk.Frame(self.tabs, padding=8); self.tabs.add(tab, text='Документ')
        self.text = tk.Text(tab, wrap='word', font=('Segoe UI', 10), relief='flat', background='#FAFAFA', padx=10, pady=8)
        self.text.tag_configure('hit', background='#FFE58F')
        actions = ttk.Frame(tab); actions.pack(side='bottom', fill='x', pady=(6, 0))
        ttk.Button(actions, text='Открыть документ', style='Accent.TButton', command=self.open_selected).pack(side='left')
        ttk.Button(actions, text='Открыть папку', command=self.open_folder).pack(side='left', padx=6)
        ttk.Button(actions, text='Скопировать', command=self.copy).pack(side='left')
        ttk.Button(actions, text='Сохранить в Word', command=self.save_word).pack(side='left', padx=6)
        ttk.Button(actions, text='Показать весь текст', command=self.show_full).pack(side='left')
        self.text.pack(fill='both', expand=True)

    def show(self, event=None):
        i = self.selected()
        if i is None:
            return
        path, snip = self.results[i]
        self.text.delete('1.0', 'end')
        self.text.insert('end', f'Файл: {path}\n\n')
        # Fragments come with [найденные] слова: highlight them instead of showing brackets.
        for k, part in enumerate(snip.replace(']', '[').split('[')):
            self.text.insert('end', part, 'hit' if k % 2 else ())

    def show_full(self):
        i = self.selected()
        if i is None:
            return
        path = self.results[i][0]
        self.with_text(path, lambda text: (self.text.delete('1.0', 'end'), self.text.insert('end', f'Файл: {path}\n\n{text}')))

    # ---------- вкладка «Контроль сроков» ----------

    def _deadlines_tab(self):
        tab = ttk.Frame(self.tabs, padding=8); self.tabs.add(tab, text='Контроль сроков')
        top = ttk.Frame(tab); top.pack(fill='x')
        ttk.Label(top, text='Сроки из всех документов: просроченные за неделю и на 30 дней вперёд.').pack(side='left')
        ttk.Button(top, text='Обновить', command=self.refresh_deadlines).pack(side='right')
        cols = ('date', 'status', 'file', 'text')
        self.deadline_view = ttk.Treeview(tab, columns=cols, show='headings', height=16)
        for c, title, w in zip(cols, ('Дата', 'Когда', 'Документ', 'Что сделать'), (90, 130, 200, 520)):
            self.deadline_view.heading(c, text=title); self.deadline_view.column(c, width=w, anchor='w')
        self.deadline_view.tag_configure('late', foreground='#B3261E')
        self.deadline_view.tag_configure('soon', foreground='#8A5300')
        self.deadline_view.pack(fill='both', expand=True, pady=(6, 0))
        self.deadline_view.bind('<Double-Button-1>', self.open_deadline)
        self.deadline_paths = {}

    def refresh_deadlines(self, report=False):
        self.status.set('Проверяю сроки во всех документах…')

        def work():
            try:
                rows = assistant.upcoming(iter_documents())
            except Exception as exc:
                msg = str(exc)
                self.events.put(('call', lambda: self.status.set(f'Не удалось проверить сроки: {msg}'))); return
            self.events.put(('call', lambda: self.show_deadlines(rows, report)))
        threading.Thread(target=work, daemon=True).start()

    def show_deadlines(self, rows, report):
        v = self.deadline_view
        v.delete(*v.get_children()); self.deadline_paths.clear()
        today = date.today()
        for d, status, path, context in rows:
            tag = 'late' if d < today else ('soon' if (d - today).days <= 3 else '')
            item = v.insert('', 'end', values=(d.strftime('%d.%m.%Y'), status, Path(path).name, context), tags=(tag,))
            self.deadline_paths[item] = path
        late = sum(1 for r in rows if r[0] < today)
        self.status.set(f'Сроков: {len(rows)}, просрочено: {late}. Двойной щелчок открывает документ.')
        if report:
            if not rows:
                self.say('Ближайших сроков в документах не нашла.')
            else:
                self.say(f'Нашла сроков: {len(rows)}, из них просрочено: {late}. Список — на вкладке «Контроль сроков». Первые:')
                for d, status, path, context in rows[:5]:
                    self.say(f'• {d.strftime("%d.%m.%Y")} ({status}) — {Path(path).name}: {context[:160]}')

    def open_deadline(self, event=None):
        item = self.deadline_view.focus()
        path = self.deadline_paths.get(item)
        if path:
            self._open(path)

    # ---------- поиск и файлы ----------

    def choose(self):
        folder = filedialog.askdirectory(title='Выберите папку с документами')
        if folder:
            self.indexer.add_folder(folder)
            self.status.set('Папка добавлена. Индексирование запущено в фоне.')

    def poll(self):
        try:
            while True:
                try:
                    event = self.events.get_nowait()
                except queue.Empty:
                    break
                try:
                    if isinstance(event, tuple) and event[0] == 'call':
                        event[1]()
                    elif self.ai_cancel is None:
                        self.status.set(str(event))         # сообщения индексатора не перекрывают «Нейросеть пишет…»
                except Exception as exc:
                    # Одна ошибка не должна останавливать всё окно: пишем в журнал и работаем дальше.
                    from iskra_ai_ui import friendly_error
                    iskra_bugreport.log_crash('Ошибка в окне', type(exc), exc, exc.__traceback__)
                    try:
                        self.say('Не получилось: ' + friendly_error(exc))
                    except Exception:
                        pass
        finally:
            self.root.after(150, self.poll)

    def find(self, report=False):
        q = self.query.get().strip()
        if not q:
            return
        self.status.set('Ищу…')

        self._search_gen = getattr(self, '_search_gen', 0) + 1
        gen = self._search_gen

        def expand(text):
            # Нейросеть подсказывает слова, только если ни обычный, ни смысловой поиск ничего не нашли —
            # и только если она свободна: поиск не должен ждать чужой ответ.
            if not self.ai_ready or iskra_ai._LOCK.locked() or gen != self._search_gen:
                return []
            self.events.put(('call', lambda: self.status.set('По словам не нашла — спрашиваю нейросеть, как это обычно пишут…')))
            try:
                return iskra_ai.search_words(text, timeout=120)
            except iskra_ai.AIError:
                return []

        def work():
            import iskra_smartsearch
            rows, how = iskra_smartsearch.search(q, root=ROOT, ai_expand=expand)
            if gen == self._search_gen:                 # пока искали, человек начал новый поиск — старый не показываем
                self.events.put(('call', lambda: self.show_results(q, rows, report, how)))
        threading.Thread(target=work, daemon=True).start()

    def show_results(self, q, rows, report, how=''):
        import iskra_memory
        rows, raised = iskra_memory.rerank(q, rows)
        if raised:
            how = (how + '; ' if how else '') + 'вверху — что вы уже открывали по этим словам'
        self.last_query = q
        self.results = rows; self.listbox.delete(0, 'end')
        for path, _ in rows:
            self.listbox.insert('end', f'{Path(path).name}   —   {Path(path).parent}')
        more = ' (показаны первые 100)' if len(rows) >= 100 else ''
        self.found_label.set(f'Найдено по «{q}»: {len(rows)}{more}' + (f' ({how})' if how else ''))
        self.status.set('Выберите документ: справа — фрагмент, помощник отвечает по нему.')
        save_history(ROOT, 'Поиск', q)
        if rows:
            self.listbox.selection_set(0); self.show()
        if report:
            if not rows:
                self.say(f'По запросу «{q}» ничего не нашла. Проверьте, добавлена ли папка с документами.')
            else:
                self.say(f'Нашла документов: {len(rows)}. Первый выбран — спросите «кратко», «сроки» или «кто».')
                for path, snip in rows[:5]:
                    self.say(f'• {Path(path).name}: {snip.replace("[", "").replace("]", "")[:160]}')

    def selected(self):
        ids = self.listbox.curselection()
        return ids[0] if ids else None

    def _open(self, path):
        try:
            os.startfile(path)
        except OSError as e:
            messagebox.showerror('Ошибка', str(e))

    def open_selected(self):
        i = self.selected()
        if i is not None:
            path = self.results[i][0]
            if getattr(self, 'last_query', ''):
                import iskra_memory
                try:
                    iskra_memory.opened(self.last_query, path)     # поиск запоминает, что вы выбрали
                except OSError:
                    pass
            self._open(path)

    def open_folder(self):
        i = self.selected()
        if i is not None:
            self._open(str(Path(self.results[i][0]).parent))

    def copy(self):
        self.root.clipboard_clear(); self.root.clipboard_append(self.text.get('1.0', 'end-1c'))
        self.status.set('Скопировано в буфер обмена.')

    def save_word(self):
        from docx import Document
        path = filedialog.asksaveasfilename(defaultextension='.docx', filetypes=[('Word', '*.docx')], initialdir=str(ROOT))
        if path:
            d = Document()
            for line in self.text.get('1.0', 'end-1c').splitlines():
                d.add_paragraph(line)
            d.save(path); self.status.set(f'Сохранено: {path}')

    def _startup_morning(self, tries=0):
        if not self._morning_enabled() or not (ROOT / 'setup.json').exists():
            return
        typing = self.query.get().strip() or self.question.get().strip()
        if typing and tries < 20:
            self.root.after(15000, lambda: self._startup_morning(tries + 1))
            return
        self.maybe_morning()

    def close(self):
        for draft, t in getattr(self, '_open_drafts', []):     # открытые редакторы — сохранить как закрытые
            try:
                if t.winfo_exists():
                    draft.close(t.get('1.0', 'end-1c'))
            except Exception:                                  # noqa: BLE001 — выход из программы не должен сорваться
                pass
        self.stop_hotkey()
        iskra_ai.shutdown()
        self.indexer.stop()
        self.inbox.stop()
        self.root.destroy()

    def folders_window(self):
        win = tk.Toplevel(self.root); win.title('Искра · папки для поиска'); win.geometry('720x380'); win.transient(self.root)
        ttk.Label(win, text='Эти папки проверяются при запуске и каждые 15 минут. Сами файлы не изменяются.').pack(anchor='w', padx=12, pady=8)
        box = tk.Listbox(win, font=('Segoe UI', 10)); box.pack(fill='both', expand=True, padx=12)

        def refresh():
            box.delete(0, 'end')
            for f in self.indexer.folders():
                box.insert('end', f + ('' if Path(f).is_dir() else '   (сейчас недоступна)'))

        def add():
            self.choose(); refresh()

        def remove():
            ix = box.curselection()
            if not ix:
                return
            folder = self.indexer.folders()[ix[0]]
            if not messagebox.askyesno('Убрать папку', f'Убрать из поиска?\n{folder}\n\nСами файлы не удаляются.', parent=win):
                return
            self.indexer.remove_folder(folder)
            removed = forget_folder(folder)
            self.status.set(f'Папка убрана из поиска, записей удалено: {removed}')
            refresh()

        buttons = ttk.Frame(win); buttons.pack(fill='x', padx=12, pady=8)
        ttk.Button(buttons, text='Добавить папку…', style='Accent.TButton', command=add).pack(side='left')
        ttk.Button(buttons, text='Убрать выбранную', command=remove).pack(side='left', padx=8)
        ttk.Button(buttons, text='Закрыть', command=win.destroy).pack(side='right')
        refresh()

    def consolidation(self):
        template=filedialog.askopenfilename(title='Общий Excel-файл: заголовки в первой строке',filetypes=[('Excel','*.xlsx')])
        if not template:return
        sources=filedialog.askopenfilenames(title='Выберите исходные Excel-файлы',filetypes=[('Excel','*.xlsx *.xlsm')])
        if not sources:return
        output=filedialog.asksaveasfilename(title='Сохранить новую сводную таблицу',defaultextension='.xlsx',initialdir=str(ROOT),filetypes=[('Excel','*.xlsx')])
        if not output:return
        win=tk.Toplevel(self.root);win.title('Сбор сводной · настройки');win.geometry('580x230')
        ttk.Label(win,text='Необязательно: название столбца с уникальным номером / кодом').pack(anchor='w',padx=12,pady=(15,5))
        value=tk.StringVar();ttk.Entry(win,textvariable=value).pack(fill='x',padx=12)
        ttk.Label(win,text='Совпадение колонок — по точным заголовкам (без учёта регистра).\nПроверьте лист «Проверка» в итоговом файле.',wraplength=540).pack(anchor='w',padx=12,pady=10)
        def run():
            try:
                result=consolidate(template,sources,output,value.get().strip() or None)
                save_history(ROOT,'Сводная Excel',f"{len(sources)} файлов, {result['rows']} строк")
                messagebox.showinfo('Сводная готова',f"Собрано строк: {result['rows']}\nЗамечаний: {result['issues']}\nПропущено файлов: {result['skipped_files']}\n\n{output}\n\nОткройте лист «Проверка».",parent=win)
                win.destroy()
            except Exception as exc:messagebox.showerror('Ошибка объединения',str(exc),parent=win)
        ttk.Button(win,text='Собрать и проверить',command=run).pack(pady=8)

    def compare_files(self):
        first=filedialog.askopenfilename(title='Старая редакция',filetypes=[('Документы','*.docx *.xlsx *.pptx *.pdf *.txt *.md *.csv')])
        if not first:return
        second=filedialog.askopenfilename(title='Новая редакция',filetypes=[('Документы','*.docx *.xlsx *.pptx *.pdf *.txt *.md *.csv')])
        if not second:return
        try:
            diff=compare_texts(extract(Path(first)),extract(Path(second)))
        except Exception as exc:
            messagebox.showerror('Ошибка сравнения',str(exc));return
        win=tk.Toplevel(self.root);win.title('Искра · сравнение редакций');win.geometry('900x650')
        area=tk.Text(win,wrap='none',font=('Consolas',10));area.pack(fill='both',expand=True,padx=10,pady=10)
        area.insert('1.0',diff)
        def save():
            path=filedialog.asksaveasfilename(parent=win,defaultextension='.txt',filetypes=[('Текст','*.txt')])
            if path:Path(path).write_text(diff,encoding='utf-8')
        ttk.Button(win,text='Сохранить сравнение',command=save).pack(pady=8)
        save_history(ROOT,'Сравнение',f'{Path(first).name} / {Path(second).name}')

    def history(self):
        win=tk.Toplevel(self.root);win.title('Искра · история действий');win.geometry('800x500')
        area=tk.Text(win,wrap='word',font=('Segoe UI',10));area.pack(fill='both',expand=True,padx=12,pady=12)
        for item in read_history(ROOT):
            area.insert('end',f"{item['timestamp']} · {item['action']} · {item['details']}\n")
        area.configure(state='disabled')

    def macros(self):
        win=tk.Toplevel(self.root);win.title('Искра · библиотека макросов (офлайн)');win.geometry('900x680')
        ttk.Label(win,text='Готовые VBA-макросы для Word и Excel',font=('Segoe UI',13,'bold')).pack(anchor='w',padx=12,pady=(12,3))
        ttk.Label(win,text='Макросы НЕ запускаются автоматически. Сначала проверьте на копии документа.',wraplength=850).pack(anchor='w',padx=12)
        area=ttk.Frame(win);area.pack(fill='both',expand=True,padx=12,pady=10)
        names=tk.Listbox(area,width=43,exportselection=False);names.pack(side='left',fill='y')
        code=tk.Text(area,wrap='none',font=('Consolas',10));code.pack(side='left',fill='both',expand=True,padx=(8,0))
        for m in MACROS:names.insert('end',f'{m.app}: {m.title}')
        status=tk.StringVar(value='Выберите макрос слева. Его код можно проверить перед сохранением.')
        ttk.Label(win,textvariable=status,wraplength=860).pack(anchor='w',padx=12,pady=4)
        def current():
            ix=names.curselection()
            return MACROS[ix[0]] if ix else None
        def show(event=None):
            m=current()
            if not m:return
            code.delete('1.0','end');code.insert('1.0',m.code)
            status.set(m.description)
        def copy():
            m=current()
            if m:
                win.clipboard_clear();win.clipboard_append(code.get('1.0','end-1c'))
                status.set('Код скопирован. Перед запуском проверьте его и используйте копию файла.')
        def save():
            m=current()
            if not m:return
            path=filedialog.asksaveasfilename(parent=win,defaultextension='.bas',initialfile='Iskra_'+m.key+'.bas',filetypes=[('Модуль VBA','*.bas')])
            if path:
                try:
                    export_macro(m,path);status.set('Модуль сохранён: '+path)
                except (OSError,ValueError) as e:messagebox.showerror('Ошибка',str(e),parent=win)
        buttons=ttk.Frame(win);buttons.pack(fill='x',padx=12,pady=10)
        ttk.Button(buttons,text='Копировать код',command=copy).pack(side='left')
        ttk.Button(buttons,text='Сохранить .bas',command=save).pack(side='left',padx=8)
        ttk.Button(buttons,text='Как подключить',command=lambda:messagebox.showinfo('Инструкция','Word / Excel: Alt+F11 → Файл → Импорт файла → выберите .bas.\nСохраните документ как .docm или .xlsm.\nМакросы могут быть запрещены политиками организации.\nВсегда сначала проверяйте на копии.',parent=win)).pack(side='left')
        names.bind('<<ListboxSelect>>',show)
        names.selection_set(0);show()


def selftest():
    """Iskra.exe --selftest: checks the bundled parts of a built program, writes data/selftest.txt."""
    import tempfile
    from docx import Document
    from iskra_spell import SpellChecker
    lines, ok = [], True

    def check(name, fn):
        nonlocal ok
        try:
            lines.append(f'OK   {name}: {fn()}')
        except Exception as exc:
            ok = False
            lines.append(f'FAIL {name}: {exc!r}')

    def docx():
        p = Path(tempfile.mkdtemp()) / 'a.docx'
        d = Document(); d.add_paragraph('ФИО: Иванов И.И.'); d.save(p)
        return extract(p).strip()

    check('словарь', lambda: [i.text for i in SpellChecker().check('Отчот направить в течении недели.')])
    check('Word', docx)
    check('формула', lambda: build_formula('сумма столбца C где в столбце A "Москва"')[0])
    check('помощник', lambda: assistant.intent('объедини телеграммы'))
    check('макрос', lambda: (lambda c: f'{len(c[1])} шага, замечаний {len(__import__("iskra_vba").lint(c[0]))}')(
        __import__('iskra_vba').generate('удалить пустые строки, выделить повторы в столбце B')))
    def telegram():
        from iskra_telegram import merge, save_docx
        a = 'Прошу в течении трёх дней представить отчёт о проверке.\nНачальник отдела И.И. Иванов'
        b = 'Прошу в течении 3 дней представить отчёт о проверке.\nНачальник отдела И.И. Иванов'
        m = merge([('1.txt', a), ('2.txt', b)], fix=__import__('iskra_spell').fix_typical)
        out = Path(tempfile.mkdtemp()) / 't.docx'
        save_docx(m.text(), out)
        return f'{len(m.points)} п., Word {out.stat().st_size} байт'

    check('телеграммы', telegram)

    def svod():
        from openpyxl import Workbook
        from iskra_svod import svod_files
        d = Path(tempfile.mkdtemp())
        for name, rows in (('a.xlsx', [['Наименование', 'По списку', 'На лицо'], ['Госпиталь', 10, 9]]),
                           ('b.xlsx', [['Организация', 'по списку', 'на лицо'], ['Модуль №1', 5, 5]])):
            wb = Workbook(); [wb.active.append(r) for r in rows]; wb.save(d / name)
        r = svod_files([d / 'a.xlsx', d / 'b.xlsx'], d / 'свод.xlsx', ['Госпиталь', 'Модуль 1'])
        return f'{r["orgs"]} орг., столбцы: {", ".join(r["columns"])}'

    def inbox():
        from iskra_inbox import Inbox
        d = Path(tempfile.mkdtemp())
        box = Inbox(d / 'data'); box.add_folder(d)
        (d / 'тлг.txt').write_text('ТЕЛЕГРАММА\nПредставить сведения до 25.12.2099.', encoding='utf-8')
        found = box.scan(settle=0)
        return f'{found[0]["kind"]}, срок {found[0]["deadline"]}'

    def control():
        from iskra_control import Control
        c = Control(Path(tempfile.mkdtemp()))
        added = c.add_from_document('ТЕЛЕГРАММА\n1. Представить сведения до 25.12.2099. Ответственный — майор С.С. Сидоров.', 'т.txt')
        return f'{len(added)} поруч., ответственный: {added[0]["responsible"]}'

    def autosvod():
        from openpyxl import Workbook
        from iskra_autosvod import AutoSvod
        d = Path(tempfile.mkdtemp())
        wb = Workbook(); [wb.active.append(r) for r in (['Наименование', 'По списку'], ['Госпиталь', 10])]; wb.save(d / 'г.xlsx')
        a = AutoSvod(d / 'data'); a.configure(d, d / 'out', True, ['Госпиталь', 'Полигон'])
        r = a.build()
        return f'не прислали: {", ".join(r["missing"])}'

    check('поручения', control)
    check('автосвод', autosvod)
    check('свод Excel', svod)
    check('входящие', inbox)

    def people():
        import iskra_people as pp
        if pp.morph() is None:
            raise RuntimeError('словарь склонений не загрузился')
        got = pp.addressee({'fio': 'Иванов И.И.', 'rank': 'старший лейтенант', 'position': 'начальник госпиталя'})
        assert got == ['Начальнику госпиталя', 'старшему лейтенанту', 'Иванову И.И.'], got
        return ' / '.join(got)
    check('справочник и падежи', people)

    def precheck():
        import iskra_precheck as pc
        issues = pc.check('1. Представить до 01.01.2020.\n3. Доложить.', today=date(2026, 10, 22))
        assert any('уже прошёл' in m for _, m in issues) and any('Нумерация' in m for _, m in issues), issues
        return f'замечаний: {len(issues)}'
    check('проверка перед отправкой', precheck)

    def pdf_render():
        import pypdfium2 as pdfium
        return f'pdfium {pdfium.PdfDocument.__name__}'
    check('чтение PDF для сканов', pdf_render)
    (ROOT / 'selftest.txt').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    return 0 if ok else 1


def selftest_ai():
    """Iskra.exe --selftest-ai: настоящая модель на этом компьютере; ответы и время пишутся в data/selftest_ai.txt."""
    import time
    lines, ok = [], True
    moved = iskra_ai.gather_parts()
    if moved:
        lines.append(f'Части модели из соседних папок перенесены к программе: {moved}')
    if not iskra_ai.available():
        (ROOT / 'selftest_ai.txt').write_text('FAIL нейросеть не найдена (engine/models)\n', encoding='utf-8')
        return 1
    lines.append(f'Модель: {iskra_ai.model_name()}')
    doc = ('тлг 154.txt', 'ТЕЛЕГРАММА № 154/2 от 20.10.2026. От кого: Штаб округа. '
           'Представить сведения о численности личного состава до 25.10.2026. Обеспечить готовность техники.')

    def check(name, fn, test):
        nonlocal ok
        t = time.time()
        try:
            out = fn()
            text = out[0] if isinstance(out, tuple) else out
            good = test(text)
            lines.append(f'{"OK  " if good else "FAIL"} {name} ({time.time() - t:.0f} с):\n{text}\n')
            ok = ok and good
        except Exception as exc:
            ok = False
            lines.append(f'FAIL {name}: {exc!r}')

    def russian(text):
        letters = [c for c in text if c.isalpha()]
        return len(text) > 40 and '<think>' not in text and sum('а' <= c.lower() <= 'я' or c in 'ёЁ' for c in letters) / max(1, len(letters)) > 0.7

    check('ответ по документу', lambda: iskra_ai.answer('Какой срок представления сведений о численности?', [doc]),
          lambda a: '25.10' in a or '25 октября' in a)
    check('составление текста', lambda: iskra_ai.write(
        'Составь телеграмму-ответ в штаб округа на телеграмму 154/2: сведения о численности личного состава будут представлены в срок', [doc]),
        lambda w: russian(w) and '154' in w)
    check('исправление ошибок', lambda: iskra_ai.edit('Прошу в течении трех дней предоставить отчет о проверке.', 'Исправить ошибки'),
          lambda e: 'течение' in e.lower())
    check('формула Excel', lambda: iskra_ai.formula('сумма значений столбца C, где в столбце A написано Госпиталь'),
          lambda f: f.startswith('='))
    check('понимание фразы', lambda: str(iskra_ai.route('надо пять табличек по организациям свести в одну')),
          lambda r: "'svod'" in r)
    check('поля из документа', lambda: str(iskra_ai.extract_fields(
        'Справка. Выдана Петрову Петру Петровичу в том, что он проходит службу в должности начальника склада.',
        ['ФИО', 'Должность'])), lambda r: 'Петров' in r)
    check('макрос VBA', lambda: (lambda c: c[0] + f'\nзамечаний проверки: {len(c[1])}')(
        iskra_ai.vba('удалить пустые строки на активном листе', 'excel')), lambda c: 'Sub ' in c and 'замечаний проверки: 0' in c)
    draft = ('ТЕЛЕГРАММА\n\nКому: Штаб округа\n\nТекст:\nНа № 154/2 от 20.10.2026. Сведения о численности личного состава '
             'будут представлены до 25.10.2026.\n\nПодпись:\nНачальник штаба\nполковник\nП.П. Сидоров')
    check('выписка из документа (когда документов много)', lambda: iskra_ai.notes('ответ на телеграмму 154', doc[0], doc[1]),
          lambda r: '25.10.2026' in r and '154' in r)
    check('уточнение прошлого ответа', lambda: iskra_ai.revise('телеграмма-ответ в штаб округа', draft, 'поменяй срок на 30.10.2026'),
          lambda r: '30.10.2026' in r and '25.10.2026' not in r and '154/2' in r and 'Начальник штаба' in r and 'Сидоров' in r)
    check('продолжение разговора', lambda: iskra_ai.followup(
        [('какой срок представления сведений по телеграмме 154?', 'До 25.10.2026.\nИсточник: тлг 154.txt')], 'а от кого она?', [doc]),
        lambda r: 'штаб' in r.lower())

    def by_template():
        import tempfile
        from docx import Document
        sample = Path(tempfile.mkdtemp()) / 'obrazec.docx'
        d = Document()
        for line in ('ТЕЛЕГРАММА', 'Командиру войсковой части 12345', '',
                     'Во исполнение указаний штаба округа приказываю представить сведения о технике до 01.03.2024.', '',
                     'Начальник штаба', 'полковник П.П. Сидоров'):
            d.add_paragraph(line)
        d.save(sample)
        iskra_ai.TELEGRAM_TEMPLATE = str(sample)
        try:
            return iskra_ai.write('Составь телеграмму начальнику госпиталя: представить сведения о численности личного '
                                  'состава до 30.10.2026', [])[0]
        finally:
            iskra_ai.TELEGRAM_TEMPLATE = None
    check('телеграмма по шаблону отдела', by_template,
          lambda r: 'Кому:' not in r and 'Сидоров' in r and '30.10.2026' in r and '01.03.2024' not in r and 'госпитал' in r.lower())
    lines.append(f'Режим: {iskra_ai.mode()}')
    iskra_ai.shutdown()
    (ROOT / 'selftest_ai.txt').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    return 0 if ok else 1


def selftest_voice(wav):
    """Iskra.exe --selftest-voice <запись.wav>: распознавание речи (английская фраза — голос Windows) → data/selftest_voice.txt."""
    import time
    import iskra_voice
    lines, ok = [f'Голос установлен: {iskra_voice.available()}'], iskra_voice.available()
    if ok:
        t = time.time()
        try:
            text = iskra_voice.transcribe(wav, lang='en')
        except Exception as exc:
            text = f'ОШИБКА {exc!r}'
        ok = 'hello' in text.lower() or 'test' in text.lower()
        lines.append(f'{"OK  " if ok else "FAIL"} распознано за {time.time() - t:.1f} с: {text}')
    (ROOT / 'selftest_voice.txt').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    return 0 if ok else 1


def selftest_ocr(folder):
    """Iskra.exe --selftest-ocr <папка>: распознавание тестовых сканов (картинка и PDF) → data/selftest_ocr.txt."""
    import iskra_ocr
    lines, ok = [f'Распознавание установлено: {iskra_ocr.available()}'], iskra_ocr.available()
    for f in sorted(Path(folder).iterdir()):
        try:
            text = extract(f)
        except Exception as exc:
            text = f'ОШИБКА {exc!r}'
        good = '217/4' in text and '28.10.2026' in text and 'численност' in text.lower()
        ok = ok and good
        lines.append(f'{"OK  " if good else "FAIL"} {f.name}:\n{text}\n')
    (ROOT / 'selftest_ocr.txt').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    return 0 if ok else 1


if __name__ == '__main__':
    import sys
    if '--selftest-voice' in sys.argv:
        sys.exit(selftest_voice(sys.argv[sys.argv.index('--selftest-voice') + 1]))
    if '--selftest-ocr' in sys.argv:
        sys.exit(selftest_ocr(sys.argv[sys.argv.index('--selftest-ocr') + 1]))
    if '--selftest-ai' in sys.argv:
        sys.exit(selftest_ai())
    if '--selftest' in sys.argv:
        sys.exit(selftest())
    root = tk.Tk()
    App(root)
    root.mainloop()
