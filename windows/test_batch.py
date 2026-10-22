import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook

from iskra_batch import discover_fields, collect_tables, mail_merge, template_fields


def card(path, fio, rank, unit, phone=None):
    d = Document()
    d.add_paragraph('КАРТОЧКА')
    d.add_paragraph(f'ФИО: {fio}')
    d.add_paragraph(f'Звание: {rank}')
    d.add_paragraph('В целях учёта: сведения проверены.')  # предложение, а не поле
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = 'Подразделение'
    t.rows[0].cells[1].text = unit
    if phone:
        d.add_paragraph(f'Телефон: {phone}')
    d.save(path)


class Batch(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())

    def test_discovers_fields_in_twenty_documents(self):
        paths = []
        for i in range(20):
            p = self.dir / f'card{i}.docx'
            card(p, f'Иванов {i}', 'сержант', f'{i} рота', phone='123' if i < 5 else None)
            paths.append(p)
        fields = discover_fields(paths)
        self.assertEqual(fields[:3], [('ФИО', 20), ('Звание', 20), ('Подразделение', 20)])
        self.assertIn(('Телефон', 5), fields)
        self.assertNotIn('В целях учёта', [f for f, _ in fields])

    def test_collects_word_tables_into_one_sheet(self):
        for i, rows in enumerate([[['1', 'Петров', '5']], [['2', 'Сидоров', '7'], ['3', 'Орлов', '2']]]):
            d = Document()
            t = d.add_table(rows=1, cols=3)
            for c, h in zip(t.rows[0].cells, ['№', 'Фамилия', 'Количество']):
                c.text = h
            for r in rows:
                for c, v in zip(t.add_row().cells, r):
                    c.text = v
            d.save(self.dir / f't{i}.docx')
        result = collect_tables([self.dir / 't0.docx', self.dir / 't1.docx'], self.dir / 'out.xlsx')
        self.assertEqual((result['sheets'], result['rows']), (1, 3))
        ws = load_workbook(self.dir / 'out.xlsx')['Таблица']
        self.assertEqual([c.value for c in ws[1]], ['№', 'Фамилия', 'Количество', 'Источник: файл', 'Источник: таблица'])
        self.assertEqual([c.value for c in ws[4]], [3, 'Орлов', 2, 't1.docx', 1])   # числа — числами (СУММ работает)

    def test_mail_merge_one_document_per_row(self):
        tpl = Document()
        tpl.add_paragraph('Справка')
        p = tpl.add_paragraph('Выдана ')
        p.add_run('{ФИО}').bold = True
        p.add_run(' в том, что он служит в {Подразделение} с {Дата}.')
        tpl.save(self.dir / 'Образец.docx')
        self.assertEqual(template_fields(self.dir / 'Образец.docx'), ['ФИО', 'Подразделение', 'Дата'])
        wb = Workbook(); ws = wb.active
        ws.append(['ФИО', 'Подразделение', 'Дата', 'Лишний'])
        ws.append(['Иванов И.И.', '1 рота', datetime(2026, 1, 15), 'x'])
        ws.append(['Петров П.П.', '2 рота', datetime(2025, 3, 1), 'y'])
        wb.save(self.dir / 'Список.xlsx')
        result = mail_merge(self.dir / 'Список.xlsx', self.dir / 'Образец.docx', self.dir / 'out', name_field='ФИО')
        self.assertEqual(result['created'], 2)
        text = '\n'.join(p.text for p in Document(self.dir / 'out' / 'Иванов И.И.docx').paragraphs)
        self.assertIn('Выдана Иванов И.И. в том, что он служит в 1 рота с 15.01.2026.', text)

    def test_mail_merge_reports_missing_columns(self):
        tpl = Document(); tpl.add_paragraph('{Должность}'); tpl.save(self.dir / 'tpl.docx')
        wb = Workbook(); wb.active.append(['ФИО']); wb.active.append(['А']); wb.save(self.dir / 'x.xlsx')
        with self.assertRaisesRegex(ValueError, 'Должность'):
            mail_merge(self.dir / 'x.xlsx', self.dir / 'tpl.docx', self.dir / 'o')


if __name__ == '__main__':
    unittest.main()
