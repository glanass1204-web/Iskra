import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt

import iskra_ai as ai
import iskra_telegram as tg


def make(path, lines):
    doc = Document()
    for text, bold, indent in lines:
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(14)
        if indent:
            p.paragraph_format.first_line_indent = Cm(1.25)
    doc.save(path)
    return path


class TemplateTest(unittest.TestCase):
    def setUp(self):
        d = Path(tempfile.mkdtemp())
        self.marked = make(d / 'с метками.docx', [('СРОЧНО', True, False), ('Командирам частей {Кому}', True, False), ('', False, False),
                                                  ('{Текст}', False, True), ('', False, False), ('{Подпись}', False, False)])
        self.sample = make(d / 'образец.docx', [('ТЕЛЕГРАММА', True, False), ('Командиру войсковой части 12345', True, False),
                                                ('', False, False),
                                                ('Во исполнение указаний штаба округа приказываю представить сведения о технике до 01.03.2024.', False, True),
                                                ('', False, False), ('Начальник штаба', False, False), ('полковник П.П. Сидоров', False, False)])
        self.out = d / 'out.docx'

    def test_template_text_and_render(self):
        self.assertTrue(tg.has_placeholders(tg.template_text(self.marked)))
        self.assertFalse(tg.has_placeholders(tg.template_text(self.sample)))
        text = tg.render_template(self.marked, {'кому': 'госпиталя', 'текст': ['1. Представить сведения.', '2. Доложить.']})
        self.assertIn('Командирам частей госпиталя', text)
        self.assertIn('1. Представить сведения.\n2. Доложить.', text)
        self.assertIn('[УТОЧНИТЬ: подпись]', text)

    def test_ai_prompt_uses_template(self):
        prompt = ai.template_format(self.sample)
        self.assertIn('Во исполнение указаний штаба округа', prompt)
        self.assertIn('НЕ переноси его содержание', prompt)
        self.assertIn('{Текст}', ai.template_format(self.marked))
        old = ai.TELEGRAM_TEMPLATE
        ai.TELEGRAM_TEMPLATE = str(self.sample)
        try:
            self.assertIn('ОБРАЗЕЦ', ai.telegram_format())
        finally:
            ai.TELEGRAM_TEMPLATE = old
        self.assertIn('Кому:', ai.telegram_format())                # без шаблона — стандартный вид

    def test_save_keeps_template_formatting(self):
        text = ('ТЕЛЕГРАММА\nНачальнику госпиталя\n\nВо исполнение указаний штаба округа приказываю представить сведения '
                'о численности личного состава до 30.10.2026 и доложить о готовности.\n\nНачальник штаба\nполковник П.П. Сидоров\n\n'
                '--- ИСТОЧНИКИ (НЕ ЧАСТЬ ТЕЛЕГРАММЫ) ---\n1. тлг.txt')
        tg.save_docx(text, self.out, str(self.sample))
        paras = Document(self.out).paragraphs
        texts = [p.text for p in paras]
        self.assertNotIn('01.03.2024', ' '.join(texts))           # содержание образца не осталось
        self.assertNotIn('ИСТОЧНИКИ', ' '.join(texts))
        head = next(p for p in paras if p.text == 'ТЕЛЕГРАММА')
        self.assertTrue(head.runs[0].bold)
        body = next(p for p in paras if p.text.startswith('Во исполнение'))
        self.assertAlmostEqual(body.paragraph_format.first_line_indent, Cm(1.25), delta=2000)
        self.assertFalse(body.runs[0].bold)

    def test_labeled_text_into_marked_template(self):
        text = 'ТЕЛЕГРАММА\n\nКому: госпиталя\n\nТекст:\nПредставить сведения до 30.10.2026.\n\nПодпись:\nНачальник штаба'
        tg.save_docx(text, self.out, str(self.marked))
        texts = [p.text for p in Document(self.out).paragraphs]
        self.assertIn('Командирам частей госпиталя', texts)
        self.assertIn('Представить сведения до 30.10.2026.', texts)


if __name__ == '__main__':
    unittest.main()
