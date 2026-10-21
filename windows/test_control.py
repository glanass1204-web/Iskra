import os
import tempfile
import time
import unittest
from datetime import date, datetime
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook

import iskra_control as control
import iskra_morning as morning
from iskra_autosvod import AutoSvod, dynamics

TODAY = date(2026, 10, 21)
TG = '''ТЕЛЕГРАММА
От кого: Штаб округа
Номер: 154/2
Текст:
Довожу до вашего сведения, что проводится проверка готовности.
1. Представить сведения о численности личного состава до 25.10.2026.
2. Начальнику госпиталя обеспечить готовность техники к 22.10.2026.
3. Провести занятия с личным составом до 15.10.2026. Ответственный — майор С.С. Сидоров.
Подпись: Начальник штаба полковник И.И. Иванов'''


class ControlTest(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp())
        self.c = control.Control(self.root)

    def test_extract(self):
        tasks = control.extract(TG, 'тлг 154.docx', TODAY)
        self.assertEqual([(t['deadline'], t['responsible']) for t in tasks],
                         [('25.10.2026', ''), ('22.10.2026', 'Начальнику госпиталя'), ('15.10.2026', 'майор С.С. Сидоров')])
        self.assertNotIn('Довожу', ' '.join(t['text'] for t in tasks))
        self.assertEqual(tasks[0]['number'], '154/2')

    def test_add_once_status_and_lists(self):
        added = self.c.add_from_document(TG, 'тлг 154.docx', today=TODAY)
        self.assertEqual(len(added), 3)
        self.assertEqual(self.c.add_from_document(TG, 'тлг 154.docx', today=TODAY), [])     # повторно не добавляются
        self.assertEqual([t['deadline'] for t in self.c.overdue(TODAY)], ['15.10.2026'])
        self.assertEqual([t['deadline'] for t in self.c.due(1, TODAY)], ['22.10.2026'])
        self.c.update(added[2]['id'], status='выполнено')
        self.assertEqual(self.c.overdue(TODAY), [])
        again = control.Control(self.root)                                                  # сохраняется между запусками
        self.assertEqual(len(again.open_items()), 2)

    def test_remind_once_a_day(self):
        self.c.add_from_document(TG, 'тлг 154.docx', today=TODAY)
        first = self.c.to_remind(TODAY, days=4)
        self.assertEqual(sorted(t['deadline'] for t in first), ['15.10.2026', '22.10.2026', '25.10.2026'])
        self.assertEqual(self.c.to_remind(TODAY, days=4), [])

    def test_ai_tasks_checked(self):
        rule = control.extract(TG, 'т', TODAY)
        ai = [{'что': 'Доложить об исполнении в штаб', 'срок': '30.10.2026', 'кто': 'генерал Петров'},
              {'что': 'Представить сведения о численности личного состава до 25.10.2026.', 'срок': '25.10.2026', 'кто': ''}]
        merged = control.merge_ai(rule, ai, TG)
        self.assertEqual(len(merged), 4)                       # повтор отброшен
        self.assertEqual((merged[3]['deadline'], merged[3]['responsible']), ('', ''))   # даты и человека нет в документе

    def test_reports(self):
        self.c.add_from_document(TG, 'тлг 154.docx', today=TODAY)
        x = self.c.report_xlsx(self.root / 'r.xlsx', today=TODAY)
        ws = load_workbook(x).active
        self.assertEqual(ws['E3'].value, 'просрочено на 6 дн.')
        d = Document(self.c.report_docx(self.root / 'r.docx', today=TODAY))
        self.assertIn('просрочено: 1', '\n'.join(p.text for p in d.paragraphs))
        self.assertEqual(len(d.tables[0].rows), 4)


def table(path, rows):
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    wb.save(path)
    return path


class AutoSvodTest(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp())
        self.box = self.root / 'Сведения'; self.box.mkdir()
        self.out = self.root / 'Своды'
        self.a = AutoSvod(self.root / 'data')
        self.a.configure(self.box, self.out, True, ['Госпиталь', 'Полигон', 'Модуль 1'])

    def test_day_by_day(self):
        today = date.today()
        yesterday = date.fromordinal(today.toordinal() - 1)
        f1 = table(self.box / 'госпиталь.xlsx', [['Наименование', 'По списку', 'На лицо'], ['Госпиталь', 120, 115]])
        f2 = table(self.box / 'полигон.xlsx', [['Наименование', 'По списку', 'На лицо'], ['Полигон', 80, 78]])
        old = time.mktime(datetime.combine(yesterday, datetime.min.time()).timetuple()) + 9 * 3600
        for f in (f1, f2):
            os.utime(f, (old, old))
        r1 = self.a.build(yesterday)
        self.assertEqual(r1['missing'], ['Модуль 1'])
        self.assertEqual(r1['changes'], [])
        self.assertIsNone(self.a.check(today))                 # сегодня ещё никто не прислал
        # Сегодня полигон прислал новые цифры, госпиталь ещё нет.
        table(self.box / 'полигон.xlsx', [['Наименование', 'По списку', 'На лицо'], ['Полигон', 80, 73]])
        r2 = self.a.check(today)
        self.assertEqual(r2['missing'], ['Госпиталь', 'Модуль 1'])
        self.assertEqual(r2['changes'], [('Полигон', 'На лицо', 78, 73, -5)])
        self.assertIn('Полигон: На лицо −5 (было 78, стало 73)', ' '.join(self.a.summary_lines(r2)))
        wb = load_workbook(r2['output'])
        self.assertEqual(wb.sheetnames[:3], ['Свод', 'Динамика', 'Не прислали'])
        self.assertIsNone(self.a.check(today))                 # ничего не изменилось — не пересобирает

    def test_dynamics(self):
        self.assertEqual(dynamics({'Модуль №1': {'x': 5}}, {'модуль 1': {'x': 7}}), [('Модуль №1', 'x', 7, 5, -2)])


class MorningTest(unittest.TestCase):
    def test_sections(self):
        root = Path(tempfile.mkdtemp())
        c = control.Control(root)
        c.add_from_document(TG, 'тлг 154.docx', today=TODAY)
        inbox = [{'kind': 'Телеграмма', 'name': 'тлг 155.docx', 'number': '№ 155', 'sender': 'Штаб', 'deadline': '21.10.2026',
                  'added': '21.10.2026 07:30', 'status': 'новый'},
                 {'kind': 'Приказ', 'name': 'приказ 7.docx', 'number': '', 'sender': '', 'deadline': '',
                  'added': '18.10.2026 10:00', 'status': 'новый'}]
        s = morning.build(inbox, c, [], ['Свод на 21.10.2026: прислали 2 файл(ов)'], today=TODAY)
        n = morning.counts(s)
        self.assertEqual(n['ПРИШЛО'], 1)
        self.assertEqual(n['СРОКИ СЕГОДНЯ И ЗАВТРА'], 2)     # поручение на 22.10 и телеграмма на 21.10
        dup = [{'kind': 'Телеграмма', 'name': 'тлг 154.docx', 'deadline': '22.10.2026', 'added': '20.10.2026 07:30', 'status': 'в работе'}]
        self.assertEqual(morning.counts(morning.build(dup, c, today=TODAY))['СРОКИ СЕГОДНЯ И ЗАВТРА'], 1)   # без повтора
        self.assertEqual(n['ПРОСРОЧЕНО'], 1)
        self.assertEqual(n['НЕ ОБРАБОТАНО (пришло раньше, ещё не открыто)'], 1)
        t = morning.text(s, TODAY)
        self.assertIn('СВОДКА НА 21.10.2026', t)
        self.assertIn('просрочено на 6 дн.', t)


if __name__ == '__main__':
    unittest.main()
