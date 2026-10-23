"""Без своего шаблона — только текст документа: без шапки («ТЕЛЕГРАММА», «Кому:», «Дата:») и без подписи."""
import tempfile
import unittest
from pathlib import Path

import iskra_ai as ai
import iskra_telegram as tg


class PlainText(unittest.TestCase):
    ME = {'signature': 'Начальник штаба\nполковник И.И. Иванов'}

    def test_shell_and_signature_removed(self):
        text = ('ТЕЛЕГРАММА\nКому: Командиру в/ч 12345\nДата: 22.10.2026\nНомер: 15\nТекст:\nНа № 154/2 от 20.10.2026\n'
                '1. Сведения представлены.\n2. Техника готова.\nПодпись:\nНачальник штаба\nполковник И.И. Иванов')
        self.assertEqual(ai.strip_shell(text, self.ME), 'На № 154/2 от 20.10.2026\n1. Сведения представлены.\n2. Техника готова.')
        bare = '1. Провести проверку.\n2. Доложить.\n\nНачальник штаба\nполковник И.И. Иванов'
        self.assertEqual(ai.strip_shell(bare, self.ME), '1. Провести проверку.\n2. Доложить.')
        self.assertEqual(ai.strip_shell('1. Провести проверку.\n\n2. Доложить к сроку.'), '1. Провести проверку.\n\n2. Доложить к сроку.')

    def test_missing_points_go_to_the_end(self):
        out, missing = ai.ensure_points('1. Провести проверку техники.\n\n2. Доложить о готовности.', ['выставить охранение'])
        self.assertEqual(missing, ['выставить охранение'])
        self.assertTrue(out.rstrip().endswith('3. Выставить охранение.'))
        self.assertIn('2. Доложить о готовности.', out)

    def test_word_plain_and_into_template(self):
        d = Path(tempfile.mkdtemp())
        from docx import Document
        tg.save_docx('1. Провести проверку техники.\n2. Доложить.', str(d / 'a.docx'))
        lines = [p.text for p in Document(str(d / 'a.docx')).paragraphs if p.text.strip()]
        self.assertEqual(lines, ['1. Провести проверку техники.', '2. Доложить.'])   # без «ТЕЛЕГРАММА» и подписи
        tpl = Document()
        for t in ('ТЕЛЕГРАММА', 'Командиру в/ч 12345', '{Текст}', 'Начальник штаба полковник И.И. Иванов'):
            tpl.add_paragraph(t)
        tpl.save(str(d / 't.docx'))
        tg.save_docx('1. Провести проверку техники.\n2. Доложить.', str(d / 'b.docx'), str(d / 't.docx'))
        lines = [p.text for p in Document(str(d / 'b.docx')).paragraphs if p.text.strip()]
        self.assertEqual(lines, ['ТЕЛЕГРАММА', 'Командиру в/ч 12345', '1. Провести проверку техники.', '2. Доложить.',
                                 'Начальник штаба полковник И.И. Иванов'])


if __name__ == '__main__':
    unittest.main()
