"""Ежедневный доклад по образцу отдела: Word-файл с метками {{дата}}, {{пришло}}, {{сроки}}, {{просрочено}},
{{не_обработано}}, {{свод}}, {{ждём_ответа}}, {{главное}}. Искра подставляет сведения из утренней сводки,
поручений, автосвода и журнала исходящих. Оформление (шрифт, шапка, подпись) остаётся вашим."""
import copy
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt

KEYS = {
    'дата': 'дата доклада',
    'пришло': 'что пришло за сутки',
    'сроки': 'сроки сегодня и завтра',
    'просрочено': 'просроченные поручения',
    'не_обработано': 'не открытые входящие',
    'свод': 'свод таблиц за день',
    'ждём_ответа': 'ждём ответа на исходящие',
    'главное': 'главное (пишет нейросеть, если включено)',
    'подпись': 'подпись (из «Мои реквизиты»)',
}
SECTION_KEYS = {'ПРИШЛО': 'пришло', 'СРОКИ СЕГОДНЯ И ЗАВТРА': 'сроки', 'ПРОСРОЧЕНО': 'просрочено',
                'НЕ ОБРАБОТАНО': 'не_обработано', 'СВОД ЗА ДЕНЬ': 'свод'}


def values_from(sections, waiting_lines=(), main='', today=None, signature=''):
    today = today or date.today()
    vals = {k: [] for k in KEYS}
    vals['дата'] = today.strftime('%d.%m.%Y')
    for title, lines in sections:
        for prefix, key in SECTION_KEYS.items():
            if title.startswith(prefix):
                vals[key] = list(lines)
    vals['ждём_ответа'] = list(waiting_lines)
    vals['главное'] = main.strip()
    vals['подпись'] = signature.strip() or '[ДОЛЖНОСТЬ]\n[звание]                                        [И.О. Фамилия]'
    return vals


def make_template(path):
    """Образец для правки в Word: шапка, разделы с метками, подпись."""
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(14)
    doc.add_paragraph('ДОКЛАД').alignment = 1
    doc.add_paragraph('о состоянии дел на {{дата}}').alignment = 1
    doc.add_paragraph('{{главное}}')
    for key, title in (('пришло', 'Поступило за сутки:'), ('сроки', 'Сроки исполнения сегодня и завтра:'),
                       ('просрочено', 'Просрочено:'), ('не_обработано', 'Не обработано:'), ('свод', 'Свод за день:'),
                       ('ждём_ответа', 'Ожидается ответ:')):
        doc.add_paragraph(title).runs[0].bold = True
        doc.add_paragraph('{{' + key + '}}')
    doc.add_paragraph('')
    doc.add_paragraph('{{подпись}}')
    doc.save(path)


def _set_text(par, text):
    runs = par.runs
    if not runs:
        par.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ''


def _fill_paragraph(par, vals):
    text = par.text
    if '{{' not in text:
        return
    stripped = text.strip()
    for key, value in vals.items():
        mark = '{{' + key + '}}'
        if mark not in text:
            continue
        if isinstance(value, list):
            if stripped == mark:
                # Метка одна в абзаце — по абзацу на строку, в том же оформлении.
                lines = value or ['нет']
                _set_text(par, ('— ' + lines[0]) if value else 'нет')
                anchor = par
                for line in lines[1:]:
                    new = copy.deepcopy(par._p)
                    anchor._p.addnext(new)
                    from docx.text.paragraph import Paragraph
                    anchor = Paragraph(new, par._parent)
                    _set_text(anchor, '— ' + line)
                return
            value = '; '.join(value) if value else 'нет'
        if stripped == mark and not value:
            par._p.getparent().remove(par._p)          # пустое «главное» — абзац не нужен
            return
        text = text.replace(mark, str(value))
    _set_text(par, text)


def fill(template, vals, out):
    doc = Document(template)
    for par in list(doc.paragraphs):
        _fill_paragraph(par, vals)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for par in list(cell.paragraphs):
                    _fill_paragraph(par, vals)
    doc.save(out)
    return out


def build(root, sections, waiting_lines=(), main='', today=None, template=None, signature=''):
    """Доклад в data/doklady/Доклад ДД.ММ.ГГГГ.docx. Нет своего образца — создаётся стандартный."""
    today = today or date.today()
    root = Path(root)
    template = Path(template) if template else root / 'doklad_obrazec.docx'
    if not template.is_file():
        make_template(template)
    out_dir = root / 'doklady'
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f'Доклад {today.strftime("%d.%m.%Y")}.docx'
    return fill(template, values_from(sections, waiting_lines, main, today, signature), out)
