"""Windows for batch Word ↔ Excel, telegram merging, spelling and the Excel formula helper."""
import os
import threading
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from iskra_batch import discover_fields, collect_tables, mail_merge, template_fields
from iskra_crossformat import merge_crossformat
from iskra_formula import build_formula, explain_formula
from iskra_index import ROOT, extract
from iskra_services import telegram_check, telegram_template
from iskra_spell import SpellChecker, summary as spell_summary
from iskra_telegram import merge as merge_telegrams_full, save_docx as telegram_docx, TEMPLATE_HELP
from iskra_spell import fix_typical
import json
from iskra_workflows import normalize_layout, draft_from_sources, save_history
import iskra_vba

DOC_TYPES = [('Word и текст', '*.docx *.txt'), ('Все документы', '*.docx *.xlsx *.pptx *.pdf *.txt *.md *.csv')]


def _window(app, title, size):
    win = tk.Toplevel(app.root)
    win.title('Искра · ' + title)
    win.geometry(size)
    win.transient(app.root)
    return win


def _open_path(path):
    try:
        os.startfile(path)
    except (OSError, AttributeError):
        pass


class ToolsMixin:
    spell = None

    def spellchecker(self):
        if self.spell is None:
            self.spell = SpellChecker()
            if self.spell.available:
                self.spell.preload()
        return self.spell

    # ---------- Word → Excel ----------

    def word_to_excel(self):
        files = filedialog.askopenfilenames(title='Выберите документы Word (можно сразу 20 и больше)', filetypes=[('Word', '*.docx')])
        if not files:
            return
        win = _window(self, 'Word → Excel', '760x560')
        ttk.Label(win, text=f'Выбрано документов: {len(files)}. Что собрать в таблицу?', font=('Segoe UI', 11, 'bold')).pack(anchor='w', padx=12, pady=(12, 6))
        mode = tk.StringVar(value='fields')
        ttk.Radiobutton(win, text='Поля вида «Название: значение» (анкеты, карточки, справки) — одна строка на документ',
                        variable=mode, value='fields').pack(anchor='w', padx=12)
        ttk.Radiobutton(win, text='Таблицы из документов — все строки всех таблиц в одну таблицу Excel',
                        variable=mode, value='tables').pack(anchor='w', padx=12)
        ttk.Label(win, text='Найденные поля (отмечены те, что есть хотя бы в половине документов). Ctrl+щелчок — выбрать/снять:',
                  wraplength=720).pack(anchor='w', padx=12, pady=(10, 4))
        frame = ttk.Frame(win); frame.pack(fill='both', expand=True, padx=12)
        box = tk.Listbox(frame, selectmode='extended', font=('Segoe UI', 10), exportselection=False)
        box.pack(side='left', fill='both', expand=True)
        sc = ttk.Scrollbar(frame, command=box.yview); sc.pack(side='right', fill='y'); box.config(yscrollcommand=sc.set)
        ttk.Label(win, text='Своё поле (если его нет в списке) — точная подпись из документов:').pack(anchor='w', padx=12, pady=(8, 2))
        extra = tk.StringVar(); ttk.Entry(win, textvariable=extra).pack(fill='x', padx=12)
        use_ai = tk.BooleanVar(value=getattr(self, 'ai_ready', False))
        ai_box = ttk.Checkbutton(win, variable=use_ai, text='Чего нет в виде «Поле: значение» — найти нейросетью '
                                 '(берётся только значение, которое дословно есть в документе; помечается на листе «Проверка»)')
        ai_box.pack(anchor='w', padx=12, pady=(6, 0))
        if not getattr(self, 'ai_ready', False):
            ai_box.state(['disabled'])
        status = tk.StringVar(value='Ищу поля в документах…')
        ttk.Label(win, textvariable=status, wraplength=720).pack(anchor='w', padx=12, pady=6)
        labels = []

        def found(fields):
            labels.clear()
            for label, n in fields:
                labels.append(label)
                box.insert('end', f'{label}   —   в {n} из {len(files)}')
                if n * 2 >= len(files):
                    box.selection_set('end')
            status.set(f'Найдено полей: {len(fields)}.' if fields else 'Поля вида «Название: значение» не найдены — выберите «Таблицы» или впишите поле.')

        def work():
            fields = discover_fields(files)
            self.events.put(('call', lambda: found(fields) if win.winfo_exists() else None))
        threading.Thread(target=work, daemon=True).start()

        def run():
            output = filedialog.asksaveasfilename(parent=win, title='Сохранить таблицу Excel', defaultextension='.xlsx',
                                                  initialdir=str(ROOT), filetypes=[('Excel', '*.xlsx')])
            if not output:
                return
            chosen = [labels[i] for i in box.curselection()] + [x.strip() for x in extra.get().split(';') if x.strip()]
            if mode.get() != 'tables' and not chosen:
                messagebox.showwarning('Поля', 'Отметьте хотя бы одно поле.', parent=win); return
            ai_on = use_ai.get() and mode.get() != 'tables'
            filled = {'n': 0}

            def filler(path, missing):
                import iskra_ai
                filled['n'] += 1
                self.events.put(('call', lambda n=filled['n'], name=Path(path).name: status.set(
                    f'Нейросеть ищет недостающие поля: {name} ({n})… Esc — остановить')))
                if self.ai_cancel_all.is_set():
                    return {}
                return iskra_ai.extract_fields(extract(Path(path)), missing, cancel=self.ai_cancel_all)

            def work():
                try:
                    if mode.get() == 'tables':
                        r = collect_tables(files, output)
                        msg = f"Строк: {r['rows']}, листов: {r['sheets']}.\nНе прочитано файлов: {len(r['failed'])}."
                    else:
                        r = merge_crossformat(files, chosen, output, 'xlsx', filler=filler if ai_on else None)
                        msg = (f"Записей: {r['rows']}, замечаний: {r['issues']}.\nПустые и противоречивые значения — «НЕ НАЙДЕНО», "
                               "подробности на листе «Проверка»." + ("\nЗначения от нейросети — с пометкой «ЗАПОЛНЕНО НЕЙРОСЕТЬЮ»." if ai_on else ''))
                    save_history(ROOT, 'Word → Excel', f'{len(files)} файлов → {Path(output).name}')
                    self.events.put(('call', lambda: finish(msg)))
                except Exception as exc:
                    err = str(exc)
                    self.events.put(('call', lambda: messagebox.showerror('Ошибка', err, parent=win) if win.winfo_exists() else None))

            def finish(msg):
                if messagebox.askyesno('Готово', msg + f'\n\n{output}\n\nОткрыть файл?', parent=win if win.winfo_exists() else None):
                    _open_path(output)
                if win.winfo_exists():
                    win.destroy()
            self.ai_cancel_all.clear()
            status.set('Собираю…')
            threading.Thread(target=work, daemon=True).start()
        ttk.Button(win, text='Собрать в Excel', style='Accent.TButton', command=run).pack(pady=(0, 12))

    # ---------- Excel → Word ----------

    def excel_to_word(self):
        win = _window(self, 'Excel → Word', '760x470')
        ttk.Label(win, text='Excel → Word', font=('Segoe UI', 11, 'bold')).pack(anchor='w', padx=12, pady=(12, 4))
        mode = tk.StringVar(value='merge')
        ttk.Radiobutton(win, text='Документы по образцу: каждая строка Excel — отдельный документ Word (справки, письма, телеграммы)',
                        variable=mode, value='merge').pack(anchor='w', padx=12)
        ttk.Radiobutton(win, text='Одна таблица Word из нескольких файлов Excel',
                        variable=mode, value='table').pack(anchor='w', padx=12)
        ttk.Label(win, wraplength=720, style='Muted.TLabel', text=(
            'Как подготовить образец: напишите документ в Word и вместо данных впишите названия столбцов Excel '
            'в фигурных скобках, например: «Выдана {ФИО} в том, что он проходит службу в {Подразделение}». '
            'Первая строка Excel — заголовки.')).pack(anchor='w', padx=12, pady=8)
        paths = {'xlsx': tk.StringVar(), 'tpl': tk.StringVar(), 'out': tk.StringVar()}
        name_field = tk.StringVar()
        grid = ttk.Frame(win); grid.pack(fill='x', padx=12)

        def pick(key, kind):
            if kind == 'dir':
                p = filedialog.askdirectory(parent=win, title='Папка для новых документов')
            else:
                types = [('Excel', '*.xlsx *.xlsm')] if kind == 'xlsx' else [('Word', '*.docx')]
                p = filedialog.askopenfilename(parent=win, filetypes=types)
            if p:
                paths[key].set(p)
                if key == 'tpl':
                    try:
                        fields = template_fields(p)
                        names.configure(values=fields)
                        status.set('Поля образца: ' + (', '.join(fields) if fields else 'не найдены — впишите {Поле}'))
                        if fields and not name_field.get():
                            name_field.set(fields[0])
                    except Exception as exc:
                        status.set(f'Не удалось прочитать образец: {exc}')

        rows = [('Таблица Excel:', 'xlsx', 'xlsx'), ('Образец Word:', 'tpl', 'docx'), ('Куда сохранить:', 'out', 'dir')]
        for r, (label, key, kind) in enumerate(rows):
            ttk.Label(grid, text=label).grid(row=r, column=0, sticky='w', pady=3)
            ttk.Entry(grid, textvariable=paths[key], width=70).grid(row=r, column=1, sticky='we', padx=6)
            ttk.Button(grid, text='Выбрать…', command=lambda k=key, t=kind: pick(k, t)).grid(row=r, column=2)
        ttk.Label(grid, text='Имя файлов из поля:').grid(row=3, column=0, sticky='w', pady=3)
        names = ttk.Combobox(grid, textvariable=name_field, width=30); names.grid(row=3, column=1, sticky='w', padx=6)
        grid.columnconfigure(1, weight=1)
        status = tk.StringVar(value='')
        ttk.Label(win, textvariable=status, wraplength=720).pack(anchor='w', padx=12, pady=8)

        def run():
            try:
                if mode.get() == 'table':
                    files = filedialog.askopenfilenames(parent=win, title='Файлы Excel', filetypes=[('Excel', '*.xlsx *.xlsm')])
                    if not files:
                        return
                    from openpyxl import load_workbook
                    wb = load_workbook(files[0], read_only=True)
                    header = [str(c) for c in next(wb.active.iter_rows(values_only=True)) if c is not None]
                    wb.close()
                    output = filedialog.asksaveasfilename(parent=win, defaultextension='.docx', filetypes=[('Word', '*.docx')], initialdir=str(ROOT))
                    if not output:
                        return
                    r = merge_crossformat(files, header, output, 'docx')
                    messagebox.showinfo('Готово', f"Строк в таблице Word: {r['rows']}\n{output}", parent=win)
                    _open_path(output)
                    return
                if not (paths['xlsx'].get() and paths['tpl'].get() and paths['out'].get()):
                    messagebox.showwarning('Не всё выбрано', 'Выберите таблицу Excel, образец Word и папку.', parent=win); return
                r = mail_merge(paths['xlsx'].get(), paths['tpl'].get(), paths['out'].get(), name_field.get().strip() or None)
                save_history(ROOT, 'Excel → Word', f"{r['created']} документов в {r['output']}")
                if messagebox.askyesno('Готово', f"Создано документов: {r['created']}\nПапка: {r['output']}\n\nОткрыть папку?", parent=win):
                    _open_path(r['output'])
            except Exception as exc:
                messagebox.showerror('Ошибка', str(exc), parent=win)
        ttk.Button(win, text='Создать', style='Accent.TButton', command=run).pack(pady=(0, 12))

    # ---------- Телеграммы ----------

    def telegrams(self, initial=None, note=None, report=None):
        win = _window(self, 'редактор телеграмм', '960x680')
        ttk.Label(win, text='Проект телеграммы · исходные документы не изменяются', font=('Segoe UI', 12, 'bold')).pack(anchor='w', padx=12, pady=8)
        buttons = ttk.Frame(win); buttons.pack(side='bottom', fill='x', padx=12, pady=8)
        status = tk.StringVar(value=note or 'Используйте утверждённый образец организации. Шаблон здесь — пример.')
        ttk.Label(win, textvariable=status, wraplength=820).pack(side='bottom', anchor='w', padx=12, pady=6)
        if report:
            box = tk.Text(win, height=min(8, len(report) + 1), wrap='word', font=('Segoe UI', 9), background='#FFF8E1', relief='flat')
            box.insert('1.0', 'Что сделано при объединении:\n' + '\n'.join('• ' + r for r in report))
            box.configure(state='disabled')
            box.pack(side='bottom', fill='x', padx=12)
        t = tk.Text(win, wrap='word', font=('Segoe UI', 11), undo=True)
        t.pack(fill='both', expand=True, padx=12)
        t.insert('1.0', initial or telegram_template())

        def check():
            issues = telegram_check(t.get('1.0', 'end-1c'))
            issues = [i for i in issues if not i.startswith('Орфография и фактическая точность')]
            messagebox.showinfo('Проверка перед отправкой', '\n'.join('• ' + x for x in issues) or 'Замечаний по оформлению нет.', parent=win)

        def save():
            path = filedialog.asksaveasfilename(parent=win, defaultextension='.docx', filetypes=[('Word', '*.docx')], initialdir=str(ROOT))
            if path:
                template = self.telegram_template_path()
                try:
                    telegram_docx(t.get('1.0', 'end-1c'), path, template)
                except Exception as exc:
                    messagebox.showerror('Ошибка', str(exc), parent=win); return
                status.set('Сохранено: ' + path + (' (по вашему шаблону)' if template else ' (стандартное оформление: Times New Roman 14, отступ 1,25 см, поля по ГОСТ)'))
                if messagebox.askyesno('Готово', 'Телеграмма сохранена. Открыть?', parent=win):
                    _open_path(path)

        def tidy():
            t.replace('1.0', 'end', normalize_layout(t.get('1.0', 'end-1c')))
            status.set('Пробелы и пустые строки приведены к единому виду.')

        def prepare():
            ix = self.selected()
            if ix is None:
                messagebox.showinfo('Источник', 'Сначала найдите и выберите документ в главном окне.', parent=win); return
            path = self.results[ix][0]
            try:
                t.replace('1.0', 'end', draft_from_sources('По выбранному документу', [(path, extract(Path(path)))]))
                save_history(ROOT, 'Проект телеграммы', Path(path).name)
                status.set('Создан черновик с выдержками из источника.')
            except Exception as exc:
                messagebox.showerror('Ошибка', str(exc), parent=win)

        ttk.Button(buttons, text='Орфография', style='Accent.TButton', command=lambda: self.spell_window(target=t)).pack(side='left')
        ttk.Button(buttons, text='Проверить оформление', command=check).pack(side='left', padx=6)
        ttk.Button(buttons, text='Убрать лишние пробелы', command=tidy).pack(side='left')
        ttk.Button(buttons, text='Черновик по выбранному файлу', command=prepare).pack(side='left', padx=6)
        if hasattr(self, 'editor_ai'):
            self.editor_ai(buttons, t, status).pack(side='left', padx=(0, 6))
        ttk.Button(buttons, text='Сохранить в Word', style='Accent.TButton', command=save).pack(side='right')
        ttk.Button(buttons, text='Шаблон оформления…', command=lambda: self.telegram_template_window(win)).pack(side='right', padx=6)
        return win, t, status

    def telegram_settings(self):
        try:
            return json.loads((ROOT / 'telegram.json').read_text(encoding='utf-8'))
        except (OSError, ValueError):
            return {}

    def telegram_template_path(self):
        p = self.telegram_settings().get('template')
        return p if p and Path(p).is_file() else None

    def telegram_template_window(self, parent=None):
        win = _window(self, 'шаблон оформления телеграмм', '680x300')
        current = self.telegram_template_path()
        state = tk.StringVar(value=current or '')
        ttk.Label(win, text=TEMPLATE_HELP, wraplength=640).pack(anchor='w', padx=12, pady=10)
        ttk.Label(win, textvariable=state, style='Muted.TLabel', wraplength=640).pack(anchor='w', padx=12)

        def save(value):
            data = self.telegram_settings(); data['template'] = value
            (ROOT / 'telegram.json').write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
            state.set(value or 'Стандартное оформление: Times New Roman 14, отступ первой строки 1,25 см, поля по ГОСТ Р 7.0.97.')

        def choose():
            p = filedialog.askopenfilename(parent=win, title='Шаблон телеграммы (.docx)', filetypes=[('Word', '*.docx')])
            if p:
                save(p)
        if not current:
            state.set('Стандартное оформление: Times New Roman 14, отступ первой строки 1,25 см, поля по ГОСТ Р 7.0.97.')
        row = ttk.Frame(win); row.pack(pady=12)
        ttk.Button(row, text='Выбрать свой шаблон…', style='Accent.TButton', command=choose).pack(side='left')
        ttk.Button(row, text='Стандартное оформление', command=lambda: save('')).pack(side='left', padx=8)
        ttk.Button(row, text='Закрыть', command=win.destroy).pack(side='left')

    def merge_telegrams_window(self, files=None):
        files = files or filedialog.askopenfilenames(title='Выберите похожие телеграммы (2–10 штук, Ctrl+щелчок)', filetypes=DOC_TYPES)
        if not files:
            return
        if len(files) < 2:
            messagebox.showinfo('Объединение', 'Выберите хотя бы две телеграммы (Ctrl+щелчок).'); return
        win = _window(self, 'объединение телеграмм', '620x360')
        ttk.Label(win, text=f'Телеграмм: {len(files)}. Как объединить?', font=('Segoe UI', 11, 'bold')).pack(anchor='w', padx=12, pady=(12, 6))
        main_only = tk.BooleanVar(value=True); dates = tk.BooleanVar(value=True); typos = tk.BooleanVar(value=True)
        ttk.Checkbutton(win, text='Оставить главное: убрать вводные фразы («Доводим до вашего сведения», приветствия)', variable=main_only).pack(anchor='w', padx=12, pady=2)
        ttk.Checkbutton(win, text='Даты к одному виду: «25 октября 2026 года» → 25.10.2026', variable=dates).pack(anchor='w', padx=12, pady=2)
        ttk.Checkbutton(win, text='Исправить типовые ошибки («в течении», «по приезду», двойные пробелы, «и и»)', variable=typos).pack(anchor='w', padx=12, pady=2)
        me = self.telegram_settings().get('me', {})
        polish = tk.BooleanVar(value=getattr(self, 'ai_ready', False))
        pb = ttk.Checkbutton(win, text='Свести нейросетью в связный текст (числа и даты сверяются с исходными телеграммами)', variable=polish)
        pb.pack(anchor='w', padx=12, pady=2)
        if not getattr(self, 'ai_ready', False):
            pb.state(['disabled'])
        mine = tk.BooleanVar(value=bool(me.get('signature')))
        if me.get('signature'):
            ttk.Checkbutton(win, text='Подпись и «От кого» — мои (из реквизитов), а не из исходных телеграмм', variable=mine).pack(anchor='w', padx=12, pady=2)
        ttk.Label(win, wraplength=590, style='Muted.TLabel', text=(
            'Одинаковые пункты остаются один раз, даже если написаны другими словами. '
            'Если сроки или числа в похожих пунктах различаются — пункт помечается [УТОЧНИТЬ]. '
            'Основание — в начало, «О выполнении доложить» — один раз в конец. Подпись — общая или [УТОЧНИТЬ ПОДПИСЬ].')).pack(anchor='w', padx=12, pady=10)

        def run():
            try:
                m = merge_telegrams_full([(Path(p).name, extract(Path(p))) for p in files], main_only=main_only.get(),
                                         dates=dates.get(), fix=fix_typical if typos.get() else None)
            except Exception as exc:
                messagebox.showerror('Ошибка', str(exc), parent=win); return
            if mine.get():
                m.signature = [x.strip() for x in me['signature'].splitlines() if x.strip()]
                if me.get('from'):
                    m.header['От кого'] = me['from']
            save_history(ROOT, 'Объединение телеграмм', ', '.join(Path(p).name for p in files)[:400])
            win.destroy()
            ed = self.telegrams(m.text(), ' '.join(m.report[:3]) + ' Проверьте пометки [УТОЧНИТЬ], [ДАТА], [НОМЕР] и нажмите «Орфография».',
                                report=m.report)
            if polish.get():
                import iskra_ai
                sources = [(Path(p).name, extract(Path(p))) for p in files]
                draft = m.text().split('--- ИСТОЧНИКИ')[0]
                self.ai_into_editor(ed, lambda cancel, update: iskra_ai.merge(sources, draft, me=me, on_text=update, cancel=cancel)[0],
                                    sources, 'Объединение нейросетью')
        ttk.Button(win, text='Объединить', style='Accent.TButton', command=run).pack(pady=(0, 12))

    # ---------- Орфография ----------

    def spell_window(self, target=None):
        checker = self.spellchecker()
        if not checker.available:
            messagebox.showerror('Орфография', 'Не найден словарь (папка dict рядом с программой).'); return
        win = _window(self, 'проверка орфографии', '980x620')
        top = ttk.Frame(win); top.pack(fill='x', padx=12, pady=8)
        status = tk.StringVar(value='Вставьте текст или откройте документ и нажмите «Проверить».')
        ttk.Label(top, textvariable=status).pack(side='left')
        pane = ttk.Panedwindow(win, orient='horizontal'); pane.pack(fill='both', expand=True, padx=12)
        if target is None:
            t = tk.Text(pane, wrap='word', font=('Segoe UI', 11), undo=True)
        else:
            t = target  # проверяем текст прямо в редакторе телеграммы
        if target is None:
            pane.add(t, weight=3)
        right = ttk.Frame(pane); pane.add(right, weight=2)
        for tag, color in (('spelling', '#C62828'), ('style', '#1565C0'), ('typo', '#8A5300')):
            t.tag_configure(tag, underline=True, foreground=color)
        t.tag_configure('current', background='#FFE58F')
        issues_box = tk.Listbox(right, font=('Segoe UI', 10), exportselection=False)
        issues_box.pack(fill='both', expand=True)
        ttk.Label(right, text='Варианты:').pack(anchor='w', pady=(6, 0))
        sugg = tk.Listbox(right, height=6, font=('Segoe UI', 10), exportselection=False); sugg.pack(fill='x')
        actions = ttk.Frame(right); actions.pack(fill='x', pady=6)
        state = {'issues': []}

        def run_check():
            text = t.get('1.0', 'end-1c')
            status.set('Проверяю…')

            def work():
                issues = checker.check(text)
                self.events.put(('call', lambda: show(issues) if win.winfo_exists() else None))
            threading.Thread(target=work, daemon=True).start()

        def show(issues):
            state['issues'] = issues
            for tag in ('spelling', 'style', 'typo', 'current'):
                t.tag_remove(tag, '1.0', 'end')
            issues_box.delete(0, 'end'); sugg.delete(0, 'end')
            for i in issues:
                t.tag_add(i.kind, f'1.0+{i.start}c', f'1.0+{i.end}c')
                kind = {'spelling': 'Орфография', 'style': 'Стиль', 'typo': 'Опечатка'}[i.kind]
                issues_box.insert('end', f'{kind}: «{i.text}» — {i.message}')
            status.set(spell_summary(issues) + ('  Красным — орфография, синим — стиль, коричневым — опечатки.' if issues else ''))

        def current():
            ix = issues_box.curselection()
            return state['issues'][ix[0]] if ix else None

        def select(event=None):
            i = current()
            if not i:
                return
            t.tag_remove('current', '1.0', 'end')
            t.tag_add('current', f'1.0+{i.start}c', f'1.0+{i.end}c'); t.see(f'1.0+{i.start}c')
            sugg.delete(0, 'end')
            if i.fix:
                sugg.insert('end', i.fix)
            elif i.kind == 'spelling':
                for s in checker.suggest(i.text):
                    sugg.insert('end', s)
            if sugg.size():
                sugg.selection_set(0)

        def apply():
            i = current()
            ix = sugg.curselection()
            if not i or not ix:
                return
            replacement = sugg.get(ix[0])
            if i.text[:1].isupper() and replacement[:1].islower():
                replacement = replacement[:1].upper() + replacement[1:]
            t.delete(f'1.0+{i.start}c', f'1.0+{i.end}c'); t.insert(f'1.0+{i.start}c', replacement)
            run_check()

        def learn():
            i = current()
            if i and i.kind == 'spelling':
                checker.add_word(i.text); run_check()

        def open_doc():
            p = filedialog.askopenfilename(parent=win, filetypes=DOC_TYPES)
            if p:
                t.delete('1.0', 'end'); t.insert('1.0', extract(Path(p))); run_check()

        issues_box.bind('<<ListboxSelect>>', select)
        sugg.bind('<Double-Button-1>', lambda e: apply())
        ttk.Button(actions, text='Исправить', style='Accent.TButton', command=apply).pack(side='left')
        ttk.Button(actions, text='Добавить в словарь', command=learn).pack(side='left', padx=6)
        ttk.Button(top, text='Проверить', style='Accent.TButton', command=run_check).pack(side='right')

        def ai_fix():
            import iskra_ai
            text = t.get('1.0', 'end-1c')
            if not text.strip():
                return
            status.set('Нейросеть исправляет ошибки… (Ctrl+Z — вернуть как было)')
            t.edit_separator()

            def done(result):
                if win.winfo_exists():
                    body, sep, tail = text.partition('--- ИСТОЧНИКИ')
                    t.delete('1.0', 'end'); t.insert('1.0', result + ('\n\n' + sep + tail if sep else ''))
                    run_check()
            self.ai_background(lambda cancel: iskra_ai.edit(text.partition('--- ИСТОЧНИКИ')[0], 'Исправить ошибки', cancel=cancel), done)
        if getattr(self, 'ai_ready', False):
            ttk.Button(top, text='Исправить всё нейросетью', command=ai_fix).pack(side='right', padx=6)
        if target is None:
            ttk.Button(top, text='Открыть документ…', command=open_doc).pack(side='right', padx=6)
        else:
            win.title('Искра · орфография телеграммы')

            def close():
                for tag in ('spelling', 'style', 'typo', 'current'):
                    t.tag_remove(tag, '1.0', 'end')
                win.destroy()
            win.protocol('WM_DELETE_WINDOW', close)
            run_check()

    # ---------- Формулы Excel ----------

    def formulas(self, request=''):
        win = _window(self, 'формулы Excel', '800x520')
        ttk.Label(win, text='Опишите задачу обычными словами или вставьте формулу, чтобы её объяснить',
                  font=('Segoe UI', 12, 'bold')).pack(anchor='w', padx=12, pady=8)
        q = tk.StringVar(value=request)
        entry = ttk.Entry(win, textvariable=q, font=('Segoe UI', 11)); entry.pack(fill='x', padx=12)
        result = tk.Text(win, wrap='word', font=('Segoe UI', 11), height=12); result.pack(fill='both', expand=True, padx=12, pady=10)
        examples = ('Примеры:\n• сумма столбца C где в столбце A "Москва"\n• количество где в B больше 100\n'
                    '• среднее по столбцу D если в колонке B равно Готово\n• подтянуть с другого листа\n'
                    '• рабочие дни между датами • стаж • найти повторы в столбце A • просрочено\n'
                    '• =ЕСЛИОШИБКА(ВПР(E2;A:B;2;ЛОЖЬ);"нет") — объяснить формулу')
        result.insert('1.0', examples)
        formula = {'value': ''}

        def go(event=None):
            text = q.get().strip()
            result.delete('1.0', 'end')
            if text.startswith('='):
                formula['value'] = ''
                result.insert('1.0', 'Что делает формула:\n' + explain_formula(text))
                return
            f, note = build_formula(text)
            formula['value'] = f or ''
            if f:
                result.insert('1.0', f + '\n\n' + note + '\n\nИз чего состоит:\n' + explain_formula(f))
            elif getattr(self, 'ai_ready', False):
                import iskra_ai
                result.insert('1.0', 'В справочнике готовой формулы нет — спрашиваю нейросеть…')

                def done(res):
                    if not win.winfo_exists():
                        return
                    formula['value'] = res[0]
                    result.delete('1.0', 'end')
                    result.insert('1.0', res[0] + '\n\n' + (res[1] or '') + '\n(формула от нейросети: число аргументов проверено, '
                                  'диапазоны проверьте сами)\n\nИз чего состоит:\n' + explain_formula(res[0]))
                self.ai_background(lambda cancel: iskra_ai.formula(text, cancel=cancel), done)
            else:
                result.insert('1.0', note + '\n\n' + examples)

        def copy():
            if formula['value']:
                win.clipboard_clear(); win.clipboard_append(formula['value'])
                messagebox.showinfo('Скопировано', 'Формула в буфере обмена — вставьте её в ячейку Excel (Ctrl+V).', parent=win)

        entry.bind('<Return>', go)
        buttons = ttk.Frame(win); buttons.pack(pady=(0, 12))
        ttk.Button(buttons, text='Подобрать / объяснить', style='Accent.TButton', command=go).pack(side='left')
        ttk.Button(buttons, text='Копировать формулу', command=copy).pack(side='left', padx=8)
        if request:
            go()
        entry.focus_set()

    # ---------- Конструктор макросов ----------

    def macro_builder(self, task=''):
        win = _window(self, 'конструктор макросов', '1100x720')
        top = ttk.Frame(win); top.pack(fill='x', padx=12, pady=(10, 4))
        ttk.Label(top, text='Что должен сделать макрос? Пишите по шагам через запятую:', font=('Segoe UI', 11, 'bold')).pack(anchor='w')
        task_box = tk.Text(top, height=3, wrap='word', font=('Segoe UI', 11)); task_box.pack(fill='x', pady=4)
        task_box.insert('1.0', task or 'удалить пустые строки, выделить повторы в столбце B красным, отсортировать по столбцу A')
        row = ttk.Frame(top); row.pack(fill='x')
        app = tk.StringVar(value='auto')
        for text, value in (('Определить само', 'auto'), ('Excel', 'excel'), ('Word', 'word')):
            ttk.Radiobutton(row, text=text, variable=app, value=value).pack(side='left', padx=(0, 10))
        ttk.Button(row, text='Написать макрос', style='Accent.TButton', command=lambda: build()).pack(side='left', padx=(10, 6))
        ai_btn = ttk.Button(row, text='Написать нейросетью', command=lambda: build_ai())
        ai_btn.pack(side='left', padx=(0, 6))
        if not getattr(self, 'ai_ready', False):
            ai_btn.state(['disabled'])
        ttk.Button(row, text='Что я умею / примеры', command=lambda: show_help()).pack(side='left')
        ttk.Button(row, text='Готовые макросы', command=self.macros).pack(side='right')

        pane = ttk.Panedwindow(win, orient='horizontal'); pane.pack(fill='both', expand=True, padx=12, pady=6)
        left = ttk.Frame(pane); pane.add(left, weight=1)
        ttk.Label(left, text='Как я поняла задачу:').pack(anchor='w')
        steps_box = tk.Text(left, wrap='word', font=('Segoe UI', 10), width=40, state='disabled', background='#FAFAFA', relief='flat')
        steps_box.tag_configure('ok', foreground='#1B5E20'); steps_box.tag_configure('warn', foreground='#8A5300')
        steps_box.tag_configure('bad', foreground='#B3261E'); steps_box.tag_configure('head', font=('Segoe UI', 10, 'bold'))
        steps_box.pack(fill='both', expand=True)
        right = ttk.Frame(pane); pane.add(right, weight=2)
        ttk.Label(right, text='Код макроса (можно править вручную):').pack(anchor='w')
        code = tk.Text(right, wrap='none', font=('Consolas', 10), undo=True)
        sx = ttk.Scrollbar(right, orient='horizontal', command=code.xview); code.configure(xscrollcommand=sx.set)
        code.tag_configure('problem', background='#FFCDD2')
        code.pack(fill='both', expand=True); sx.pack(fill='x')

        status = tk.StringVar(value='Опишите задачу и нажмите «Написать макрос».')
        ttk.Label(win, textvariable=status, wraplength=1060).pack(anchor='w', padx=12)
        buttons = ttk.Frame(win); buttons.pack(fill='x', padx=12, pady=8)

        def steps_text(parts):
            steps_box.configure(state='normal'); steps_box.delete('1.0', 'end')
            for text, tag in parts:
                steps_box.insert('end', text + '\n', tag)
            steps_box.configure(state='disabled')

        def build():
            text = task_box.get('1.0', 'end').strip()
            try:
                vba, steps, unknown = iskra_vba.generate(text, None if app.get() == 'auto' else app.get())
            except ValueError as exc:
                steps_text([(str(exc), 'bad')]); status.set(str(exc)); return
            code.delete('1.0', 'end'); code.insert('1.0', vba)
            kind = 'Excel' if steps[0].op.app == 'excel' else 'Word'
            parts = [(f'Макрос для {kind}:', 'head')]
            for i, s in enumerate(steps, 1):
                parts.append((f'{i}. {s.title}', 'ok'))
                if s.warning:
                    parts.append(('   ⚠ ' + s.warning, 'warn'))
            if unknown:
                parts.append(('', 'ok'))
                parts.append(('Не поняла (в макрос не вошло):', 'head'))
                parts += [('• ' + u, 'bad') for u in unknown]
                parts.append(('Перефразируйте словами из списка «Что я умею»' +
                              (' или нажмите «Написать нейросетью» — она напишет макрос целиком.' if getattr(self, 'ai_ready', False) else '.'), 'bad'))
            steps_text(parts)
            problems = iskra_vba.lint(vba)
            status.set(f'Готово: шагов {len(steps)}.' + (f' Не понято частей: {len(unknown)}.' if unknown else '') +
                       (' Проверка кода: замечаний нет.' if not problems else f' Замечаний в коде: {len(problems)}.') +
                       f' Запускать в {kind}: Alt+F11 → Insert → Module → вставить → F5.')
            save_history(ROOT, 'Макрос', text[:300])

        def build_ai():
            import iskra_ai
            text = task_box.get('1.0', 'end').strip()
            target = app.get() if app.get() != 'auto' else iskra_vba.detect_app(text)
            steps_text([('Нейросеть пишет макрос для ' + ('Excel' if target == 'excel' else 'Word') + '…', 'head'),
                        ('Код проверяется так же, как макросы из шаблонов; при ошибках нейросеть исправляет его ещё раз.', 'ok')])

            def done(res):
                vba, problems = res
                if not win.winfo_exists():
                    return
                code.delete('1.0', 'end'); code.insert('1.0', vba)
                mark(problems)
                steps_text([('Макрос написан нейросетью.', 'head'),
                            ('Проверка кода: замечаний нет.' if not problems else f'Осталось замечаний: {len(problems)} — строки подсвечены.',
                             'ok' if not problems else 'warn'),
                            ('Перед запуском сохраните копию файла и прочитайте код: нейросеть может ошибиться.', 'warn')])
                status.set('Макрос от нейросети. Запуск: Alt+F11 → Insert → Module → вставить → F5. Ошибка — «Была ошибка при запуске».')
                save_history(ROOT, 'Макрос (нейросеть)', text[:300])
            self.ai_background(lambda cancel: iskra_ai.vba(text, target, cancel=cancel), done)

        def show_help():
            excel, word = iskra_vba.capabilities()
            parts = [('Excel умеет:', 'head')] + [('• ' + x, 'ok') for x in excel]
            parts += [('', 'ok'), ('Word умеет:', 'head')] + [('• ' + x, 'ok') for x in word]
            parts += [('Для Word можно добавить «во всех документах папки» — обработаются все файлы, копии сохранятся отдельно.', 'warn')]
            parts += [('', 'ok'), ('Примеры (скопируйте в поле задачи):', 'head')] + [('• ' + e.split(': ', 1)[1], 'ok') for e in iskra_vba.EXAMPLES]
            steps_text(parts)

        def mark(problems):
            code.tag_remove('problem', '1.0', 'end')
            for line, _ in problems:
                if line:
                    code.tag_add('problem', f'{line}.0', f'{line}.end')

        def check():
            problems = iskra_vba.lint(code.get('1.0', 'end-1c'))
            mark(problems)
            if not problems:
                status.set('Проверка кода: типичных ошибок нет. (Полную проверку делает только сам Excel/Word при запуске.)')
            else:
                status.set(f'Найдено замечаний: {len(problems)} — строки подсвечены.')
                messagebox.showinfo('Проверка кода', '\n'.join((f'Строка {n}: ' if n else '') + m for n, m in problems[:20]), parent=win)

        def fix():
            fixed, changes = iskra_vba.auto_fix(code.get('1.0', 'end-1c'))
            code.delete('1.0', 'end'); code.insert('1.0', fixed)
            left_over = iskra_vba.lint(fixed); mark(left_over)
            msg = ('Исправлено:\n' + '\n'.join('• ' + c for c in changes)) if changes else 'Автоматически исправлять нечего.'
            if left_over:
                msg += '\n\nОсталось исправить вручную:\n' + '\n'.join((f'Строка {n}: ' if n else '') + m for n, m in left_over[:10])
            messagebox.showinfo('Исправление', msg, parent=win)

        def error_help():
            dlg = _window(self, 'ошибка при запуске макроса', '700x460')
            ttk.Label(dlg, text='Вставьте текст ошибки из окна Excel/Word (Ctrl+C в окне ошибки копирует текст):').pack(anchor='w', padx=12, pady=8)
            err = tk.Text(dlg, height=4, wrap='word', font=('Segoe UI', 10)); err.pack(fill='x', padx=12)
            out = tk.Text(dlg, wrap='word', font=('Segoe UI', 10), state='disabled', background='#FAFAFA', relief='flat')
            out.pack(fill='both', expand=True, padx=12, pady=8)

            def explain():
                title, cause, what = iskra_vba.explain_error(err.get('1.0', 'end'))
                out.configure(state='normal'); out.delete('1.0', 'end')
                out.insert('end', title + '\n\nПочему: ' + cause + '\n\nЧто сделать: ' + what)
                fixed, changes = iskra_vba.auto_fix(code.get('1.0', 'end-1c'))
                if changes:
                    out.insert('end', '\n\nВ коде можно исправить автоматически:\n' + '\n'.join('• ' + c for c in changes) +
                               '\nНажмите «Исправить код» и вставьте макрос заново.')
                out.configure(state='disabled')
                if getattr(self, 'ai_ready', False) and err.get('1.0', 'end').strip():
                    import iskra_ai
                    out.configure(state='normal'); out.insert('end', '\n\nНейросеть разбирает ошибку в вашем коде…'); out.configure(state='disabled')

                    def done(text):
                        if dlg.winfo_exists():
                            out.configure(state='normal')
                            out.delete('end-2l linestart', 'end')
                            out.insert('end', '\n\nРАЗБОР НЕЙРОСЕТИ:\n' + text); out.configure(state='disabled')
                    self.ai_background(lambda cancel: iskra_ai.vba_error(code.get('1.0', 'end-1c'), err.get('1.0', 'end'), cancel=cancel), done)
            row2 = ttk.Frame(dlg); row2.pack(pady=(0, 10))
            ttk.Button(row2, text='Объяснить', style='Accent.TButton', command=explain).pack(side='left')
            ttk.Button(row2, text='Исправить код', command=lambda: (fix(), dlg.destroy())).pack(side='left', padx=8)

        def copy():
            win.clipboard_clear(); win.clipboard_append(code.get('1.0', 'end-1c'))
            status.set('Код скопирован. В Excel/Word: Alt+F11 → Insert → Module → Ctrl+V → F5.')

        def save():
            path = filedialog.asksaveasfilename(parent=win, defaultextension='.bas', initialfile='Iskra_Macro.bas',
                                                initialdir=str(ROOT), filetypes=[('Модуль VBA', '*.bas')])
            if path:
                Path(path).write_bytes(iskra_vba.bas_bytes(code.get('1.0', 'end-1c')))
                status.set('Сохранено: ' + path + '. Подключение: Alt+F11 → File → Import File.')

        def how():
            messagebox.showinfo('Как запустить макрос', (
                '1. Сохраните копию файла — действия макроса нельзя отменить.\n'
                '2. В Excel или Word нажмите Alt+F11.\n'
                '3. Insert → Module (Вставка → Модуль), вставьте код (Ctrl+V).\n'
                '4. Встаньте внутрь Sub Iskra_Macro и нажмите F5.\n'
                '5. Если появилась ошибка — скопируйте её текст и нажмите здесь «Была ошибка».\n\n'
                'Чтобы макрос сохранился в файле, сохраните его как .xlsm / .docm.\n'
                'Макросы могут быть запрещены политиками организации.'), parent=win)

        ttk.Button(buttons, text='Копировать код', style='Accent.TButton', command=copy).pack(side='left')
        ttk.Button(buttons, text='Сохранить .bas', command=save).pack(side='left', padx=6)
        ttk.Button(buttons, text='Проверить код', command=check).pack(side='left')
        ttk.Button(buttons, text='Исправить автоматически', command=fix).pack(side='left', padx=6)
        ttk.Button(buttons, text='Была ошибка при запуске…', command=error_help).pack(side='left')
        ttk.Button(buttons, text='Как запустить', command=how).pack(side='right')
        if task:
            build()
