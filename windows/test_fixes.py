"""Регрессионные тесты на ошибки, найденные в версии 0.9."""
import json
import os
import tempfile
import unittest
from pathlib import Path

os.environ['ISKRA_DATA_DIR'] = tempfile.mkdtemp()

from openpyxl import Workbook, load_workbook
import iskra_index
from iskra_consolidate import consolidate
from iskra_crossformat import merge_crossformat
from iskra_services import AutoIndexer, save_settings


def write(path, head, rows):
    w = Workbook(); s = w.active; s.append(head)
    for r in rows:
        s.append(r)
    w.save(path)


class Fixes(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())

    def test_empty_template_column_does_not_shift_data(self):
        write(self.dir / 'tpl.xlsx', ['Организация', None, 'Количество'], [])
        write(self.dir / 'src.xlsx', ['Количество', 'Организация'], [[5, 'Альфа']])
        consolidate(self.dir / 'tpl.xlsx', [self.dir / 'src.xlsx'], self.dir / 'out.xlsx', 'Количество')
        row = [c.value for c in load_workbook(self.dir / 'out.xlsx')['Сводная'][2]]
        self.assertEqual(row[:4], ['Альфа', None, 5, 'src.xlsx'])

    def test_text_is_copied_exactly_and_formulas_stay_text(self):
        write(self.dir / 'tpl.xlsx', ['Пункт'], [])
        write(self.dir / 'src.xlsx', ['Пункт'], [['- первый'], ['@почта']])
        # Formula text in a source cell (as text) must not become a live formula in the result.
        w = load_workbook(self.dir / 'src.xlsx'); w.active['A4'] = '=1+1'; w.active['A4'].data_type = 's'; w.save(self.dir / 'src.xlsx')
        consolidate(self.dir / 'tpl.xlsx', [self.dir / 'src.xlsx'], self.dir / 'out.xlsx')
        sheet = load_workbook(self.dir / 'out.xlsx')['Сводная']
        self.assertEqual([sheet.cell(r, 1).value for r in (2, 3, 4)], ['- первый', '@почта', '=1+1'])
        self.assertEqual(sheet.cell(4, 1).data_type, 's')

    def test_crossformat_keeps_text_exactly(self):
        write(self.dir / 'c.xlsx', ['Номер'], [['-17']])
        merge_crossformat([self.dir / 'c.xlsx'], ['Номер'], self.dir / 'out.xlsx')
        self.assertEqual(load_workbook(self.dir / 'out.xlsx')['Сводная']['B2'].value, '-17')

    def test_utf16_text_files(self):
        (self.dir / 'u16.txt').write_text('Приказ номер 5', encoding='utf-16')
        (self.dir / 'cp.txt').write_bytes('Приказ номер 6'.encode('cp1251'))
        self.assertEqual(iskra_index.extract(self.dir / 'u16.txt'), 'Приказ номер 5')
        self.assertEqual(iskra_index.extract(self.dir / 'cp.txt'), 'Приказ номер 6')

    def test_search_handles_yo_endings_and_file_names(self):
        docs = self.dir / 'docs'; docs.mkdir()
        (docs / 'Ёлочный_отчёт.txt').write_text('Отчёт о ёлке. Заключены договоры поставки.', encoding='utf-8')
        iskra_index.index_folder(str(docs.resolve()), lambda m: None)
        for q in ['отчет', 'отчёт', 'договоров', 'Ёлочный', 'елке', 'поставка']:
            self.assertEqual(len(iskra_index.search_docs(q)), 1, q)
        self.assertEqual(iskra_index.forget_folder(str(docs)), 1)
        self.assertEqual(iskra_index.search_docs('отчет'), [])

    def test_offline_folder_is_not_forgotten(self):
        root = self.dir / 'settings'
        save_settings(root, {'folders': [str(root / 'сетевой_диск')], 'interval_minutes': 15})
        ix = AutoIndexer(root, lambda *a: None, lambda *a: None)
        ix.add_folder(str(self.dir))
        folders = json.loads((root / 'settings.json').read_text(encoding='utf-8'))['folders']
        self.assertEqual(len(folders), 2)


if __name__ == '__main__':
    unittest.main()
