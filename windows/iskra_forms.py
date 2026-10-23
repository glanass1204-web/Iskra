"""Бланки военных документов по Инструкции по делопроизводству в ВС РФ (приказ МО РФ от 04.04.2017 № 170):
отпускной билет (приложение 30). Только оформление — сведения вписывает человек, ничего не выдумывается."""
import re
from datetime import date, timedelta

ONES = ['ноль', 'одни', 'двое', 'трое', 'четверо', 'пять', 'шесть', 'семь', 'восемь', 'девять', 'десять',
        'одиннадцать', 'двенадцать', 'тринадцать', 'четырнадцать', 'пятнадцать', 'шестнадцать', 'семнадцать',
        'восемнадцать', 'девятнадцать']
TENS = ['', '', 'двадцать', 'тридцать', 'сорок', 'пятьдесят', 'шестьдесят', 'семьдесят', 'восемьдесят', 'девяносто']
HUNDREDS = ['', 'сто', 'двести', 'триста']
LAST_COMPOUND = {1: 'одни', 2: 'двое', 3: 'трое', 4: 'четверо'}   # «сутки» — только собирательные числительные


def days_words(n):
    """15 → «пятнадцать суток», 21 → «двадцать одни сутки», 32 → «тридцать двое суток» (сроком на …)."""
    n = int(n)
    if n <= 0 or n > 399:
        return f'{n} суток'
    parts = []
    if n >= 100:
        parts.append(HUNDREDS[n // 100])
    rest = n % 100
    if 0 < rest < 20:
        word = ONES[rest]
        parts.append(word)
    elif rest:
        parts.append(TENS[rest // 10])
        if rest % 10:
            parts.append(LAST_COMPOUND.get(rest % 10, ONES[rest % 10]))
    unit = 'сутки' if rest % 10 == 1 and rest != 11 else 'суток'
    return ' '.join(p for p in parts if p) + ' ' + unit


def parse_date(text):
    m = re.fullmatch(r'\s*(\d{1,2})\.(\d{1,2})\.(\d{4}|\d{2})\s*', text or '')
    if not m:
        return None
    d, mo, y = map(int, m.groups())
    y = y + 2000 if y < 100 else y
    try:
        return date(y, mo, d)
    except ValueError:
        return None


def leave_dates(start, days):
    """Отпуск с start на days суток: (последний день, день явки). Первый день отпуска входит в срок."""
    end = start + timedelta(days=int(days) - 1)
    return end, end + timedelta(days=1)


MONTHS_GEN = ('января', 'февраля', 'марта', 'апреля', 'мая', 'июня', 'июля', 'августа', 'сентября', 'октября',
              'ноября', 'декабря')


def quoted(d):
    """date → «"01" ноября 2026 г.» как в бланке."""
    return f'"{d.day:02d}" {MONTHS_GEN[d.month - 1]} {d.year} г.' if d else '"___"____________ 20__ г.'


def leave_ticket_text(f):
    """Текст отпускного билета по полям f: who, kind_place, days, start, end, back_name, unit, back_date, vpd,
    commander. Пустое поле — пометка [УТОЧНИТЬ: …]."""
    def v(key, what):
        return (f.get(key) or '').strip() or f'[УТОЧНИТЬ: {what}]'
    days = f.get('days')
    lines = ['ОТПУСКНОЙ БИЛЕТ', '', v('who', 'воинское звание, фамилия, имя, отчество'),
             'убыл ' + v('kind_place', 'вид отпуска и населённый пункт, куда убыл'),
             'сроком на ' + (days_words(days) if str(days or '').strip().isdigit() else '[УТОЧНИТЬ: количество суток]'),
             f'с {quoted(f.get("start"))} по {quoted(f.get("end"))}',
             f'По окончании срока отпуска {v("back_name", "фамилия и инициалы")} обязан(а) явиться к месту военной службы '
             f'в {v("unit", "наименование воинской части или населённого пункта")} {quoted(f.get("back_date"))}',
             'Для проезда выданы воинские перевозочные документы — требования за № ' + ((f.get('vpd') or '').strip() or '______'),
             'Действителен по предъявлении документа, удостоверяющего личность.', '',
             'Командир (начальник) ' + ((f.get('commander') or '').strip() or '__________')]
    return '\n'.join(lines)


def save_leave_ticket(f, path):
    """Отпускной билет в Word: формат А5, угловой штамп воинской части, поля как в приложении 30."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(14.8), Cm(21.0)                       # А5
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(2), Cm(1), Cm(1.5), Cm(1.5)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(2)

    def para(text, align=None, bold=False, size=None):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if align is not None:
            p.alignment = align
        return p
    stamp = (f.get('stamp') or '').strip()
    para(stamp or 'Угловой штамп\nвоинской части', size=10)
    para('')
    para('ОТПУСКНОЙ БИЛЕТ', WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    for line in leave_ticket_text(f).splitlines()[2:]:
        para(line, WD_ALIGN_PARAGRAPH.JUSTIFY if len(line) > 60 else None)
    para('М.П.', size=10)
    doc.save(path)
    return path
