"""Offline Excel workbook consolidation, preserving provenance and checking results."""
from pathlib import Path
from collections import Counter
from datetime import datetime
import re
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill


def append_as_values(ws, row):
    """Write text that starts with '=' as text, not as a formula (protects against formula injection).
    Other leading characters (+ - @) are harmless in .xlsx, so the text is kept exactly as in the source."""
    ws.append(row)
    for cell in ws[ws.max_row]:
        if cell.data_type == 'f':
            cell.data_type = 's'


def key(value):
    return re.sub(r'\s+', ' ', str(value or '').strip()).casefold()


def consolidate(template_path, source_paths, output_path, key_column=None):
    """Merge rows into a new workbook following template column headings.

    Never alters sources/template. Every accepted row has provenance; duplicate
    identifiers are flagged rather than overwritten. Rows without identifiers
    are allowed but warned about. Unmapped columns are reported.
    """
    template_path = Path(template_path).resolve()
    output_path = Path(output_path).resolve()
    sources = [Path(p).resolve() for p in source_paths]
    if not sources: raise ValueError('Выберите хотя бы один файл-источник')
    if len(set(sources)) != len(sources): raise ValueError('В списке есть повторяющиеся файлы')
    if output_path in sources or output_path == template_path:
        raise ValueError('Нельзя сохранять результат поверх исходного файла или шаблона')
    if not template_path.is_file(): raise FileNotFoundError(template_path)
    wb = load_workbook(template_path, read_only=True, data_only=False)
    try:
        sheet = wb.active
        headers = [str(c.value).strip() if c.value is not None else '' for c in sheet[1]]
    finally: wb.close()
    if not any(headers): raise ValueError('В первой строке шаблона должны быть заголовки столбцов')
    nonempty = [key(h) for h in headers if h]
    if len(nonempty) != len(set(nonempty)):
        raise ValueError('В шаблоне есть повторяющиеся названия столбцов')
    if key_column and key(key_column) not in nonempty:
        raise ValueError('Ключевой столбец не найден в шаблоне')
    result = Workbook(); target = result.active; target.title = 'Сводная'
    target.append(headers + ['Источник: файл', 'Источник: лист', 'Источник: строка'])
    report = result.create_sheet('Проверка')
    report.append(['Тип', 'Файл', 'Лист', 'Строка', 'Описание'])
    template_map = {key(h): i for i,h in enumerate(headers) if h}
    seen = {}; total = 0; issues = 0; skipped = 0
    for source in sources:
        if source.suffix.lower() not in ('.xlsx','.xlsm'):
            report.append(['Пропущен',source.name,'','', 'Поддерживаются .xlsx и .xlsm']); issues += 1; skipped += 1; continue
        if not source.exists():
            report.append(['Ошибка',str(source),'','', 'Файл не найден']); issues += 1; skipped += 1; continue
        try:
            book = load_workbook(source,read_only=True,data_only=True)
            try:
                for sh in book:
                    rows = sh.iter_rows(values_only=True)
                    first = next(rows,None)
                    if first is None: continue
                    source_map = {key(h): i for i,h in enumerate(first) if key(h)}
                    matched = [h for h in template_map if h in source_map]
                    if not matched:
                        report.append(['Нет совпадений',source.name,sh.title,1,'Названия столбцов не совпадают с шаблоном']); issues += 1; continue
                    extra = [str(h) for h in first if key(h) and key(h) not in template_map]
                    if extra:
                        report.append(['Лишние столбцы',source.name,sh.title,1,', '.join(extra)[:1000]]); issues += 1
                    missing = [headers[i] for h,i in template_map.items() if h not in source_map]
                    if missing:
                        report.append(['Нет столбцов',source.name,sh.title,1,', '.join(missing)[:1000]]); issues += 1
                    for rownum, values in enumerate(rows,start=2):
                        if not any(v is not None and str(v).strip() for v in values): continue
                        # One value per template column, including unnamed ones, so columns never shift.
                        data = [values[source_map[key(h)]] if h and key(h) in source_map and source_map[key(h)] < len(values) else None
                                for h in headers]
                        if key_column:
                            id_value = data[template_map[key(key_column)]]
                            identifier = key(id_value)
                            if not identifier:
                                report.append(['Пустой ключ',source.name,sh.title,rownum,'Нет идентификатора']); issues += 1
                            elif identifier in seen:
                                report.append(['Дубликат ключа',source.name,sh.title,rownum,'Ранее: '+seen[identifier]]); issues += 1
                            else: seen[identifier] = f'{source.name} / {sh.title} / {rownum}'
                        if all(v is None or str(v).strip()=='' for v in data):
                            report.append(['Пустая строка',source.name,sh.title,rownum,'Нет данных в столбцах шаблона']); issues += 1; continue
                        append_as_values(target, data+[source.name,sh.title,rownum]); total += 1
            finally: book.close()
        except Exception as exc:
            report.append(['Ошибка чтения',source.name,'','',str(exc)[:500]]); issues += 1; skipped += 1
    target.freeze_panes='A2';target.auto_filter.ref=target.dimensions
    for ws in (target,report):
        for cell in ws[1]:
            cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24518A')
    # Independent validation: compare accepted row count with output worksheet.
    assert target.max_row-1 == total, 'Несовпадение количества записанных строк'
    report.append(['ИТОГ','','','',f'Строк: {total}; замечаний: {issues}; пропущено файлов: {skipped}'])
    output_path.parent.mkdir(parents=True,exist_ok=True)
    result.save(output_path)
    # Re-open generated file to verify it is readable and counts match.
    check=load_workbook(output_path,read_only=True,data_only=True)
    try:
        if check['Сводная'].max_row-1 != total: raise RuntimeError('Проверка результата не пройдена')
    finally: check.close()
    return {'rows':total,'issues':issues,'skipped_files':skipped,'output':str(output_path)}
