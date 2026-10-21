import tempfile
import unittest
from pathlib import Path

from docx import Document

from iskra_spell import fix_typical
from iskra_telegram import merge, merge_telegrams, numeric_dates, save_docx, similar, split_editor_text

A = '''ТЕЛЕГРАММА
Кому: Командиру в/ч 11111
От кого: Начальник штаба
Дата: 20.10.2026
Номер: 15
Доводим до вашего сведения, что во исполнение приказа командующего № 101 проводится проверка готовности.
1. Представить списки личного состава до 25 октября 2026 года.
2. Обеспечить готовность техники к 01.11.2026.
О выполнении доложить.
Начальник штаба
полковник И.И. Иванов'''
B = '''ТЕЛЕГРАММА
Кому: Командиру в/ч 22222
От кого: Начальник штаба
Во исполнение приказа командующего № 101 проводится проверка готовности.
Списки личного состава представить до 25.10.2026.
Обеспечить  готовность техники и вооружения к 01.11.2026.
В течении недели провести проверку караулов.
О выполнении доложить.
Начальник штаба
полковник И.И. Иванов'''
C = '''ТЕЛЕГРАММА
Кому: Командиру в/ч 11111
Здравия желаю!
Представить списки личного состава до 27.10.2026.
Провести проверку караулов в течение недели.
О выполнении доложить.
Подпись: Начальник штаба
полковник И.И. Иванов'''


class Merge(unittest.TestCase):
    def setUp(self):
        self.m = merge([('a.docx', A), ('b.docx', B), ('c.docx', C)], fix=fix_typical)
        self.text = self.m.text()

    def test_points_without_repeats_even_when_reworded(self):
        self.assertEqual([p for p, _ in self.m.points], [
            'Представить списки личного состава до 25.10.2026. [УТОЧНИТЬ: в «c.docx»: 27.10.2026]',
            'Обеспечить готовность техники и вооружения к 01.11.2026.',
            'В течение недели провести проверку караулов.',
        ])

    def test_basis_first_closing_once_intro_dropped(self):
        self.assertEqual(self.m.basis, ['Во исполнение приказа командующего № 101 проводится проверка готовности.'])
        self.assertEqual(self.m.closing, ['О выполнении доложить.'])
        self.assertNotIn('Доводим', self.text)
        self.assertNotIn('Здравия', self.text)

    def test_header_and_signature(self):
        self.assertIn('Кому: Командиру в/ч 11111; Командиру в/ч 22222', self.text)
        self.assertIn('Дата: [ДАТА]', self.text)
        self.assertEqual(self.m.signature, ['Начальник штаба', 'полковник И.И. Иванов'])

    def test_report(self):
        report = '\n'.join(self.m.report)
        self.assertIn('пунктов: 3', report)
        self.assertIn('27.10.2026', report)
        self.assertIn('Исправлено типовых ошибок', report)

    def test_different_signatures_are_marked(self):
        text, report = merge_telegrams([('a', A), ('d', C.replace('И.И. Иванов', 'П.П. Петров'))])
        self.assertIn('[УТОЧНИТЬ ПОДПИСЬ:', text)

    def test_helpers(self):
        self.assertEqual(numeric_dates('до 5 ноября 2026 г. и до 3 марта'), 'до 05.11.2026 и до 3 марта')
        self.assertTrue(similar('Обеспечить готовность техники к 01.11.2026.', 'Обеспечить готовность техники и вооружения к 01.11.2026.'))
        self.assertFalse(similar('Обеспечить готовность техники.', 'Провести проверку караулов.'))
        self.assertEqual(fix_typical('Согласно с приказом, по приезду доложить.'), ('Согласно с приказом, по приезде доложить.', 1))


class Word(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())
        self.text = merge([('a.docx', A), ('b.docx', B)]).text()

    def test_editor_text_round_trip(self):
        title, header, body, signature = split_editor_text(self.text)
        self.assertEqual(title, 'ТЕЛЕГРАММА')
        self.assertEqual(header['От кого'], 'Начальник штаба')
        self.assertTrue(body[1].startswith('1. Представить'))
        self.assertEqual(signature, ['Начальник штаба', 'полковник И.И. Иванов'])
        self.assertNotIn('ИСТОЧНИКИ', ' '.join(body + signature))

    def test_standard_layout(self):
        path = save_docx(self.text, self.dir / 'out.docx')
        d = Document(path)
        texts = [p.text for p in d.paragraphs]
        self.assertEqual(texts[0], 'ТЕЛЕГРАММА')
        self.assertIn('полковник\tИ.И. Иванов', texts)
        body = [p for p in d.paragraphs if p.text.startswith('1. ')][0]
        self.assertAlmostEqual(body.paragraph_format.first_line_indent.cm, 1.25, places=2)
        self.assertEqual(d.styles['Normal'].font.name, 'Times New Roman')

    def test_user_template_keeps_its_formatting(self):
        from docx.shared import Cm, Pt
        t = Document()
        t.add_paragraph('Кому: {Кому}')
        p = t.add_paragraph('{Текст}')
        p.paragraph_format.first_line_indent = Cm(1.5)
        p.runs[0].font.size = Pt(13)
        t.add_paragraph('{Подпись}')
        t.save(self.dir / 'tpl.docx')
        d = Document(save_docx(self.text, self.dir / 'out.docx', template=self.dir / 'tpl.docx'))
        texts = [p.text for p in d.paragraphs]
        self.assertEqual(texts[0], 'Кому: Командиру в/ч 11111; Командиру в/ч 22222')
        point = [p for p in d.paragraphs if p.text.startswith('2. ')][0]
        self.assertAlmostEqual(point.paragraph_format.first_line_indent.cm, 1.5, places=2)
        self.assertEqual(point.runs[0].font.size.pt, 13)
        self.assertEqual(texts[-2:], ['Начальник штаба', 'полковник И.И. Иванов'])


if __name__ == '__main__':
    unittest.main()
