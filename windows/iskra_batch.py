"""Batch Word ↔ Excel: find fields in many Word files, collect Word tables into one Excel workbook,
and create one Word document per Excel row from a template. Sources are never modified."""
import re
from collections import Counter, OrderedDict
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from iskra_crossformat import norm, append_as_values

PLACEHOLDER = re.compile(r'\{([^{}\n]{1,60})\}')
HEADER_FILL = PatternFill('solid', fgColor='24518A')
# «В целях учёта: …», «Согласно приказу: …» — начало предложения, а не название поля.
NOT_A_LABEL = re.compile(r'^(в|во|на|по|для|с|со|о|об|при|согласно|и|а|но|прошу|приказываю|сообщаю)\s', re.IGNORECASE)


def _style_header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = HEADER_FILL
    ws.freeze_panes = 'A2'
    for col in ws.columns:
        width = max(len(str(c.value or '')) for c in col[:200])
        ws.column_dimensions[col[0].column_letter].width = min(60, max(10, width + 2))


# ---------- Word → Excel: поля ----------

def _labels_in(doc):
    """Labels of explicitly marked fields: «Название: значение» lines and two-column tables."""
    labels = []
    for p in doc.paragraphs:
        line = p.text.strip()
        if ':' in line:
            label, value = line.split(':', 1)
            label = label.strip()
            if (value.strip() and 1 <= len(label) <= 50 and len(label.split()) <= 6 and not label.endswith(('.', '!', '?'))
                    and not NOT_A_LABEL.match(label)):
                labels.append(label)
    for t in doc.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) == 2 and cells[0] and cells[1] and len(cells[0]) <= 50:
                labels.append(cells[0].rstrip(':'))
    return labels


def discover_fields(paths):
    """[(label, in how many documents)] — the most common first. Lets the user tick fields instead of typing them."""
    counts = Counter()
    shown = {}
    for path in paths:
        try:
            doc = Document(path)
        except Exception:
            continue
        seen = set()
        for label in _labels_in(doc):
            key = norm(label)
            shown.setdefault(key, label)
            if key not in seen:
                seen.add(key)
                counts[key] += 1
    return [(shown[k], n) for k, n in counts.most_common()]


# ---------- Word → Excel: таблицы ----------

def collect_tables(paths, output):
    """All tables from many Word files into one workbook. Tables with the same header row go to the same sheet;
    each row keeps its source (file, table number)."""
    output = Path(output).resolve()
    paths = [Path(p).resolve() for p in paths]
    if output in paths:
        raise ValueError('Нельзя перезаписывать исходный документ')
    groups = OrderedDict()
    failed = []
    for path in paths:
        try:
            doc = Document(path)
        except Exception as exc:
            failed.append((path.name, str(exc)[:200]))
            continue
        for ti, table in enumerate(doc.tables, 1):
            rows = [[c.text.strip() for c in r.cells] for r in table.rows]
            rows = [r for r in rows if any(r)]
            if len(rows) < 2:
                continue
            header = rows[0]
            key = tuple(norm(h) for h in header)
            group = groups.setdefault(key, {'header': header, 'rows': []})
            for r in rows[1:]:
                r = (r + [''] * len(header))[:len(header)]
                group['rows'].append(r + [path.name, ti])
    if not groups:
        raise ValueError('В выбранных документах нет таблиц с заголовком и данными')
    wb = Workbook()
    wb.remove(wb.active)
    ordered = sorted(groups.values(), key=lambda g: len(g['rows']), reverse=True)
    for i, g in enumerate(ordered, 1):
        name = 'Таблица' if i == 1 else f'Таблица {i}'
        ws = wb.create_sheet(name)
        ws.append(g['header'] + ['Источник: файл', 'Источник: таблица'])
        for r in g['rows']:
            append_as_values(ws, r, numbers=True)          # из таблиц Word: числа — числами (СУММ работает)
        _style_header(ws)
        ws.auto_filter.ref = ws.dimensions
    if failed:
        ws = wb.create_sheet('Не прочитано')
        ws.append(['Файл', 'Причина'])
        for f in failed:
            ws.append(list(f))
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return {'sheets': len(ordered), 'rows': sum(len(g['rows']) for g in ordered), 'failed': [f[0] for f in failed],
            'output': str(output)}


# ---------- Excel → Word: документы по образцу ----------

def template_fields(template):
    doc = Document(template)
    found = []
    for p in _all_paragraphs(doc):
        for m in PLACEHOLDER.finditer(p.text):
            if m.group(1) not in found:
                found.append(m.group(1))
    return found


def _all_paragraphs(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs
            for t in part.tables:
                for row in t.rows:
                    for cell in row.cells:
                        yield from cell.paragraphs


def _fill_paragraph(p, values):
    text = p.text
    if '{' not in text:
        return
    new = PLACEHOLDER.sub(lambda m: values.get(norm(m.group(1)), m.group(0)), text)
    if new == text:
        return
    # Word splits text into runs; the whole paragraph goes into the first run, keeping its formatting.
    runs = p.runs
    if not runs:
        return
    runs[0].text = new
    for r in runs[1:]:
        r.text = ''


def _value(v):
    if v is None:
        return ''
    if hasattr(v, 'strftime'):
        return v.strftime('%d.%m.%Y')
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _safe_name(s):
    s = re.sub(r'[\\/:*?"<>|\r\n\t]+', ' ', s).strip().rstrip('.')
    return s[:80] or 'Документ'


def mail_merge(xlsx, template, out_dir, name_field=None):
    """One Word document per Excel row: {Поле} in the template is replaced by the value from the column «Поле»."""
    xlsx, template, out_dir = Path(xlsx), Path(template), Path(out_dir)
    if template.suffix.lower() != '.docx':
        raise ValueError('Образец должен быть в формате .docx')
    fields = template_fields(template)
    if not fields:
        raise ValueError('В образце нет полей. Впишите в него названия столбцов в фигурных скобках, например {ФИО}.')
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        rows = wb.active.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else '' for h in next(rows, [])]
        data = [r for r in rows if any(v is not None and str(v).strip() for v in r)]
    finally:
        wb.close()
    columns = {norm(h): i for i, h in enumerate(header) if h}
    missing = [f for f in fields if norm(f) not in columns]
    if missing:
        raise ValueError('В Excel нет столбцов: ' + ', '.join(missing) + '. Заголовки должны совпадать с полями образца.')
    out_dir.mkdir(parents=True, exist_ok=True)
    created = []
    for n, row in enumerate(data, 1):
        values = {k: _value(row[i]) if i < len(row) else '' for k, i in columns.items()}
        doc = Document(template)
        for p in _all_paragraphs(doc):
            _fill_paragraph(p, values)
        base = _safe_name(values.get(norm(name_field), '') if name_field else '') if name_field else ''
        base = base if base and base != 'Документ' else f'{template.stem} {n:03d}'
        target = out_dir / f'{base}.docx'
        k = 2
        while target.exists():
            target = out_dir / f'{base} ({k}).docx'
            k += 1
        doc.save(target)
        created.append(target)
    return {'created': len(created), 'files': [str(p) for p in created], 'fields': fields, 'output': str(out_dir)}
