import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import iskra_highlight as H

TMP = Path(tempfile.mkdtemp())
HEAD = ['№', 'Город', 'Расстояние, км', 'Ответственный', 'Численность']
ROWS = [['1', 'Москва', '120', 'Иванов И.И.', '15'], ['2', 'Тверь', '180', 'Петров П.П.', '10'],
        ['3', 'Клин', '90', 'Сидоров С.С.', '8'], ['Итого', '', '', '', '33']]


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def make(name, head, rows, marks=(), cell_shade=(), paragraph=None):
    """marks: {(строка, столбец)} — выделить салатовым; cell_shade — залить ячейку 92D050."""
    doc = Document()
    doc.add_paragraph('ПЛАН ПЕРЕВОЗКИ')
    t = doc.add_table(rows=0, cols=len(head))
    for r_i, row in enumerate([head] + rows):
        cells = t.add_row().cells
        for c_i, v in enumerate(row):
            run = cells[c_i].paragraphs[0].add_run(v)
            if (r_i, c_i) in marks:
                run.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN
            if (r_i, c_i) in cell_shade:
                shade(cells[c_i], '92D050')
    p = doc.add_paragraph('Срок прибытия ')
    date = p.add_run(paragraph or '25.10.2026')
    if paragraph:
        date.font.highlight_color = WD_COLOR_INDEX.GREEN
    path = TMP / name
    doc.save(path)
    return path


def table_text(path):
    t = Document(str(path)).tables[0]
    out = []
    for row in t.rows:
        seen, cells = set(), []
        for c in row.cells:
            if id(c._tc) not in seen:
                seen.add(id(c._tc)); cells.append(c.text)
        out.append(cells)
    return out


class Merge(unittest.TestCase):
    def test_three_documents_into_template(self):
        template = make('образец.docx', HEAD, ROWS)
        a_rows = [r[:] for r in ROWS]; a_rows[1][2] = '185'; a_rows[1][3] = 'Смирнов А.А.'
        a = make('рота1.docx', HEAD, a_rows, marks={(2, 2)}, cell_shade={(2, 3)})
        # другой формат: без №, столбцы в другом порядке, своя подпись столбца
        head_b = ['Город', 'Ответственный', 'Численность, чел.', 'Расстояние (км)']
        b_rows = [['Москва', 'Иванов И.И.', '15', '120'], ['Клин', 'Сидоров С.С.', '12', '90'],
                  ['Дмитров', 'Козлов К.К.', '5', '70']]
        b = make('рота2.docx', head_b, b_rows, marks={(2, 2), (3, 0), (3, 1), (3, 2), (3, 3)})
        c_rows = [r[:] for r in ROWS]; c_rows[1][2] = '190'
        c = make('рота3.docx', HEAD, c_rows, marks={(2, 2)}, paragraph='27.10.2026')
        d = make('без_пометок.docx', HEAD, ROWS)
        out = TMP / 'итог.docx'
        rep = H.merge(template, [a, b, c, d], out)
        rows = table_text(out)
        tver = next(r for r in rows if r[1] == 'Тверь')
        self.assertIn('[УТОЧНИТЬ: 185 (рота1.docx) или 190 (рота3.docx)]', tver[2])
        self.assertEqual(tver[3], 'Смирнов А.А.')                      # залитая ячейка — тоже изменение
        klin = next(r for r in rows if r[1] == 'Клин')
        self.assertEqual(klin[4], '12')                                # из файла другого формата — в свой столбец
        names = [r[1] for r in rows]
        self.assertLess(names.index('Дмитров'), [r[0] for r in rows].index('Итого'))   # новая строка — перед «Итого»
        dm = rows[names.index('Дмитров')]
        self.assertEqual((dm[2], dm[3], dm[4]), ('70', 'Козлов К.К.', '5'))
        self.assertEqual(rows[-1][4], '42')                            # 15 + 10 + 12 + 5
        para = [p.text for p in Document(str(out)).paragraphs if p.text.startswith('Срок')][0]
        self.assertEqual(para, 'Срок прибытия 27.10.2026')
        self.assertEqual(rep.files_without_marks, ['без_пометок.docx'])
        self.assertEqual(len(rep.conflicts), 1)
        self.assertIn('33 → 42', ' '.join(rep.totals))
        self.assertEqual(rep.unplaced, [])
        # изменения выделены в итоговом документе
        cell = Document(str(out)).tables[0].rows[3].cells[4]
        self.assertTrue(any(r.font.highlight_color == WD_COLOR_INDEX.BRIGHT_GREEN for p in cell.paragraphs for r in p.runs))
        self.assertIn('Перенесено изменений', rep.text())

    def test_unknown_row_is_reported(self):
        template = make('о2.docx', HEAD, ROWS)
        rows = [r[:] for r in ROWS]; rows[2] = ['3', 'Казань', '800', 'Орлов О.О.', '4']
        src = make('чужой.docx', HEAD, rows, marks={(3, 4)})
        rep = H.merge(template, [src], TMP / 'и2.docx')
        self.assertEqual(len(rep.unplaced), 1)
        self.assertIn('Казань', rep.unplaced[0])

    def test_two_level_header_and_other_layout(self):
        def doc_two_level(name, rows, marks=()):
            d = Document()
            t = d.add_table(rows=2 + len(rows), cols=4)
            h = t.rows[0].cells
            h[0].text = 'Город'; h[1].text = 'Расстояние, км'; h[3].text = 'Численность'
            h[1].merge(h[2]); t.cell(0, 0).merge(t.cell(1, 0)); t.cell(0, 3).merge(t.cell(1, 3))
            t.cell(1, 1).text = 'туда'; t.cell(1, 2).text = 'обратно'
            for i, row in enumerate(rows, 2):
                for k, v in enumerate(row):
                    run = t.cell(i, k).paragraphs[0].add_run(v)
                    if (i, k) in marks:
                        run.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN
            d.save(TMP / name)
            return TMP / name
        base = [['Москва', '120', '125', '15'], ['Тверь', '180', '180', '10']]
        template = doc_two_level('шаблон2.docx', base)
        changed = [r[:] for r in base]; changed[1][2] = '200'
        src = doc_two_level('изм2.docx', changed, marks={(3, 2)})
        flat = make('плоский.docx', ['Город', 'Численность', 'Расстояние туда', 'Расстояние обратно'],
                    [['Москва', '17', '120', '125']], marks={(1, 1)})
        out = TMP / 'и3.docx'
        rep = H.merge(template, [src, flat], out)
        t = Document(str(out)).tables[0]
        self.assertEqual(t.cell(3, 2).text, '200')
        self.assertEqual(t.cell(3, 1).text, '180')
        self.assertEqual(t.cell(2, 3).text, '17')                      # из файла с другой шапкой
        self.assertEqual(rep.unplaced, [])

    def test_screenshot_case_green_font_and_coloured_rows(self):
        """Как на снимке: шапка, строка-раздел, строки «77 км»/«105 км» с заливкой образца, изменения — зелёным шрифтом."""
        from docx.shared import RGBColor
        heads = ['Наименование ПУ', 'Подчиненные в/ч', 'Удаление от линии фронта', 'Старший на ПУ', 'Телефоны',
                 'Количество ОС/ВВСТ', 'Положение', 'Характер ПУ, количество запасных районов', 'Время нахождения в районе']

        def build(name, kp, green=(), extra=None):
            d = Document()
            t = d.add_table(rows=6 + (1 if extra else 0) + 3, cols=9)
            for k, h in enumerate(heads):
                t.cell(0, k).text = h
            sec = t.cell(1, 0).merge(t.cell(1, 8)); sec.text = 'СЕВЕРОДОНЕЦКОЕ НАПРАВЛЕНИЕ'
            for r, km in ((2, '77 км'), (3, '105 км')):
                t.cell(r, 0).merge(t.cell(r, 1))
                t.cell(r, 2).text = km
                for k in range(9):
                    shade(t.cell(r, k), 'C2D69B' if k % 2 else '9BBB59')   # «оливковая» раскраска образца
            for k, v in enumerate(kp):
                cell = t.cell(4, k)
                for i, line in enumerate(v.split('\n')):
                    par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
                    run = par.add_run(line)
                    if (k, i) in green or (k, -1) in green:
                        run.font.color.rgb = RGBColor(0x00, 0xB0, 0x50)
            if extra:
                for k, v in enumerate(extra):
                    run = t.cell(5, k).paragraphs[0].add_run(v)
                    run.font.color.rgb = RGBColor(0x00, 0xB0, 0x50)
            d.save(TMP / name)
            return TMP / name
        was = ['КП 5447 обр СпН', '', '56 км', 'НШ гв п/п-к Залупа В.А.', '12188 4444\n18197 4400\n11741 4444', '27/2',
               'с-в окраина ЗАЛУПИНСК', 'Заброшенное помещение в промышленной зоне, 1 запасной район', '4 сут. с 19.10.2026']
        now = was[:]; now[2] = '43 км'; now[5] = '27/4'; now[6] = 'с-в окраина ЗАЛУПИНСК\nПЕЧОРИНСК'
        template = build('ПУ_образец.docx', was)
        new_kp = ['КП 101 мсп', '', '30 км', 'НШ м-р Иванов И.И.', '11111 2222', '12/3', 'южнее ЛИСИЧАНСК', 'палатки', '2 сут.']
        sent = build('ПУ_рота.docx', now, green={(0, -1), (2, -1), (5, -1), (6, 1)}, extra=new_kp)
        out = TMP / 'ПУ_итог.docx'
        rep = H.merge(template, [sent], out)
        t = Document(str(out)).tables[0]
        self.assertEqual(t.cell(4, 2).text, '43 км')
        self.assertEqual(t.cell(4, 5).text, '27/4')
        self.assertEqual(t.cell(4, 6).text, 'с-в окраина ЗАЛУПИНСК\nПЕЧОРИНСК')
        self.assertEqual(t.cell(4, 0).text, 'КП 5447 обр СпН')            # отмечено, текст тот же — перенесено
        self.assertTrue(any(r.font.highlight_color for r in t.cell(4, 0).paragraphs[0].runs))   # и выделено как изменение
        self.assertEqual(t.cell(2, 2).text, '77 км')                      # заливка образца — не изменение
        self.assertEqual(len(t.rows), 9)                                   # новые строки не добавлены лишние
        self.assertEqual(t.cell(5, 0).text, 'КП 101 мсп')                 # новая — в первую пустую строку раздела
        self.assertEqual(t.cell(5, 3).text, 'НШ м-р Иванов И.И.')
        self.assertEqual(rep.unplaced, [])
        self.assertEqual(len(rep.placed), 5)
        self.assertTrue(any('то же значение' in p for p in rep.placed))
        from datetime import date
        days = t.cell(4, 8)                                  # положение сменилось («ПЕЧОРИНСК») — сутки заново, зелёным
        self.assertEqual(days.text, f'1 сут. с {date.today():%d.%m.%Y}')
        self.assertTrue(any(r.font.highlight_color for r in days.paragraphs[0].runs))
        self.assertEqual(len(rep.added_rows), 1)
        self.assertIn('СЕВЕРОДОНЕЦКОЕ', rep.added_rows[0])

    def test_full_table_many_sections_no_empty_rows(self):
        """Реальный документ: всё заполнено, несколько направлений, одинаковые старшие/телефоны в разных строках."""
        from docx.shared import RGBColor
        heads = ['Наименование ПУ', 'Удаление от линии фронта', 'Старший на ПУ', 'Количество ОС/ВВСТ', 'Положение']
        sections = {'СЕВЕРОДОНЕЦКОЕ НАПРАВЛЕНИЕ': [['КП 5447', '56 км', 'НШ Залупа В.А.', '27/2', 'ЗАЛУПИНСК'],
                                                  ['КП 5448', '61 км', 'НШ Залупа В.А.', '12/1', 'ЛИСИЧАНСК'],
                                                  ['ЗКП 5447', '80 км', 'ЗНШ Петров П.П.', '9/0', 'КРЕМЕННАЯ']],
                    'ЛИМАНСКОЕ НАПРАВЛЕНИЕ': [['КП 7001', '40 км', 'НШ Сидоров С.С.', '30/5', 'ЛИМАН'],
                                              ['КП 7002', '45 км', 'НШ Сидоров С.С.', '18/2', 'ЯМПОЛЬ']]}

        def build(name, changes=(), extra=None):
            d = Document()
            t = d.add_table(rows=1, cols=5)
            for k, h in enumerate(heads):
                t.cell(0, k).text = h
            for sec, rows in sections.items():
                r = t.add_row(); c = r.cells[0].merge(r.cells[4]); c.text = sec
                for row in rows + ([extra[1]] if extra and extra[0] == sec else []):
                    cells = t.add_row().cells
                    for k, v in enumerate(row):
                        new = dict(changes).get((row[0], k))
                        run = cells[k].paragraphs[0].add_run(new or v)
                        if new or (extra and row is extra[1]):
                            run.font.color.rgb = RGBColor(0x00, 0xB0, 0x50)
            d.save(TMP / name)
            return TMP / name
        template = build('полный_образец.docx')
        a = build('полный_1.docx', changes={('КП 5448', 1): '58 км', ('КП 7002', 4): 'ЯМПОЛЬ, СЕВЕРСК'})
        b = build('полный_2.docx', changes={('КП 5447', 3): '27/4'},
                  extra=('ЛИМАНСКОЕ НАПРАВЛЕНИЕ', ['КП 7003', '50 км', 'НШ Орлов О.О.', '5/1', 'ДРОБЫШЕВО']))
        out = TMP / 'полный_итог.docx'
        rep = H.merge(template, [a, b], out)
        rows = [[c.text for c in r.cells] for r in Document(str(out)).tables[0].rows]
        by = {r[0]: r for r in rows}
        self.assertEqual(by['КП 5448'][1], '58 км')
        self.assertEqual(by['КП 5447'][1], '56 км')                     # тот же старший — но не та строка
        self.assertEqual(by['КП 5447'][3], '27/4')
        self.assertEqual(by['КП 7002'][4], 'ЯМПОЛЬ, СЕВЕРСК')
        names = [r[0] for r in rows]
        self.assertEqual(names.index('КП 7003'), names.index('КП 7002') + 1)   # новая — в конце своего раздела
        self.assertEqual(rep.unplaced, [])
        self.assertEqual(len(rep.placed), 3)

    def test_unmarked_differences_and_accept(self):
        template = make('у_образец.docx', HEAD, ROWS)
        rows = [r[:] for r in ROWS]
        rows[0][2] = '125'                                     # отмечено зелёным
        rows[1][4] = '11'                                      # изменено, но не отмечено
        rows.insert(3, ['4', 'Ржев', '230', 'Орлов О.О.', '7'])  # новая строка без отметки
        src = make('у_рота.docx', HEAD, rows, marks={(1, 2)})
        out = TMP / 'у_итог.docx'
        rep = H.merge(template, [src], out)
        self.assertEqual(len(rep.placed), 1)
        kinds = {(u['object'], u['column']): u for u in rep.unmarked}
        self.assertEqual(kinds[('Тверь', 'Численность')]['new'], '11')
        self.assertEqual(kinds[('Тверь', 'Численность')]['old'], '10')
        self.assertIn(('Ржев', 'новая строка'), kinds)
        self.assertEqual(table_text(out)[2][4], '10')          # без выбора — не перенесено
        rep2 = H.merge(template, [src], out, accept=[u['id'] for u in rep.unmarked])
        rows_out = table_text(out)
        self.assertEqual(rows_out[2][4], '11')
        self.assertIn('Ржев', [r[1] for r in rows_out])
        self.assertEqual(rep2.unmarked, [])
        self.assertEqual(rows_out[-1][4], '41')                 # 15 + 11 + 8 + 7
        change = next(c for c in rep2.changes if c.get('column') == 'Численность')
        self.assertEqual((change['object'], change['old'], change['new']), ('Тверь', '10', '11'))
        self.assertTrue(any(c.get('new_row') for c in rep2.changes))
        snap = rep2.snapshot[0]['rows']
        self.assertEqual(next(r for r in snap if r['object'] == 'Ржев')['cells']['Численность'], '7')

    def test_days_recount_and_threshold(self):
        from datetime import date
        d = Document()
        t = d.add_table(rows=3, cols=2)
        t.cell(0, 0).text = 'ПУ'; t.cell(0, 1).text = 'Время нахождения в районе'
        t.cell(1, 0).text = 'КП 1'
        r = t.cell(1, 1).paragraphs[0]; r.add_run('4 сут. с 19.10.2026')
        t.cell(2, 0).text = 'КП 2'; t.cell(2, 1).text = '1 сут. с 23.10.26'
        path = TMP / 'сутки.docx'
        d.save(path)
        rep = H.recount_days_file(path, today=date(2026, 10, 25), threshold=5)
        t = Document(str(path)).tables[0]
        self.assertEqual(t.cell(1, 1).text, '7 сут. с 19.10.2026')     # 19…25 октября — 7 суток
        self.assertEqual(t.cell(2, 1).text, '3 сут. с 23.10.26')
        self.assertEqual(len(rep.days), 2)
        self.assertEqual(rep.long_stay, ['КП 1 — 7 сут. с 19.10.2026'])
        self.assertEqual(t.cell(1, 1).paragraphs[0].runs[0].font.highlight_color, WD_COLOR_INDEX.YELLOW)

    def test_colors(self):
        self.assertTrue(H._green_hex('92D050'))
        self.assertTrue(H._green_hex('00FF00'))
        self.assertTrue(H._green_hex('C6EFCE'))
        self.assertFalse(H._green_hex('FFFF00'))
        self.assertFalse(H._green_hex('D9D9D9'))


class SameValue(unittest.TestCase):
    def test_same_name_marked_green_or_lime_is_transferred(self):
        """«Авдеевка» → «Авдеевка», но написано зелёным (шрифт), салатовым маркером или салатовой заливкой — переносится."""
        from docx.shared import RGBColor
        from docx.enum.text import WD_COLOR_INDEX

        def build(name, how=None, template_green=False):
            d = Document()
            t = d.add_table(rows=3, cols=3)
            for k, h in enumerate(['Наименование ПУ', 'Положение', 'Удаление']):
                t.cell(0, k).text = h
            for r, (kp, place, km) in enumerate((('КП 1', 'Авдеевка', '20 км'), ('КП 2', 'Покровск', '35 км')), 1):
                t.cell(r, 0).text, t.cell(r, 2).text = kp, km
                run = t.cell(r, 1).paragraphs[0].add_run(place)
                if r == 1 and how == 'font':
                    run.font.color.rgb = RGBColor(0x00, 0xB0, 0x50)
                if r == 1 and how == 'marker':
                    run.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN
                if r == 1 and how == 'fill':
                    shade(t.cell(r, 1), '92D050')
            if template_green:
                shade(t.cell(2, 2), '92D050')              # в образце прошлой недели — своя салатовая ячейка
            d.save(TMP / name)
            return TMP / name
        for how in ('font', 'marker', 'fill'):
            template = build('образец_авд.docx', template_green=True)
            out = TMP / f'итог_{how}.docx'
            rep = H.merge(template, [build(f'присланный_{how}.docx', how)], out)
            cell = Document(str(out)).tables[0].cell(1, 1)
            self.assertEqual(cell.text, 'Авдеевка', how)
            self.assertTrue(any(r.font.highlight_color for r in cell.paragraphs[0].runs), how)
            self.assertEqual(len(rep.placed), 1, how)
            self.assertIn('Авдеевка', rep.placed[0])



class ThreeDocuments(unittest.TestCase):
    """Шаблон — как должно выглядеть; документ — куда вносить; три присланных — откуда брать. Изменения — зелёным шрифтом."""

    def build(self, name, rows, green=(), font=None, size=None, old_green=()):
        from docx.shared import Pt, RGBColor
        d = Document()
        t = d.add_table(rows=1 + len(rows), cols=3)
        for k, h in enumerate(['Наименование ПУ', 'Положение', 'Удаление']):
            t.cell(0, k).text = h
        for r, values in enumerate(rows, 1):
            for k, v in enumerate(values):
                run = t.cell(r, k).paragraphs[0].add_run(v)
                if font:
                    run.font.name, run.font.size = font, Pt(size)
                if (r, k) in green or (r, k) in old_green:
                    run.font.color.rgb = RGBColor(0x00, 0xB0, 0x50)
        d.save(TMP / name)
        return TMP / name

    def test_template_target_and_three_sources(self):
        rows = [['КП 1', 'Авдеевка', '20 км'], ['КП 2', 'Покровск', '35 км'], ['КП 3', 'Курахово', '15 км']]
        template = self.build('шаблон3.docx', [['', '', '']] * 3, font='Times New Roman', size=12)
        target = self.build('документ3.docx', rows, font='Arial', size=9, old_green={(2, 2)})   # отметка с прошлого раза
        s1 = self.build('рота1.docx', [rows[0][:1] + ['Авдеевка', '22 км'], rows[1], rows[2]], green={(1, 1), (1, 2)})
        s2 = self.build('рота2.docx', [rows[0], ['КП 2', 'Мирноград', '35 км'], rows[2]], green={(2, 1)})
        s3 = self.build('рота3.docx', [rows[0], rows[1], ['КП 3', 'Курахово', '12 км']])      # без отметки
        out = TMP / 'итог3.docx'
        rep = H.merge(template, [s1, s2, s3], out, mark=H.FONT, base=target, take_all=True)
        t = Document(str(out)).tables[0]

        def green(r, k):
            return any(run.font.color.rgb is not None and str(run.font.color.rgb) == '00B050'
                       for run in t.cell(r, k).paragraphs[0].runs)
        self.assertEqual([t.cell(r, 1).text for r in (1, 2, 3)], ['Авдеевка', 'Мирноград', 'Курахово'])
        self.assertEqual([t.cell(r, 2).text for r in (1, 2, 3)], ['22 км', '35 км', '12 км'])
        self.assertTrue(green(1, 1) and green(1, 2) and green(2, 1) and green(3, 2))   # все изменения — зелёным шрифтом
        self.assertFalse(green(1, 0) or green(2, 2))           # не менялось — обычным; старая отметка снята
        self.assertFalse(any(r.font.highlight_color for r in t.cell(1, 1).paragraphs[0].runs))
        run = t.cell(2, 0).paragraphs[0].runs[0]
        self.assertEqual((run.font.name, run.font.size.pt), ('Times New Roman', 12))   # вид — как в шаблоне
        self.assertEqual(len(rep.placed), 4)
        self.assertEqual(rep.unmarked, [])


    def test_strict_template_font_everywhere(self):
        """Шаблон: Times New Roman 14 через стиль «Обычный», пустые ячейки без своего шрифта. Документ: Arial 9 и
        Calibri в отдельных ячейках. Итог — везде Times New Roman 14: таблица, новая строка, заголовок над таблицей."""
        from docx.shared import Pt, RGBColor
        from docx.oxml.ns import qn as q

        def doc(name, normal, size, title, rows, extra_font=None, green=(), spare=0):
            d = Document()
            d.styles['Normal'].font.name, d.styles['Normal'].font.size = normal, Pt(size)
            d.add_paragraph(title)
            t = d.add_table(rows=1 + len(rows) + spare, cols=3)
            for k, h in enumerate(['Наименование ПУ', 'Положение', 'Удаление']):
                t.cell(0, k).text = h
            for r, values in enumerate(rows, 1):
                for k, v in enumerate(values):
                    run = t.cell(r, k).paragraphs[0].add_run(v)
                    if extra_font and k == 1:
                        run.font.name = extra_font
                    if (r, k) in green or (r, -1) in green:
                        run.font.color.rgb = RGBColor(0x00, 0xB0, 0x50)
            d.add_paragraph('Начальник штаба')
            d.save(TMP / name)
            return TMP / name
        rows = [['КП 1', 'Авдеевка', '20 км'], ['КП 2', 'Покровск', '35 км']]
        template = doc('шаблон_строгий.docx', 'Times New Roman', 14, 'Сведения о ПУ', [['', '', ''], ['', '', '']], spare=1)
        target = doc('документ_разный.docx', 'Arial', 9, 'Сведения о ПУ', rows, extra_font='Calibri', spare=1)
        sent = doc('рота_строгий.docx', 'Courier New', 10, 'Сведения о ПУ',
                   rows + [['КП 9', 'Мирноград', '40 км']], green={(1, 2), (3, -1)})
        out = TMP / 'итог_строгий.docx'
        H.merge(template, [sent], out, mark=H.FONT, base=target)
        result = Document(str(out))
        self.assertIn('КП 9', [c.text for row in result.tables[0].rows for c in row.cells])   # новая строка есть
        fonts = set()
        for r in result.element.body.iter(q('w:r')):
            if not ''.join(t.text or '' for t in r.iter(q('w:t'))).strip():
                continue
            rpr = r.find(q('w:rPr'))
            f = rpr.find(q('w:rFonts')) if rpr is not None else None
            z = rpr.find(q('w:sz')) if rpr is not None else None
            fonts.add((f.get(q('w:ascii')) if f is not None else None, z.get(q('w:val')) if z is not None else None))
        self.assertEqual(fonts, {('Times New Roman', '28')})


class StayRestart(unittest.TestCase):
    def test_moved_unit_days_restart_unchanged_unit_days_recounted(self):
        """Было: 56 км, ЗАЛУПИНСК, «4 сут. с 19.10.2026». Стало: 43 км, ЗАЛУПИНСК + ПЕЧОРИНСК (зелёным) —
        сутки «1 сут. с 23.10.2026» зелёным. Строка без смены положения — сутки просто пересчитаны, без зелёного."""
        from datetime import date
        from docx.shared import RGBColor
        heads = ['Наименование ПУ', 'Удаление от линии фронта', 'Положение', 'Время нахождения в районе']

        def build(name, rows, green=()):
            d = Document()
            t = d.add_table(rows=1 + len(rows), cols=4)
            for k, h in enumerate(heads):
                t.cell(0, k).text = h
            for r, values in enumerate(rows, 1):
                for k, v in enumerate(values):
                    for i, line in enumerate(v.split('\n')):
                        par = t.cell(r, k).paragraphs[0] if i == 0 else t.cell(r, k).add_paragraph()
                        run = par.add_run(line)
                        if (r, k, i) in green:
                            run.font.color.rgb = RGBColor(0x00, 0xB0, 0x50)
            d.save(TMP / name)
            return TMP / name
        was = [['КП 5447 обр СпН', '56 км', 'с-в окраина ЗАЛУПИНСК', '4 сут. с 19.10.2026'],
               ['КП 7001', '40 км', 'ЛИМАН', '2 сут. с 21.10.2026']]
        now = [['КП 5447 обр СпН', '43 км', 'с-в окраина ЗАЛУПИНСК\nПЕЧОРИНСК', '4 сут. с 19.10.2026'], was[1]]
        doc = build('было_сут.docx', was)
        sent = build('стало_сут.docx', now, green={(1, 1, 0), (1, 2, 1)})
        out = TMP / 'итог_сут.docx'
        H.merge(doc, [sent], out, mark=H.FONT, base=doc, today=date(2026, 10, 23))
        t = Document(str(out)).tables[0]

        def green(cell):
            return any(r.font.color.rgb is not None and str(r.font.color.rgb) == '00B050'
                       for p in cell.paragraphs for r in p.runs)
        self.assertEqual(t.cell(1, 3).text, '1 сут. с 23.10.2026')
        self.assertTrue(green(t.cell(1, 3)))
        self.assertEqual(t.cell(1, 2).text, 'с-в окраина ЗАЛУПИНСК\nПЕЧОРИНСК')
        self.assertTrue(green(t.cell(1, 1)))
        self.assertEqual(t.cell(2, 3).text, '3 сут. с 21.10.2026')     # не переезжали — пересчёт
        self.assertFalse(green(t.cell(2, 3)))
        place = t.cell(1, 2).paragraphs                                  # старый пункт — чёрным, новый снизу — зелёным
        self.assertFalse(any(str(r.font.color.rgb) == '00B050' for r in place[0].runs if r.font.color.rgb))
        self.assertTrue(all(str(r.font.color.rgb) == '00B050' for r in place[1].runs))
        # На следующий день поменялся только КП 7001 — вчерашний зелёный у КП 5447 становится чёрным.
        today_doc = TMP / 'итог_сут.docx'
        next_now = [now[0], ['КП 7001', '35 км', 'ЛИМАН', '3 сут. с 21.10.2026']]
        sent2 = build('стало_сут2.docx', next_now, green={(2, 1, 0)})
        out2 = TMP / 'итог_сут2.docx'
        H.merge(today_doc, [sent2], out2, mark=H.FONT, base=today_doc, today=date(2026, 10, 24))
        t2 = Document(str(out2)).tables[0]
        self.assertFalse(green(t2.cell(1, 1)) or green(t2.cell(1, 2)) or green(t2.cell(1, 3)))
        self.assertEqual(t2.cell(1, 3).text, '2 сут. с 23.10.2026')     # пересчёт — чёрным
        self.assertTrue(green(t2.cell(2, 1)))
        self.assertEqual(t2.cell(2, 1).text, '35 км')


class StyledTemplate(unittest.TestCase):
    def test_template_paragraphs_with_own_styles(self):
        """Отчёт об ошибке 24.10: шаблон пользователя, у абзацев свои стили (Standard, Text body…) без шрифта
        в тексте — перенос падал (AttributeError: 'NoneType' object has no attribute 'part')."""
        from docx.shared import Pt
        from docx.enum.style import WD_STYLE_TYPE
        from docx.oxml.ns import qn as q
        d = Document()
        st = d.styles.add_style('Text body', WD_STYLE_TYPE.PARAGRAPH)
        st.font.name, st.font.size = 'Times New Roman', Pt(10)
        d.add_paragraph('Сведения о ПУ', style=st)
        t = d.add_table(rows=3, cols=2)
        for k, h in enumerate(['Наименование ПУ', 'Положение']):
            t.cell(0, k).paragraphs[0].style = st
            t.cell(0, k).paragraphs[0].add_run(h)
        for r in (1, 2):
            for k in range(2):
                t.cell(r, k).paragraphs[0].style = st
        t.cell(1, 0).paragraphs[0].add_run('КП 5447')
        t.cell(1, 1).paragraphs[0].add_run('ЗАЛУПИНСК')
        d.save(TMP / 'шаблон_стили.docx')
        sent = Document()
        st2 = sent.add_table(rows=2, cols=2)
        st2.cell(0, 0).text, st2.cell(0, 1).text, st2.cell(1, 0).text = 'Наименование ПУ', 'Положение', 'КП 5447'
        from docx.shared import RGBColor
        run = st2.cell(1, 1).paragraphs[0].add_run('ПЕЧОРИНСК')
        run.font.color.rgb = RGBColor(0x00, 0xB0, 0x50)
        sent.save(TMP / 'рота_стили.docx')
        out = TMP / 'итог_стили.docx'
        H.merge(TMP / 'шаблон_стили.docx', [TMP / 'рота_стили.docx'], out, mark=H.FONT, base=TMP / 'шаблон_стили.docx')
        res = Document(str(out))
        self.assertEqual(res.tables[0].cell(1, 1).text, 'ПЕЧОРИНСК')
        rpr = res.tables[0].cell(1, 1).paragraphs[0].runs[0]._r.find(q('w:rPr'))
        self.assertEqual(rpr.find(q('w:rFonts')).get(q('w:ascii')), 'Times New Roman')
        self.assertEqual(rpr.find(q('w:sz')).get(q('w:val')), '20')          # 10 пт — из стиля шаблона


if __name__ == '__main__':
    unittest.main()
