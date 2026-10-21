import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from iskra_svod import fold, read_tables, combine, svod_files, describe

ORDER = ['Госпиталь', 'Полигон', 'Модуль 1', 'Модуль 2', 'Модуль 3', 'Учебный центр', 'Склад']


def book(path, rows, merges=()):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    for m in merges:
        ws.merge_cells(m)
    wb.save(path)
    return path


class SvodTest(unittest.TestCase):
    def setUp(self):
        self.d = Path(tempfile.mkdtemp())
        d = self.d
        # 1. Заголовок документа, двухстрочная шапка с объединением, строка нумерации, «Итого».
        self.a = book(d / 'a.xlsx', [
            ['Сведения о численности личного состава организаций', None, None, None],
            ['№ п/п', 'Наименование организации', 'Личный состав', None],
            [None, None, 'по списку', 'на лицо'],
            [1, 2, 3, 4],
            [1, 'Госпиталь', 120, 115],
            [2, 'Полигон', 80, 78],
            [None, 'Итого', 200, 193],
        ], merges=('A1:D1', 'A2:A3', 'B2:B3', 'C2:D2'))
        # 2. Простая шапка, «Модуль №1», числа текстом, примечание.
        self.b = book(d / 'b.xlsx', [
            ['Наименование', 'По списку', 'На лицо', 'Примечание'],
            ['Модуль №1', 40, 38, 'в командировке 2'],
            ['МОДУЛЬ 2', '35', '35', None],
        ])
        # 3. «На лицо» больше «по списку» — ошибка в данных.
        self.c = book(d / 'c.xlsx', [
            ['Организация', 'по списку', 'на лицо'],
            ['Учебный центр', 50, 52],
        ])
        # 4. Итог в файле не сходится со строками.
        self.e = book(d / 'e.xlsx', [
            ['Подразделение', 'По списку', 'На лицо'],
            ['Склад', 10, 9],
            ['Итого', 20, 9],
        ])
        # 5. Опечатка в названии и повтор «Модуль-1» (складывается с b.xlsx).
        self.f = book(d / 'f.xlsx', [
            ['Наименование организации', 'По списку', 'На лицо'],
            ['Полегон', 5, 5],
            ['Модуль-1', 2, 2],
        ])
        self.files = [self.a, self.b, self.c, self.e, self.f]

    def test_fold(self):
        self.assertEqual(fold('Модуль №1'), fold('модуль 1'))
        self.assertEqual(fold('МОДУЛЬ-1'), fold('модуль 1'))
        self.assertEqual(fold('«Госпиталь»'), 'госпиталь')
        self.assertNotEqual(fold('Модуль 1'), fold('Модуль 2'))

    def test_two_row_header(self):
        t = read_tables(self.a)[0]
        self.assertEqual(t.columns, ['Личный состав / по списку', 'Личный состав / на лицо'])
        self.assertEqual([r[0] for r in t.rows], ['Госпиталь', 'Полигон'])
        self.assertEqual(t.totals, {'Личный состав / по списку': 200, 'Личный состав / на лицо': 193})

    def test_combine(self):
        tables = [t for p in self.files for t in read_tables(p)]
        s = combine(tables, ORDER)
        # «по списку» из простых шапок и «Личный состав / по списку» — один столбец.
        self.assertEqual(s.columns, ['Личный состав / по списку', 'Личный состав / на лицо', 'Примечание'])
        names = [o.name for o in s.orgs]
        self.assertEqual(names[:7], ORDER)
        self.assertEqual(names[7:], ['Полегон'])
        m1 = s.orgs[2]
        self.assertEqual(s.value(m1, 'Личный состав / по списку'), (42, False))
        self.assertEqual(s.value(s.orgs[3], 'Личный состав / на лицо'), (35, False))
        kinds = {(r[0], r[1]) for r in s.report}
        self.assertIn(('Нет данных', 'Модуль 3'), kinds)
        self.assertIn(('Ошибка в данных', 'Учебный центр'), kinds)
        self.assertIn(('Сложено', 'Модуль 1'), kinds)
        self.assertTrue(any(r[0] == 'Нет в перечне' and 'Полигон' in r[3] for r in s.report))
        self.assertTrue(any(r[0] == 'Итог не сходится' and r[2] == 'e.xlsx' for r in s.report))
        self.assertFalse(any(r[0] == 'Итог не сходится' and r[2] == 'a.xlsx' for r in s.report))

    def test_check_mode_marks_conflict(self):
        g = book(self.d / 'g.xlsx', [['Наименование', 'По списку'], ['Госпиталь', 119]])
        tables = [t for p in (self.a, g) for t in read_tables(p)]
        s = combine(tables, mode='check')
        hosp = s.orgs[0]
        self.assertEqual(s.value(hosp, 'Личный состав / по списку'), (119, True))
        self.assertTrue(any(r[0] == 'Расхождение' and r[1] == 'Госпиталь' for r in s.report))

    def test_save_and_form(self):
        form = book(self.d / 'form.xlsx', [
            ['№', 'Наименование', 'По списку', 'На лицо'],
            [1, 'Госпиталь', None, None],
            [2, 'Модуль 1', None, None],
            [3, 'Модуль 3', None, None],
            [None, 'Итого', '=SUM(C2:C4)', '=SUM(D2:D4)'],
        ])
        out = self.d / 'свод.xlsx'
        r = svod_files(self.files + [self.d / 'нет.xls'], out, ORDER, form=form)
        self.assertEqual(r['orgs'], 8)
        self.assertTrue(any(f[0] == 'нет.xls' for f in r['failed']))
        wb = load_workbook(out)
        ws = wb['Свод']
        self.assertEqual(ws['B2'].value, 'Наименование организации')
        self.assertEqual([ws.cell(i, 2).value for i in range(3, 11)], ORDER + ['Полегон'])
        self.assertEqual(ws['C3'].value, 120)
        total = [ws.cell(ws.max_row, j).value for j in range(1, 6)]
        self.assertEqual(total[1], 'ИТОГО')
        self.assertEqual(total[2], '=SUM(C3:C10)')
        self.assertIn(total[4], ('', None))            # примечание не суммируется
        self.assertIn('Проверка', wb.sheetnames)
        self.assertGreater(wb['Исходные данные'].max_row, 5)
        f = load_workbook(r['form_output'])['Sheet']
        self.assertEqual((f['C2'].value, f['D2'].value), (120, 115))
        self.assertEqual((f['C3'].value, f['D3'].value), (42, 40))
        self.assertEqual(f['C5'].value, '=SUM(C2:C4)')
        self.assertEqual(r['form_missing'], ['Модуль 3'])
        self.assertIn('Полигон', r['form_extra'])
        self.assertIn('Нет данных: 1', describe(r))


if __name__ == '__main__':
    unittest.main()
