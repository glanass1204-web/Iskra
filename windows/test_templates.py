"""Несколько шаблонов: свой для каждого вида документа; «напиши рапорт…» — сам берётся шаблон рапорта."""
import tempfile
import unittest
from datetime import date
from pathlib import Path

from docx import Document

import iskra_ai as ai
import iskra_forms as forms

TMP = Path(tempfile.mkdtemp())


def template(name, lines):
    d = Document()
    for line in lines:
        d.add_paragraph(line)
    d.save(TMP / name)
    return str(TMP / name)


class Choose(unittest.TestCase):
    def setUp(self):
        self.old = dict(ai.TEMPLATES), ai.TELEGRAM_TEMPLATE
        self.tlg = template('тлг.docx', ['ТЕЛЕГРАММА', 'Командирам подразделений', 'Текст.', 'Начальник штаба'])
        self.rap = template('рапорт.docx', ['Командиру войсковой части 12345', 'РАПОРТ', 'Текст рапорта.', 'капитан Иванов'])
        self.spr = template('справка.docx', ['СПРАВКА-ДОКЛАД', 'Текст.'])
        ai.TEMPLATES = {'Телеграмма': self.tlg, 'Рапорт': self.rap, 'Справка-доклад': self.spr}

    def tearDown(self):
        ai.TEMPLATES, ai.TELEGRAM_TEMPLATE = self.old

    def test_kind_and_template_by_task(self):
        self.assertEqual(ai.template_for('напиши рапорт на отпуск'), self.rap)
        self.assertEqual(ai.template_for('составь указание командирам о проверке техники'), self.tlg)   # указание — по тлг
        self.assertEqual(ai.template_for('подготовь справку-доклад о технике'), self.spr)               # свой вид
        self.assertIsNone(ai.template_for('подготовь отпускной билет'))       # своего бланка нет — стандартный

    def test_write_uses_kind_template(self):
        prompts = []
        old = ai.run
        ai.run = lambda prompt, **kw: prompts.append(prompt) or 'Командиру войсковой части 12345\nРАПОРТ\nПрошу предоставить отпуск.'
        try:
            ai.write('напиши рапорт на отпуск', [])
            self.assertIn('Текст рапорта', prompts[-1])           # образец оформления — из шаблона рапорта
            self.assertNotIn('Командирам подразделений', prompts[-1])
            ai.write('напиши телеграмму командирам о проверке', [])
            self.assertIn('Командирам подразделений', prompts[-1])
        finally:
            ai.run = old
        self.assertEqual(ai.TELEGRAM_TEMPLATE, self.old[1])         # после работы — как было

    def test_combat_order_not_written(self):
        with self.assertRaises(ai.AIError):
            ai.write('напиши боевое распоряжение', [])


class LeaveTicket(unittest.TestCase):
    def test_own_blank_with_labels(self):
        blank = template('отпускной.docx', ['{Штамп}', 'ОТПУСКНОЙ БИЛЕТ', '{ФИО} убыл {Вид отпуска}',
                                             'сроком на {Сроком на} с {С} по {По}', 'Командир {Командир}'])
        f = {'stamp': 'в/ч 12345', 'who': 'капитан Иванов Иван Иванович', 'kind_place': 'в основной отпуск в г. Тверь',
             'days': '15', 'start': date(2026, 11, 1), 'end': date(2026, 11, 15), 'commander': 'полковник Петров'}
        out = TMP / 'билет.docx'
        forms.save_leave_ticket(f, out, blank)
        text = [p.text for p in Document(str(out)).paragraphs]
        self.assertEqual(text[0], 'в/ч 12345')
        self.assertEqual(text[2], 'капитан Иванов Иван Иванович убыл в основной отпуск в г. Тверь')
        self.assertEqual(text[3], 'сроком на пятнадцать суток с "01" ноября 2026 г. по "15" ноября 2026 г.')
        self.assertNotIn('{', ' '.join(text))


if __name__ == '__main__':
    unittest.main()
