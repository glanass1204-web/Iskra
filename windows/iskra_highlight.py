"""Перенос подсвеченных изменений в один документ по образцу.
Исполнители присылают свои варианты документа (таблица Word: города, километраж, фамилии, численность) и
отмечают изменения зелёным или салатовым: цветом шрифта, маркером выделения или заливкой ячейки. Заливка, которая
есть и в образце (раскраска строк таблицы), изменением не считается. Искра берёт образец, находит в
каждом присланном файле подсвеченное и ставит в образец на то же место:
- таблица — по заголовкам столбцов (порядок столбцов в присланном может быть другим) и по строке: строка ищется
  по неподсвеченным ячейкам (город, фамилия, № п/п), а не по номеру;
- целиком подсвеченная строка, которой нет в образце, — новая строка (перед «Итого», если оно есть);
- подсвеченное вне таблиц — в похожий абзац образца;
- два файла поменяли одно и то же по-разному — [УТОЧНИТЬ: … или …], выделено жёлтым;
- «Итого/Всего» пересчитывается, если в образце оно было суммой столбца.
Нейросеть не нужна: всё переносится дословно, ничего не придумывается. Исходные файлы не меняются."""
import copy
import difflib
import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

GREEN_NAMES = {'green', 'brightgreen', 'darkgreen'}      # «green» в Word — ярко-зелёный (салатовый)
TOTAL = re.compile(r'^\s*(итого|всего)\b', re.IGNORECASE)
NUMBER = re.compile(r'^-?\d[\d\s ]*([.,]\d+)?$')


@dataclass
class Report:
    placed: list = field(default_factory=list)      # что перенесено
    added_rows: list = field(default_factory=list)
    conflicts: list = field(default_factory=list)
    unplaced: list = field(default_factory=list)    # подсвечено, но места в образце не нашлось
    totals: list = field(default_factory=list)
    files_without_marks: list = field(default_factory=list)
    unmarked: list = field(default_factory=list)    # отличия без отметки: [{id, file, object, column, old, new}]
    changes: list = field(default_factory=list)     # что поменялось (для доклада и истории)
    days: list = field(default_factory=list)        # пересчитанные «N сут.»
    long_stay: list = field(default_factory=list)   # дольше порога
    snapshot: list = field(default_factory=list)    # состояние таблиц после сбора

    def text(self):
        out = [f'Перенесено изменений: {len(self.placed)}' + (f', новых строк: {len(self.added_rows)}' if self.added_rows else '') + '.']
        if self.conflicts:
            out.append(f'Разные значения в разных файлах ({len(self.conflicts)}) — в документе [УТОЧНИТЬ], выделено жёлтым:')
            out += ['  • ' + c for c in self.conflicts]
        if self.unplaced:
            out.append(f'Не нашла, куда поставить ({len(self.unplaced)}) — перенесите вручную:')
            out += ['  • ' + u for u in self.unplaced]
        if self.totals:
            out.append('Пересчитано «Итого»: ' + '; '.join(self.totals))
        if self.long_stay:
            out.append('Дольше порога в районе (выделено жёлтым): ' + '; '.join(self.long_stay))
        if self.days:
            out.append(f'Пересчитано «сут.» на сегодня: {len(self.days)}.')
        if self.unmarked:
            out.append(f'Отличия БЕЗ зелёной отметки ({len(self.unmarked)}) — не перенесены, можно выбрать галочками:')
            out += [f'  • {u["object"]} / {u["column"]}: «{u["old"]}» → «{u["new"]}» ({u["file"]})' for u in self.unmarked[:100]]
        if self.files_without_marks:
            out.append('В этих файлах нет зелёной/салатовой подсветки: ' + ', '.join(self.files_without_marks))
        if self.placed:
            out.append('Что перенесено:')
            out += ['  • ' + p for p in self.placed[:200]]
        return '\n'.join(out)


# ---------- цвет ----------

def _green_hex(value):
    if not value or not re.fullmatch(r'[0-9A-Fa-f]{6}', value):
        return False
    r, g, b = (int(value[i:i + 2], 16) for i in (0, 2, 4))
    return g >= 120 and g >= max(r, b) + 25              # 00FF00, 92D050 (салатовый), C6EFCE, 00B050…


def _shd_green(parent):
    shd = parent.find(qn('w:shd')) if parent is not None else None
    return shd is not None and _green_hex(shd.get(qn('w:fill')))


def _green_font(value):
    """Зелёный цвет шрифта: 00B050 (стандартный зелёный Word), 008000, 00FF00, 92D050, тёмно-зелёный 006100."""
    if not value or not re.fullmatch(r'[0-9A-Fa-f]{6}', value):
        return False
    r, g, b = (int(value[i:i + 2], 16) for i in (0, 2, 4))
    return g >= 90 and g >= max(r, b) + 40


def run_green(r):
    """Фрагмент отмечен как изменение: зелёный/салатовый маркер, зелёная заливка текста или зелёный цвет шрифта."""
    rpr = r.find(qn('w:rPr'))
    if rpr is None:
        return False
    hl = rpr.find(qn('w:highlight'))
    if hl is not None and (hl.get(qn('w:val')) or '').lower() in GREEN_NAMES:
        return True
    color = rpr.find(qn('w:color'))
    if color is not None and _green_font(color.get(qn('w:val'))):
        return True
    return _shd_green(rpr)


def _run_text(r):
    out = []
    for el in r:
        if el.tag == qn('w:t'):
            out.append(el.text or '')
        elif el.tag == qn('w:tab'):
            out.append('\t')
        elif el.tag in (qn('w:br'), qn('w:cr')):
            out.append('\n')
    return ''.join(out)


def segments(p_el, all_green=False):
    """Абзац → [(текст, подсвечен)]; соседние одинаковые склеиваются."""
    out = []
    for r in p_el.iter(qn('w:r')):
        text = _run_text(r)
        if not text:
            continue
        g = all_green or run_green(r)
        if out and out[-1][1] == g:
            out[-1] = (out[-1][0] + text, g)
        else:
            out.append((text, g))
    p_pr = p_el.find(qn('w:pPr'))
    if out and p_pr is not None and _shd_green(p_pr):
        out = [(t, True) for t, _ in out]
    return out


def norm(text):
    text = (text or '').lower().replace('ё', 'е').replace(' ', ' ')
    return re.sub(r'\s+', ' ', text).strip(' .,;:-–—')


# ---------- таблицы ----------

def _fill(parent):
    shd = parent.find(qn('w:shd')) if parent is not None else None
    return (shd.get(qn('w:fill')) or '').upper() if shd is not None else ''


class Cell:
    def __init__(self, cell):
        self.cell = cell
        self.fill = _fill(cell._tc.find(qn('w:tcPr')))
        self.marked = [segments(p._p) for p in cell.paragraphs]          # отметки самого текста
        self.resolve()

    def resolve(self, template_fills=()):
        """Заливка ячейки — изменение, только если она зелёная и такой заливки нет в образце."""
        shaded = _green_hex(self.fill) and self.fill not in template_fills
        self.paragraphs = [[(t, g or shaded) for t, g in segs] for segs in self.marked]
        self.text = '\n'.join(''.join(t for t, _ in segs) for segs in self.paragraphs).strip()
        marked = ''.join(t for segs in self.paragraphs for t, g in segs if g)
        self.green = bool(marked.strip()) or (shaded and not self.text)
        self.all_green = self.green and not ''.join(t for segs in self.paragraphs for t, g in segs if not g).strip()


def fills(table):
    """Все заливки ячеек таблицы образца (раскраска строк — не изменение)."""
    return {f for f in (_fill(tc.find(qn('w:tcPr'))) for tc in table._tbl.iter(qn('w:tc'))) if f}


def is_section(row):
    """Строка-раздел: одна ячейка на всю ширину («СЕВЕРОДОНЕЦКОЕ НАПРАВЛЕНИЕ»)."""
    return len(row) > 1 and all(c is row[0] for c in row) and bool(norm(row[0].text))


def is_empty(row):
    return all(not norm(c.text) for c in row)


def section_of(g, i):
    for j in range(i, g.head - 1, -1):
        if is_section(g[j]):
            return norm(g[j][0].text)
    return ''


class Grid(list):
    """Таблица по позициям: объединённая ячейка повторяется во всех своих столбцах (столбцы совпадают по строкам).
    head — сколько строк занимает шапка (двухуровневые заголовки, строка «1 2 3 …»), headers — заголовок каждого столбца."""
    head = 1
    headers = ()


def grid(table):
    cache = {}
    g = Grid()
    for row in table.rows:
        cells = []
        for c in row.cells:
            key = id(c._tc)
            if key not in cache:
                cache[key] = Cell(c)
            cells.append(cache[key])
        g.append(cells)
    # Шапка: первая строка; следующая — тоже шапка, если в предыдущей есть объединённые по горизонтали ячейки
    # («Расстояние, км» над «туда | обратно») или ячейка продолжается вниз (объединение по вертикали).
    head = 1
    while head < min(len(g), 4):
        prev, row = g[head - 1], g[head]
        spans = any(prev[k] is prev[k + 1] for k in range(len(prev) - 1))
        down = any(c is d for c, d in zip(row, prev))
        if not (spans or down):
            break
        head += 1
    if head < len(g) and _is_numbering(g[head]):
        head += 1                                          # строка «1 2 3 4 5» под шапкой
    g.head = head
    width = max((len(r) for r in g), default=0)
    headers = []
    for k in range(width):
        parts = []
        for row in g[:g.head]:
            if k < len(row) and row[k].text and not _is_numbering(row) and row[k].text not in parts:
                parts.append(row[k].text)
        headers.append(' '.join(parts))
    g.headers = headers
    return g


def _similar(a, b):
    a, b = norm(a), norm(b)
    if not a or not b:
        return 0.0
    return 1.0 if a == b else difflib.SequenceMatcher(None, a, b).ratio()


def column_map(src_head, dst_head):
    """Столбец присланной таблицы → столбец образца (по заголовкам; порядок может отличаться)."""
    mapping, used = {}, set()
    for j, h in enumerate(src_head):
        best, score = None, 0.0
        for k, d in enumerate(dst_head):
            if k in used:
                continue
            s = _similar(h, d)
            if s > score:
                best, score = k, s
        if best is not None and score >= 0.6:
            mapping[j] = best
            used.add(best)
    if len(mapping) < len(src_head) and len(src_head) == len(dst_head):
        for j in range(len(src_head)):                    # заголовки не совпали — по порядку, если столбцов столько же
            if j not in mapping and j not in used:
                mapping[j] = j
                used.add(j)
    return mapping


def table_score(src, dst):
    if not src or not dst:
        return 0.0
    head = sum(max((_similar(h, d) for d in dst.headers), default=0) for h in src.headers) / max(1, len(src.headers))
    return head + (0.2 if len(src.headers) == len(dst.headers) else 0)


def _is_numbering(row):
    texts = [norm(c.text) for c in row]
    return len(texts) > 1 and texts == [str(i) for i in range(1, len(texts) + 1)]


def match_row(src_row, mapping, dst_rows, start=None, include_green=False):
    """Строка образца для строки присланной таблицы — по неподсвеченным ячейкам. (индекс, как нашли) или (None, '')."""
    start = getattr(dst_rows, 'head', 1) if start is None else start
    keys, seen = [], set()
    for j, c in enumerate(src_row):
        if j in mapping and (include_green or not c.all_green) and norm(c.text) and id(c) not in seen:
            seen.add(id(c))                                # объединённая ячейка — один ключ, не несколько
            keys.append((j, c))
    if not keys:
        return None, ''
    scores = []
    for i in range(start, len(dst_rows)):
        row = dst_rows[i]
        s = 0.0
        for j, c in keys:
            k = mapping[j]
            if k < len(row):
                sim = _similar(c.text, row[k].text)
                s += 1 if sim == 1 else (0.5 if sim >= 0.85 else 0)
        scores.append((s, i))
    scores.sort(key=lambda x: -x[0])
    if not scores or scores[0][0] < max(1, len(keys) / 2):
        return None, ''                                    # совпал один № п/п, а город и фамилия другие — не та строка
    if len(scores) > 1 and scores[1][0] == scores[0][0]:
        return None, 'неоднозначно'
    return scores[0][1], ''


# ---------- запись ----------

def _base_rpr(p_el):
    for r in p_el.iter(qn('w:r')):
        rpr = r.find(qn('w:rPr'))
        return copy.deepcopy(rpr) if rpr is not None else None
    return None


def _clean_rpr(rpr):
    if rpr is None:
        return None
    for tag in ('w:highlight', 'w:shd'):
        for el in rpr.findall(qn(tag)):
            rpr.remove(el)
    for el in rpr.findall(qn('w:color')):
        if _green_font(el.get(qn('w:val'))):
            rpr.remove(el)                                 # старая зелёная отметка не переходит на новый текст
    return rpr


def clear_marks(doc):
    """Снять прошлые отметки (зелёный/салатовый шрифт и маркер) в документе, куда вносятся изменения, — видно
    будет только новое. Заливку ячеек не трогаем: это может быть раскраска таблицы."""
    for r in doc.element.body.iter(qn('w:r')):
        rpr = r.find(qn('w:rPr'))
        if rpr is None:
            continue
        for el in rpr.findall(qn('w:highlight')):
            if (el.get(qn('w:val')) or '').lower() in GREEN_NAMES:
                rpr.remove(el)
        for el in rpr.findall(qn('w:color')):
            if _green_font(el.get(qn('w:val'))):
                rpr.remove(el)


def _style_font(style):
    """(имя шрифта, размер в полупунктах) по цепочке стилей."""
    name = size = None
    seen = set()
    while style is not None and (name is None or size is None) and id(style.element) not in seen:
        seen.add(id(style.element))                    # стили могут ссылаться друг на друга по кругу
        rpr = style.element.find(qn('w:rPr'))
        if rpr is not None:
            f = rpr.find(qn('w:rFonts'))
            if name is None and f is not None:
                name = f.get(qn('w:ascii')) or f.get(qn('w:hAnsi'))
            z = rpr.find(qn('w:sz'))
            if size is None and z is not None:
                size = z.get(qn('w:val'))
        style = style.base_style
    return name, size


def _doc_default_font(doc):
    name, size = _style_font(doc.styles['Normal'])
    defaults = doc.styles.element.find(qn('w:docDefaults'))
    rpr = defaults.find(qn('w:rPrDefault') + '/' + qn('w:rPr')) if defaults is not None else None
    if rpr is not None:
        f, z = rpr.find(qn('w:rFonts')), rpr.find(qn('w:sz'))
        name = name or (f.get(qn('w:ascii')) or f.get(qn('w:hAnsi')) if f is not None else None)
        size = size or (z.get(qn('w:val')) if z is not None else None)
    return name or 'Times New Roman', size or '24'


def _para_style(p_el, doc):
    """Стиль абзаца шаблона по его коду (у абзаца из ячейки нет «родителя», через python-docx его не узнать)."""
    ps = p_el.find(qn('w:pPr') + '/' + qn('w:pStyle')) if p_el is not None else None
    sid = ps.get(qn('w:val')) if ps is not None else None
    try:
        if sid:
            for st in doc.styles:
                if st.style_id == sid:
                    return st
        return doc.styles['Normal']
    except (KeyError, ValueError, AttributeError):
        return None


def _exact_rpr(r_el, p_el, doc):
    """Свойства фрагмента шаблона с ЯВНО прописанными шрифтом и размером (а не «как в стиле»): в другом документе
    стиль «Обычный» может быть другим — шрифт должен совпасть с шаблоном буквально."""
    rpr = r_el.find(qn('w:rPr')) if r_el is not None else None
    if rpr is None and p_el is not None:
        mark = p_el.find(qn('w:pPr') + '/' + qn('w:rPr'))       # пустая ячейка: шрифт знака абзаца
        rpr = mark
    rpr = _clean_rpr(copy.deepcopy(rpr)) if rpr is not None else OxmlElement('w:rPr')
    for el in rpr.findall(qn('w:color')):
        rpr.remove(el)
    for tag in ('w:ins', 'w:del', 'w:rPrChange'):
        for el in rpr.findall(qn(tag)):
            rpr.remove(el)
    f, z = rpr.find(qn('w:rFonts')), rpr.find(qn('w:sz'))
    name = (f.get(qn('w:ascii')) or f.get(qn('w:hAnsi'))) if f is not None else None
    size = z.get(qn('w:val')) if z is not None else None
    if (name is None or size is None) and p_el is not None:
        sn, ss = _style_font(_para_style(p_el, doc))
        dn, ds = _doc_default_font(doc)
        name, size = name or sn or dn, size or ss or ds
    if f is None:
        f = OxmlElement('w:rFonts'); rpr.insert(0, f)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        f.set(qn(attr), name or 'Times New Roman')
    for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        f.attrib.pop(qn(attr), None)
    for tag in ('w:sz', 'w:szCs'):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); rpr.append(el)
        el.set(qn('w:val'), str(size or '24'))
    return rpr


def _restyle(p_el, src_ppr, rpr):
    """Абзац документа: выравнивание/отступы и шрифт — как у шаблона; текст и цветные отметки остаются."""
    if src_ppr is not None:
        old = p_el.find(qn('w:pPr'))
        if old is not None:
            p_el.remove(old)
        p_el.insert(0, copy.deepcopy(src_ppr))
    for r in p_el.iter(qn('w:r')):
        keep = []
        old = r.find(qn('w:rPr'))
        if old is not None:
            keep = [copy.deepcopy(el) for tag in ('w:color', 'w:highlight') for el in old.findall(qn(tag))]
            r.remove(old)
        new = copy.deepcopy(rpr)
        for el in keep:
            new.append(el)
        r.insert(0, new)


def _look(p_el, doc):
    """(pPr, rPr) абзаца шаблона: по первому фрагменту с текстом."""
    if p_el is None:
        return None, _exact_rpr(None, None, doc)
    first = next((r for r in p_el.iter(qn('w:r')) if _run_text(r).strip()), None)
    if first is None:                                    # пустой фрагмент со своим шрифтом
        first = next((r for r in p_el.iter(qn('w:r')) if r.find(qn('w:rPr')) is not None), None)
    return p_el.find(qn('w:pPr')), _exact_rpr(first, p_el, doc)


def _has_own_font(p_el):
    return p_el is not None and (any(r.find(qn('w:rPr')) is not None for r in p_el.iter(qn('w:r')))
                                 or p_el.find(qn('w:pPr') + '/' + qn('w:rPr')) is not None)


def _copy_look(dst_cell, src_cell, doc, fallback=None):
    """Вид ячейки шаблона (шрифт, размер, жирность, выравнивание, отступы) — на ячейку документа."""
    src_p = src_cell.paragraphs[0]._p if src_cell.paragraphs else None
    ppr, rpr = _look(src_p, doc)
    if fallback is not None and not _has_own_font(src_p):
        rpr = fallback                                     # пустая ячейка шаблона без своего шрифта — шрифт таблицы
    for p in dst_cell.paragraphs:
        _restyle(p._p, ppr, rpr)


def _table_font(table, doc):
    """Самый частый шрифт данных в таблице шаблона — для пустых ячеек без своего оформления."""
    from collections import Counter
    seen = Counter()
    keep = {}
    runs = [r for r in table._tbl.iter(qn('w:r')) if _run_text(r).strip()] or \
        [r for r in table._tbl.iter(qn('w:r')) if r.find(qn('w:rPr')) is not None]
    for r in runs:
        rpr = _exact_rpr(r, r.getparent(), doc)
        f, z = rpr.find(qn('w:rFonts')), rpr.find(qn('w:sz'))
        key = (f.get(qn('w:ascii')), z.get(qn('w:val')))
        seen[key] += 1
        keep.setdefault(key, rpr)
    if not seen:
        return _exact_rpr(None, None, doc)
    rpr = copy.deepcopy(keep[seen.most_common(1)[0][0]])
    for tag in ('w:b', 'w:bCs', 'w:i', 'w:iCs', 'w:u'):
        for el in rpr.findall(qn(tag)):
            rpr.remove(el)
    return rpr


def apply_format(doc, template):
    """Строго по шаблону: поля и ориентация страницы, шрифт, размер, начертание и выравнивание — в таблицах
    (по заголовкам столбцов) и в абзацах вне таблиц. Шрифт прописывается явно, а не «как в стиле»."""
    tpl = Document(str(template))
    for d, t in zip(doc.sections, tpl.sections):
        for a in ('orientation', 'page_width', 'page_height', 'left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
            setattr(d, a, getattr(t, a))
    from docx.shared import Pt
    name, size = _doc_default_font(tpl)                  # пустые строки и новый текст — тоже шрифтом шаблона
    doc.styles['Normal'].font.name, doc.styles['Normal'].font.size = name, Pt(int(size) / 2)
    tpl_grids = [grid(t) for t in tpl.tables]
    tpl_fonts = [_table_font(t, tpl) for t in tpl.tables]
    for table in doc.tables:
        g = grid(table)
        if not g or not tpl_grids:
            continue
        best = max(range(len(tpl_grids)), key=lambda i: table_score(g, tpl_grids[i]))
        t = tpl_grids[best]
        if table_score(g, t) < 0.3:
            continue
        mapping = column_map(g.headers, t.headers)
        sections = [i for i in range(t.head, len(t)) if is_section(t[i])]
        data = [i for i in range(t.head, len(t)) if not is_section(t[i]) and not TOTAL.match(t[i][0].text or '')]
        totals = [i for i in range(t.head, len(t)) if TOTAL.match(t[i][0].text or '')]
        done = set()
        for r, row in enumerate(g):
            if r < g.head:
                src_r = min(r, t.head - 1)
            elif is_section(row):
                src_r = sections[0] if sections else None
            elif TOTAL.match(row[0].text or ''):
                src_r = totals[0] if totals else (data[-1] if data else None)
            else:
                src_r = data[min(r - g.head, len(data) - 1)] if data else None
            for j, c in enumerate(row):
                if id(c) in done:
                    continue
                done.add(id(c))
                k = mapping.get(j, j)
                if src_r is not None and k < len(t[src_r]):
                    _copy_look(c.cell, t[src_r][k].cell, tpl, fallback=tpl_fonts[best])
                else:
                    for p in c.cell.paragraphs:
                        _restyle(p._p, None, tpl_fonts[best])
    # Абзацы вне таблиц: такой же по тексту абзац шаблона, иначе тот же по счёту, иначе основной шрифт шаблона.
    tpl_pars = [p for p in tpl.paragraphs if p.text.strip()]
    body = _look(None, tpl)[1]
    if tpl_pars:
        from collections import Counter
        common = Counter((p.text.strip() and _look(p._p, tpl)[1].find(qn('w:rFonts')).get(qn('w:ascii')),
                          _look(p._p, tpl)[1].find(qn('w:sz')).get(qn('w:val'))) for p in tpl_pars).most_common(1)[0][0]
        body = next(_look(p._p, tpl)[1] for p in tpl_pars
                    if (_look(p._p, tpl)[1].find(qn('w:rFonts')).get(qn('w:ascii')),
                        _look(p._p, tpl)[1].find(qn('w:sz')).get(qn('w:val'))) == common)
        for el in [x for tag in ('w:b', 'w:bCs', 'w:i', 'w:iCs', 'w:u') for x in body.findall(qn(tag))]:
            body.remove(el)
    doc_pars = [p for p in doc.paragraphs if p.text.strip()]
    for n, p in enumerate(doc_pars):
        same = max(tpl_pars, key=lambda q: difflib.SequenceMatcher(None, norm(p.text), norm(q.text)).ratio(), default=None)
        if same is not None and difflib.SequenceMatcher(None, norm(p.text), norm(same.text)).ratio() >= 0.6:
            src = same
        elif len(doc_pars) == len(tpl_pars):
            src = tpl_pars[n]
        else:
            src = None
        if src is not None:
            ppr, rpr = _look(src._p, tpl)
            _restyle(p._p, ppr, rpr)
        else:
            _restyle(p._p, None, body)


FONT = 'font'                                            # mark=FONT — изменения зелёным шрифтом, а не маркером
GREEN_RGB = '00B050'


def _set_segments(p, segs, mark, color=WD_COLOR_INDEX.BRIGHT_GREEN):
    """Заменить текст абзаца, сохранив оформление его первого фрагмента; подсвеченное — выделить
    (mark=FONT — зелёным шрифтом; конфликты и «дольше порога» — всегда маркером своего цвета)."""
    base = _clean_rpr(_base_rpr(p._p))
    for r in list(p._p.iter(qn('w:r'))):
        r.getparent().remove(r)
    for text, green in segs:
        run = p.add_run(text)
        if base is not None:
            old = run._r.find(qn('w:rPr'))
            if old is not None:
                run._r.remove(old)
            run._r.insert(0, copy.deepcopy(base))
        if mark and green:
            if mark == FONT and color == WD_COLOR_INDEX.BRIGHT_GREEN:
                from docx.shared import RGBColor
                run.font.color.rgb = RGBColor.from_string(GREEN_RGB)
            else:
                run.font.highlight_color = color


def set_cell(cell, paragraphs, mark, color=WD_COLOR_INDEX.BRIGHT_GREEN):
    """paragraphs: [[(текст, подсвечен)]] — содержимое ячейки."""
    ps = cell.paragraphs
    for extra in ps[1:]:
        extra._p.getparent().remove(extra._p)
    first = ps[0]
    _set_segments(first, paragraphs[0] if paragraphs else [], mark, color)
    for segs in paragraphs[1:]:
        p = cell.add_paragraph()
        p_pr = first._p.find(qn('w:pPr'))
        if p_pr is not None:
            p._p.insert(0, copy.deepcopy(p_pr))
        _set_segments(p, segs, mark, color)


def _number(text):
    t = norm(text).replace(' ', '').replace(',', '.')
    if not t or not NUMBER.match(norm(text)):
        return None
    try:
        return float(t)
    except ValueError:
        return None


def _fmt(n):
    return str(int(n)) if abs(n - round(n)) < 1e-9 else f'{n:.2f}'.rstrip('0').rstrip('.').replace('.', ',')


# ---------- главное ----------

def _row_label(row):
    """Чем строка называется: первая непустая ячейка, не номер («КП 5447 обр СпН»)."""
    return next((x.text for x in row if norm(x.text) and not re.match(r'^\d+\.?$', norm(x.text))), '')


def _all_green(c):
    c.paragraphs = [[(t, True) for t, _ in segs] for segs in c.paragraphs]
    c.green = True


def merge(template, sources, output, mark=True, accept=(), days=True, threshold=None, today=None, base=None,
          take_all=False):
    """template — шаблон .docx (как должно выглядеть); base — документ, в который вносить изменения (нет —
    вносим в сам шаблон); sources — присланные .docx; output — куда сохранить. Возвращает Report.
    mark — выделить перенесённое (FONT — зелёным шрифтом). accept — отличия без отметки, которые человек выбрал
    перенести (id из report.unmarked); take_all — переносить все отличия, а не только отмеченные.
    days — пересчитать «N сут. с ДАТА» на сегодня; threshold — дольше скольких суток выделить."""
    report = Report()
    accept = set(accept)
    doc = Document(str(base or template))
    if base or mark == FONT:
        clear_marks(doc)                                 # видно только новые изменения
    dst_tables = [grid(t) for t in doc.tables]
    dst_fills = [fills(t) for t in doc.tables]
    totals_before = [_totals(g) for g in dst_tables]
    # цель → {значение: [ячейка-источник, [файлы]]}
    cell_values, cell_info = {}, {}
    par_values = {}
    new_rows = {}                                        # индекс таблицы → [(ячейки по столбцам образца, файл, раздел)]
    dst_pars = [p for p in doc.paragraphs]

    def put(best, dst, target, col, c, name):
        key = (best, target, col)
        cell_values.setdefault(key, {}).setdefault(norm(c.text), [c, []])[1].append(name)
        cell_info[key] = (_row_label(dst[target]), dst.headers[col])

    def add_row(best, g, r_index, row, mapping, name):
        cells = {}
        for j, c in enumerate(row):
            if j in mapping and c not in cells.values():
                cells[mapping[j]] = c
        signature = tuple(norm(c.text) for c in cells.values())
        if signature not in {tuple(norm(c.text) for c in x[0].values()) for x in new_rows.get(best, [])}:
            new_rows.setdefault(best, []).append((cells, name, section_of(g, r_index)))

    for path in sources:
        name = Path(path).name
        src = Document(str(path))
        found_marks = False
        for t_index, table in enumerate(src.tables):
            g = grid(table)
            if not g:
                continue
            best = max(range(len(dst_tables)), key=lambda i: table_score(g, dst_tables[i]), default=None)
            if best is not None:
                for row in g:
                    for c in row:
                        c.resolve(dst_fills[best])            # раскраска строк, как в образце, — не изменение
            marked_table = any(c.green for row in g for c in row)
            found_marks = found_marks or marked_table
            if best is None or table_score(g, dst_tables[best]) < 0.3:
                if marked_table:
                    for row in g:
                        if any(c.green for c in row):
                            report.unplaced.append(f'{name}: строка «{" | ".join(c.text for c in row)}» — в образце нет такой таблицы')
                continue
            dst = dst_tables[best]
            mapping = column_map(g.headers, dst.headers)
            for r_index, row in enumerate(g[g.head:], g.head):
                if is_section(row) or is_empty(row):
                    continue
                marked = any(c.green for c in row)
                target, why = match_row(row, mapping, dst)
                if target is None and marked and all(c.all_green or not norm(c.text) for c in row):
                    # Вся строка зелёная: это либо новая строка, либо старая, целиком отмеченная, — сверяем по всем ячейкам.
                    target, why = match_row(row, mapping, dst, include_green=True)
                    if target is None:
                        add_row(best, g, r_index, row, mapping, name)
                        continue
                if target is None:
                    if marked:
                        report.unplaced.append(f'{name}: строка «{" | ".join(c.text for c in row)}» — '
                                               + ('подходит к нескольким строкам образца' if why else 'нет такой строки в образце'))
                    elif sum(1 for c in set(row) if norm(c.text)) >= 2:
                        uid = (str(path), t_index, r_index, 'row')
                        if uid in accept or take_all:
                            for c in set(row):
                                _all_green(c)              # без отметки, но берём — в итоге видно как изменение
                            add_row(best, g, r_index, row, mapping, name)
                        else:
                            report.unmarked.append({'id': uid, 'file': name, 'object': _row_label(row),
                                                    'column': 'новая строка', 'old': '',
                                                    'new': ' | '.join(dict.fromkeys(c.text for c in row if c.text))})
                    continue
                done = set()
                for j, c in enumerate(row):
                    if id(c) in done or j not in mapping:
                        if c.green and j not in mapping and id(c) not in done:
                            report.unplaced.append(f'{name}: «{c.text}» — нет столбца «{g.headers[j]}» в образце')
                        done.add(id(c))
                        continue
                    done.add(id(c))
                    goal = dst[target][mapping[j]]
                    # Зелёная/салатовая заливка — отметка, если у этой же ячейки образца её нет (в образце прошлой
                    # недели зелёные ячейки есть, но это не значит, что новая заливка — «раскраска таблицы»).
                    if not c.green and _green_hex(c.fill) and c.fill != goal.fill and norm(c.text):
                        _all_green(c)
                    if norm(c.text) == norm(goal.text) and not c.green:
                        continue                               # значение как в образце и не отмечено — менять нечего
                    col = next(k for k, x in enumerate(dst[target]) if x is goal)   # объединённая ячейка — один адрес
                    uid = (str(path), t_index, r_index, j)
                    if c.green or uid in accept or take_all:   # отмечено — переносится, даже если текст тот же («Авдеевка»)
                        if not c.green:
                            _all_green(c)                      # взято без отметки — в итоге тоже зелёным
                        put(best, dst, target, col, c, name)
                    elif norm(c.text):                         # пустая ячейка у исполнителя — «не заполнил», не изменение
                        report.unmarked.append({'id': uid, 'file': name, 'object': _row_label(dst[target]),
                                                'column': dst.headers[col], 'old': goal.text, 'new': c.text})
        for p in src.paragraphs:
            segs = segments(p._p)
            if not any(g for _, g in segs):
                continue
            found_marks = True
            text = ''.join(t for t, _ in segs)
            best, score = None, 0.0
            for i, d in enumerate(dst_pars):
                s = difflib.SequenceMatcher(None, norm(text), norm(d.text)).ratio() if d.text.strip() else 0
                if s > score:
                    best, score = i, s
            if best is None or score < 0.5:
                report.unplaced.append(f'{name}: абзац «{text[:120]}» — нет похожего абзаца в образце')
                continue
            par_values.setdefault(best, {}).setdefault(norm(text), [segs, []])[1].append(name)
        if not found_marks:
            report.files_without_marks.append(name)

    moved = set()                                        # строки, где сменилось положение, — сутки заново
    for (t, r, c), variants in cell_values.items():
        cell = dst_tables[t][r][c].cell
        label, header = cell_info[(t, r, c)]
        old = dst_tables[t][r][c].text
        where = f'{label} / {header}'
        if len(variants) == 1:
            src_cell, files = next(iter(variants.values()))
            set_cell(cell, src_cell.paragraphs, mark)
            same = ' — то же значение, отмечено как изменение' if norm(old) == norm(src_cell.text) else ''
            report.placed.append(f'{where}: «{src_cell.text}» ({", ".join(files)}){same}')
            report.changes.append({'object': label, 'column': header, 'old': old, 'new': src_cell.text, 'files': files})
            if POSITION.search(header or '') and norm(old) != norm(src_cell.text):
                moved.add((t, r))
        else:
            text = '[УТОЧНИТЬ: ' + ' или '.join(f'{v[0].text} ({", ".join(v[1])})' for v in variants.values()) + ']'
            set_cell(cell, [[(text, True)]], True, WD_COLOR_INDEX.YELLOW)
            report.conflicts.append(f'{where}: {text[11:-1]}')
    for t, r in sorted(moved):
        _restart_days(dst_tables[t], r, {c for (tt, rr, c) in cell_values if (tt, rr) == (t, r)}, mark, today, report)
    for i, variants in par_values.items():
        p = dst_pars[i]
        if len(variants) == 1:
            segs, files = next(iter(variants.values()))
            _set_segments(p, segs, mark)
            report.placed.append(f'абзац «{"".join(t for t, _ in segs)[:80]}» ({", ".join(files)})')
        else:
            text = '[УТОЧНИТЬ: ' + ' или '.join(f'«{"".join(t for t, _ in v[0])[:100]}» ({", ".join(v[1])})' for v in variants.values()) + ']'
            _set_segments(p, [(text, True)], True, WD_COLOR_INDEX.YELLOW)
            report.conflicts.append('абзац: ' + text[11:-1])
    for t, rows in new_rows.items():
        _add_rows(doc.tables[t], dst_tables[t], rows, mark, report)
    for t, table in enumerate(doc.tables):
        if t < len(totals_before) and totals_before[t]:
            _recount(grid(table), totals_before[t], report, mark)
    if days:
        update_days(doc, today=today, threshold=threshold, report=report)
    try:
        apply_format(doc, template)                      # строго как в шаблоне: шрифт везде одинаковый
    except Exception as exc:                             # noqa: BLE001 — изменения важнее: документ всё равно сохраняем
        report.unplaced.append(f'Оформление по шаблону применено не полностью ({type(exc).__name__}) — '
                               'изменения перенесены, проверьте шрифт. Пришлите «Сообщить об ошибке».')
        try:
            import sys
            import iskra_bugreport
            iskra_bugreport.log_crash('Оформление по шаблону', *sys.exc_info())
        except Exception:                                # noqa: BLE001
            pass
    report.snapshot = snapshot(doc)
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    return report


def _add_rows(table, dst, rows, mark, report):
    """Новые строки: в первую пустую строку своего раздела («СЕВЕРОДОНЕЦКОЕ НАПРАВЛЕНИЕ»); пустой нет — копия
    последней строки данных раздела (то же оформление), перед «Итого»."""
    from docx.table import _Row
    used = set()
    for cells, name, section in rows:
        start, stop = dst.head, len(dst)
        if section:
            at = next((i for i in range(dst.head, len(dst)) if is_section(dst[i]) and norm(dst[i][0].text) == section), None)
            if at is not None:
                start = at + 1
                stop = next((i for i in range(start, len(dst)) if is_section(dst[i])), len(dst))
        empty = next((i for i in range(start, stop) if is_empty(dst[i]) and i not in used), None)
        if empty is not None:
            used.add(empty)
            target_cells, seen = [], set()
            for k, c in enumerate(dst[empty]):
                if id(c) not in seen:
                    seen.add(id(c))
                    target_cells.append((k, c.cell))
        else:
            last = next((i for i in range(stop - 1, start - 1, -1) if dst[i] and not is_empty(dst[i])
                         and not TOTAL.match(_label(dst[i])) and not is_section(dst[i])), None)
            if last is None:
                report.unplaced.append(f'{name}: новая строка «{" | ".join(c.text for c in cells.values())}» — '
                                       'в таблице образца нет строк-образцов')
                continue
            tr = copy.deepcopy(table.rows[last]._tr)
            table.rows[last]._tr.addnext(tr)
            target_cells, seen = [], set()
            for k, cell in enumerate(_Row(tr, table).cells):
                if id(cell._tc) not in seen:
                    seen.add(id(cell._tc))
                    target_cells.append((k, cell))
        for k, cell in target_cells:
            src = cells.get(k)
            set_cell(cell, src.paragraphs if src else [[('', False)]], mark)
        where = f' (раздел «{section.upper()}»)' if section else ''
        report.added_rows.append(f'{" | ".join(c.text for c in cells.values() if c.text)}{where} ({name})')
        report.changes.append({'object': next((c.text for c in cells.values() if norm(c.text)), ''), 'new_row': True,
                               'values': {dst.headers[k]: c.text for k, c in cells.items() if c.text and k < len(dst.headers)},
                               'section': section, 'files': [name]})


def _label(row):
    return next((c.text for c in row if norm(c.text)), '')


def _totals(g):
    """Строки «Итого» и столбцы, где итог был суммой строк выше (только такие пересчитываются)."""
    out = []
    data_start = g.head
    for i, row in enumerate(g):
        if i < g.head or not row or not TOTAL.match(_label(row)):
            continue
        body = [r for r in g[data_start:i] if not _is_numbering(r)]
        cols = []
        for k in range(len(row)):
            if row[k] is row[0] or (k and row[k] is row[k - 1]):
                continue                                   # подпись «Итого» и объединённые ячейки
            total = _number(row[k].text)
            values = [_number(r[k].text) if k < len(r) else None for r in body]
            nums = [v for v in values if v is not None]
            if total is None or not nums or len(nums) < len([r for r in body if k < len(r) and norm(r[k].text)]):
                continue
            if abs(sum(nums) - total) < 0.01:
                cols.append(k)
        if cols:
            out.append((i, cols, len(body)))
        data_start = i + 1
    return out


def _recount(g, before, report, mark):
    """Пересчитать «Итого» по тем столбцам, где оно было суммой. Итоговые строки сопоставляются по порядку."""
    total_rows = [i for i, row in enumerate(g) if i >= g.head and row and TOTAL.match(_label(row))]
    start = g.head
    for i, (_, cols, _) in zip(total_rows, before):
        row = g[i]
        body = [r for r in g[start:i] if not _is_numbering(r)]
        start = i + 1
        for k in cols:
            pairs = [(r, _number(r[k].text)) for r in body if k < len(r)]
            if any(n is None and norm(r[k].text) for r, n in pairs):
                continue                                   # в столбце появился не-число — не трогаем
            total = sum(n for _, n in pairs if n is not None)
            old = _number(row[k].text)
            if old is not None and abs(old - total) > 0.001:
                set_cell(row[k].cell, [[(_fmt(total), True)]], mark)
                report.totals.append(f'{g.headers[k]}: {_fmt(old)} → {_fmt(total)}')


# ---------- «N сут. с ДАТА» ----------

DAYS = re.compile(r'(\d+)(\s*сут(?:ок|ки|\.)?\s*[сc]\s*)(\d{1,2})\.(\d{1,2})\.(\d{2,4})', re.IGNORECASE)


POSITION = re.compile(r'положени|местонахожд|район\s+(?:располож|сосредот)|\bн\.?\s?п\.?\b|населенн|населённ', re.IGNORECASE)
STAY = re.compile(r'время\s+нахожд|сут(?:ок|ки)?\s+в\s+район|нахождени\w*\s+в\s+район', re.IGNORECASE)


def _restart_days(g, r, placed_cols, mark, today, report):
    """Положение сменилось — «Время нахождения в районе» начинается заново: «1 сут. с <сегодня>», как изменение.
    Если подразделение само прислало новые сутки — их и оставляем."""
    from datetime import date
    today = today or date.today()
    row = g[r]
    col = next((k for k, h in enumerate(g.headers) if STAY.search(h or '')), None)
    if col is None:
        col = next((k for k, c in enumerate(row) if DAYS.search(c.text or '')), None)
    if col is None or col in placed_cols or any(row[col] is row[k] for k in placed_cols):
        return
    cell = row[col]
    old = cell.text
    m = DAYS.search(old or '')
    when = today.strftime('%d.%m.%Y')
    new = (old[:m.start()] + '1' + m.group(2) + when + old[m.end():]) if m else f'1 сут. с {when}'
    set_cell(cell.cell, [[(new, True)]], mark)
    report.placed.append(f'{_row_label(row)} / {g.headers[col]}: «{new}» (положение сменилось — сутки заново)')
    report.changes.append({'object': _row_label(row), 'column': g.headers[col], 'old': old, 'new': new, 'files': []})


def days_since(day, month, year, today):
    """Сутки в районе с даты прибытия включительно: с 19.10 по 22.10 — 4 сут."""
    from datetime import date
    year = int(year)
    year += 2000 if year < 100 else 0
    try:
        start = date(year, int(month), int(day))
    except ValueError:
        return None
    n = (today - start).days + 1
    return n if n >= 1 else None


def _paragraph_days(p, today, threshold):
    """Пересчитать число суток в абзаце. Число правится внутри своего фрагмента — оформление не меняется.
    Возвращает [(было, стало)]."""
    runs = list(p._p.iter(qn('w:r')))
    texts = [_run_text(r) for r in runs]
    full = ''.join(texts)
    out = []
    for m in DAYS.finditer(full):
        n = days_since(m.group(3), m.group(4), m.group(5), today)
        if n is None:
            continue
        start, end = m.span(1)
        pos, run = 0, None
        for r, t in zip(runs, texts):
            if pos <= start and end <= pos + len(t):
                run = r
                break
            pos += len(t)
        if run is None:
            continue                                     # число разорвано на фрагменты — не трогаем
        long = threshold is not None and n > threshold
        if n != int(m.group(1)):
            for t_el in run.iter(qn('w:t')):
                text = t_el.text or ''
                local = start - pos
                if text[local:local + (end - start)] == m.group(1):
                    t_el.text = text[:local] + str(n) + text[local + (end - start):]
                    break
                pos += len(text)
            out.append((m.group(0), f'{n}{m.group(2)}{m.group(3)}.{m.group(4)}.{m.group(5)}', long))
        elif long:
            out.append((m.group(0), m.group(0), long))
        if long:
            rpr = run.find(qn('w:rPr'))
            if rpr is None:
                rpr = OxmlElement('w:rPr'); run.insert(0, rpr)
            for el in rpr.findall(qn('w:highlight')):
                rpr.remove(el)
            hl = OxmlElement('w:highlight'); hl.set(qn('w:val'), 'yellow'); rpr.append(hl)
    return out


def update_days(doc, today=None, threshold=None, report=None):
    """Все «N сут. с ДАТА» в документе — на сегодня; дольше порога — выделить жёлтым. Возвращает число правок."""
    from datetime import date
    today = today or date.today()
    report = report if report is not None else Report()
    changed = 0
    for table in doc.tables:
        g = grid(table)
        seen = set()
        for row in g:
            for c in row:
                if id(c.cell._tc) in seen:
                    continue
                seen.add(id(c.cell._tc))
                for p in c.cell.paragraphs:
                    for old, new, long in _paragraph_days(p, today, threshold):
                        if old != new:
                            changed += 1
                            report.days.append(f'{_row_label(row)}: {old} → {new}')
                        if long:
                            report.long_stay.append(f'{_row_label(row)} — {new}')
    for p in doc.paragraphs:
        for old, new, long in _paragraph_days(p, today, threshold):
            if old != new:
                changed += 1
                report.days.append(f'{old} → {new}')
            if long:
                report.long_stay.append(new)
    return changed


def recount_days_file(path, output=None, today=None, threshold=None):
    """Пересчитать сутки в готовом документе (без переноса изменений). Сохраняет в output (по умолчанию — тот же файл)."""
    doc = Document(str(path))
    report = Report()
    update_days(doc, today=today, threshold=threshold, report=report)
    doc.save(str(output or path))
    return report


# ---------- снимок для истории ----------

def snapshot(doc):
    """Состояние таблиц: [{'headers': [...], 'rows': [{'object', 'section', 'cells': {заголовок: текст}}]}]."""
    out = []
    for table in doc.tables:
        g = grid(table)
        rows = []
        for i in range(g.head, len(g)):
            row = g[i]
            if is_section(row) or is_empty(row) or TOTAL.match(_label(row)):
                continue
            cells, seen = {}, set()
            for k, c in enumerate(row):
                if id(c) in seen or k >= len(g.headers):
                    continue
                seen.add(id(c))
                cells[g.headers[k] or f'столбец {k + 1}'] = c.text
            rows.append({'object': _row_label(row), 'section': section_of(g, i), 'cells': cells})
        out.append({'headers': list(g.headers), 'rows': rows})
    return out
