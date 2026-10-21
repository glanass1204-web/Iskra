"""Контроль исполнения поручений: из входящих документов выписываются поручения (что сделать, срок, кто отвечает),
ведётся общий список с отметками «выполнено», напоминания о сроках, отчёт «что просрочено» в Word и Excel.
Список хранится в data/control.json."""
import json
import os
import re
import threading
import uuid
from datetime import date, datetime, timedelta
from pathlib import Path

import iskra_assistant as assistant
from iskra_telegram import INSTRUCTION, INTRO_PREFIX, parse

STATUSES = ('в работе', 'выполнено')
RESPONSIBLE = re.compile(r'(?:ответственн\w*(?:\s+за\s+[^—:\n]{3,60}?)?\s*(?:[—:-]|назначить|является)\s*'
                         r'|исполнител\w*\s*[—:-]\s*)([^\n;]{2,120})', re.IGNORECASE)
PERSON = re.compile(r'[А-ЯЁ]\.\s?[А-ЯЁ]\.\s?[А-ЯЁ][а-яё]+(?:-[А-ЯЁ][а-яё]+)?|[А-ЯЁ][а-яё]+(?:-[А-ЯЁ][а-яё]+)?\s+[А-ЯЁ]\.\s?[А-ЯЁ]\.')
ADDRESSEE = re.compile(r'^\s*(?:\d{1,2}[.)]\s*)?((?:начальник|командир|руководител|заместител|директор)\w*\s+[^:,.\n]{2,50}?)'
                       r'(?=\s+(?:обеспечить|представить|провести|направить|доложить|организовать|подготовить|выделить|'
                       r'исполнить|проверить|назначить|принять)|[:,])', re.IGNORECASE)
DATE_FMT = '%d.%m.%Y'


def _d(value):
    try:
        return datetime.strptime(value, DATE_FMT).date()
    except (TypeError, ValueError):
        return None


def _short(text, limit=220):
    s = re.sub(r'\s+', ' ', text).strip()
    return s if len(s) <= limit else s[:limit].rsplit(' ', 1)[0] + '…'


def responsible_in(text):
    """«Ответственный — майор С.С. Сидоров» → «майор С.С. Сидоров» (до фамилии включительно)."""
    m = RESPONSIBLE.search(text)
    if not m:
        return ''
    tail = m.group(1)
    person = PERSON.search(tail)
    if person:
        return _short(tail[:person.end()], 80)
    return _short(re.split(r'\.\s', tail)[0].rstrip('.'), 80)


def extract(text, source='', today=None):
    """Поручения из документа по правилам: предложения с «представить / обеспечить / провести…»,
    срок — дата в этом предложении (или единственный срок документа), ответственный — «Ответственный — …»
    в том же пункте, адресат пункта («Начальнику госпиталя обеспечить…») или общий ответственный документа."""
    today = today or date.today()
    t = parse(source, text)
    paragraphs = t.body or [p for p in text.split('\n') if p.strip()]
    body = '\n'.join(paragraphs)
    common = assistant.deadlines(body, today)
    common_deadline = common[0].value if len({c.value for c in common}) == 1 else ''
    number = (t.fields.get('Номер') or [''])[0]
    # Общий ответственный — только если он назван отдельным пунктом, без поручения рядом.
    doc_resp = next((responsible_in(p) for p in paragraphs if RESPONSIBLE.search(p) and not INSTRUCTION.search(
        RESPONSIBLE.sub('', p))), '')
    out, seen = [], set()
    for paragraph in paragraphs:
        para_resp = responsible_in(paragraph)
        a = ADDRESSEE.match(paragraph)
        addressee = _short(a.group(1), 80) if a else ''
        for s in assistant.sentences(paragraph):
            s = INTRO_PREFIX.sub('', s).strip()
            if not INSTRUCTION.search(s) or len(s) < 15 or RESPONSIBLE.match(s):
                continue
            key = s.lower()[:120]
            if key in seen:
                continue
            seen.add(key)
            dl = assistant.deadlines(s, today)
            out.append({'text': _short(re.sub(r'^\d{1,2}[.)]\s*', '', s)),
                        'deadline': dl[0].value if dl else common_deadline,
                        'responsible': responsible_in(s) or para_resp or addressee or doc_resp,
                        'source': source, 'number': number})
    return out


def merge_ai(rule_tasks, ai_tasks, text):
    """Поручения от нейросети дополняют правила. Срок принимается, только если такая дата есть в документе."""
    from iskra_ai import supported
    flat = ' '.join(text.split()).lower().replace('ё', 'е')
    dates = {d.strftime(DATE_FMT) for _, _, d in assistant.parse_dates(text)}
    have = {t['text'].lower()[:60] for t in rule_tasks}
    out = list(rule_tasks)
    for t in ai_tasks:
        what = _short(str(t.get('что') or t.get('text') or ''))
        if not what or what.lower()[:60] in have:
            continue
        dl = str(t.get('срок') or '').strip()
        who = str(t.get('кто') or '').strip()
        out.append({'text': what, 'deadline': dl if dl in dates else '',
                    'responsible': who if who and supported(who, flat) else '', 'ai': True})
    return out


class Control:
    def __init__(self, root):
        self.path = Path(root) / 'control.json'
        self._lock = threading.Lock()
        try:
            self.items = json.loads(self.path.read_text(encoding='utf-8'))
        except (OSError, ValueError):
            self.items = []

    def _save(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        tmp = self.path.with_suffix('.tmp')
        tmp.write_text(json.dumps(self.items, ensure_ascii=False, indent=1), encoding='utf-8')
        os.replace(tmp, self.path)

    def all(self):
        with self._lock:
            return [dict(i) for i in self.items]

    def add(self, text, deadline='', responsible='', source='', number='', path='', ai=False):
        item = {'id': uuid.uuid4().hex[:10], 'text': text.strip(), 'deadline': deadline.strip(),
                'responsible': responsible.strip(), 'source': source, 'number': number, 'path': path,
                'status': 'в работе', 'created': date.today().strftime(DATE_FMT), 'done': '', 'note': '',
                'ai': bool(ai), 'reminded': ''}
        with self._lock:
            # Тот же пункт из того же документа повторно не добавляется.
            for i in self.items:
                if i['source'] == source and i['text'].lower()[:80] == item['text'].lower()[:80]:
                    return None
            self.items.append(item)
            self._save()
        return item

    def add_from_document(self, text, source='', path='', ai_tasks=None, today=None):
        tasks = extract(text, source, today)
        if ai_tasks:
            tasks = merge_ai(tasks, ai_tasks, text)
        added = []
        for t in tasks:
            item = self.add(t['text'], t.get('deadline', ''), t.get('responsible', ''), source,
                            t.get('number', ''), path, t.get('ai', False))
            if item:
                added.append(item)
        return added

    def update(self, item_id, **fields):
        with self._lock:
            for i in self.items:
                if i['id'] == item_id:
                    i.update(fields)
                    if fields.get('status') == 'выполнено' and not i.get('done'):
                        i['done'] = date.today().strftime(DATE_FMT)
                    if fields.get('status') == 'в работе':
                        i['done'] = ''
            self._save()

    def remove(self, item_id):
        with self._lock:
            self.items = [i for i in self.items if i['id'] != item_id]
            self._save()

    # ---------- выборки ----------

    def open_items(self):
        return [i for i in self.all() if i['status'] != 'выполнено']

    def overdue(self, today=None):
        today = today or date.today()
        return sorted((i for i in self.open_items() if _d(i['deadline']) and _d(i['deadline']) < today),
                      key=lambda i: _d(i['deadline']))

    def due(self, days=3, today=None):
        """Не выполненные со сроком от сегодня до today+days."""
        today = today or date.today()
        return sorted((i for i in self.open_items() if _d(i['deadline']) and today <= _d(i['deadline']) <= today + timedelta(days=days)),
                      key=lambda i: _d(i['deadline']))

    def to_remind(self, today=None, days=3):
        """Что напомнить сейчас: просроченное и со сроком в ближайшие дни; не чаще раза в день на поручение."""
        today = today or date.today()
        stamp = today.strftime(DATE_FMT)
        items = [i for i in self.overdue(today) + self.due(days, today) if i.get('reminded') != stamp]
        with self._lock:
            ids = {i['id'] for i in items}
            for i in self.items:
                if i['id'] in ids:
                    i['reminded'] = stamp
            if items:
                self._save()
        return items

    # ---------- отчёты ----------

    def report_rows(self, only_open=True, today=None):
        today = today or date.today()
        rows = []
        for i in sorted(self.all() if not only_open else self.open_items(),
                        key=lambda x: (_d(x['deadline']) or date.max, x['text'])):
            d = _d(i['deadline'])
            if i['status'] == 'выполнено':
                state = f'выполнено {i["done"]}'
            elif d and d < today:
                state = f'просрочено на {(today - d).days} дн.'
            elif d:
                state = assistant.deadline_status(d, today)
            else:
                state = 'без срока'
            rows.append([i['deadline'] or '—', i['text'], i['responsible'] or '—',
                         (i['source'] + (f' № {i["number"]}' if i['number'] else '')).strip() or '—', state])
        return rows

    def report_xlsx(self, path, only_open=True, today=None):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = 'Контроль исполнения'
        ws.append([f'Контроль исполнения поручений на {(today or date.today()).strftime(DATE_FMT)}'])
        ws['A1'].font = Font(bold=True, size=13)
        ws.append(['Срок', 'Поручение', 'Ответственный', 'Документ', 'Состояние'])
        for c in ws[2]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='24518A')
        red = PatternFill('solid', fgColor='FFCDD2')
        for row in self.report_rows(only_open, today):
            ws.append(row)
            if row[4].startswith('просрочено'):
                for c in ws[ws.max_row]:
                    c.fill = red
            for c in ws[ws.max_row]:
                c.alignment = Alignment(wrap_text=True, vertical='top')
        for col, width in zip('ABCDE', (12, 70, 28, 30, 22)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = 'A3'
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        return path

    def report_docx(self, path, only_open=True, today=None):
        from docx import Document
        from docx.shared import Cm, Pt
        doc = Document()
        sec = doc.sections[0]
        sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(2), Cm(1), Cm(2), Cm(2)
        style = doc.styles['Normal']
        style.font.name = 'Times New Roman'
        style.font.size = Pt(12)
        today = today or date.today()
        rows = self.report_rows(only_open, today)
        late = sum(1 for r in rows if r[4].startswith('просрочено'))
        doc.add_paragraph().add_run(f'СПРАВКА\nо ходе исполнения поручений на {today.strftime(DATE_FMT)}').bold = True
        doc.add_paragraph(f'На контроле: {len(rows)}, из них просрочено: {late}.')
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Table Grid'
        for cell, title in zip(table.rows[0].cells, ('Срок', 'Поручение', 'Ответственный', 'Документ', 'Состояние')):
            cell.text = title
            cell.paragraphs[0].runs[0].bold = True
        for row in rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = value
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        doc.save(path)
        return path
