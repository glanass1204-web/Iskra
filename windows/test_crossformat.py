import tempfile
from pathlib import Path
from docx import Document
from openpyxl import Workbook,load_workbook
from iskra_crossformat import merge_crossformat,MISSING

def test():
 with tempfile.TemporaryDirectory() as t:
  p=Path(t);d=Document();d.add_paragraph('Организация: Альфа');d.add_paragraph('Дата: 21.10.2026');d.add_paragraph('Номер: 42');d.save(p/'a.docx')
  d=Document();d.add_paragraph('Организация: Бета');d.add_paragraph('Номер: 11');d.save(p/'b.docx')
  w=Workbook();s=w.active;s.append(['Организация','Дата','Номер']);s.append(['Гамма','20.10.2026',17]);w.save(p/'c.xlsx')
  src=[p/'a.docx',p/'b.docx',p/'c.xlsx'];heads=['Организация','Дата','Номер']
  x=merge_crossformat(src,heads,p/'out.xlsx');assert x['rows']==3 and x['issues']==1,x
  r=load_workbook(p/'out.xlsx',read_only=True);assert r['Сводная']['C3'].value==MISSING;assert r['Сводная']['B2'].value=='Альфа';assert r['Источники'].max_row==10;r.close()
  y=merge_crossformat(src,heads,p/'out.docx','docx');assert y['rows']==3
  doc=Document(p/'out.docx');assert len(doc.tables[0].rows)==4
  try:merge_crossformat(src,heads,p/'a.docx','docx')
  except ValueError:pass
  else:raise AssertionError('overwriting source permitted')
  print('PASS: Word->Excel, Excel->Word, missing flag, provenance, output verification, overwrite protection')
if __name__=='__main__':test()
