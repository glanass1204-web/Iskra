"""Свод нескольких таблиц Excel в одну по наименованию организации.

Шапка каждой таблицы находится сама (в том числе двухстрочная: «Личный состав» → «по списку» / «на лицо»),
организации сопоставляются по названию («Модуль №1» = «модуль 1»), числа складываются, итог считается формулой.
Порядок строк — по перечню организаций, если он задан. Всё спорное попадает на лист «Проверка».
Исходные файлы не изменяются."""
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAME_HEADER = re.compile(r'наименован|организац|подразделен|учрежден|объект|воинск|в/ч|получател|адресат|структурн', re.IGNORECASE)
TOTAL_ROW = re.compile(r'^\s*(итого|всего|итог)\b', re.IGNORECASE)
ROW_NUMBER = re.compile(r'^(№|n)?\s*(п\s*/?\s*п\.?)?$', re.IGNORECASE)
DASH = {'-', '–', '—', 'х', 'x', 'Х'}

BLUE = PatternFill('solid', fgColor='24518A')
YELLOW = PatternFill('solid', fgColor='FFF2CC')
ORANGE = PatternFill('solid', fgColor='FCE4D6')
GRAY = PatternFill('solid', fgColor='E7E6E6')
THIN = Side(style='thin', color='808080')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def fold(text):
    """Название для сравнения: регистр, ё/е, №, кавычки, дефисы и лишние пробелы не важны."""
    s = str(text or '').casefold().replace('ё', 'е').replace('№', ' ')
    s = re.sub(r'(?<=[^\W\d_])-(?=\d)|(?<=\d)-(?=[^\W\d_])', ' ', s)
    s = re.sub(r'[«»"\'“”„()\[\].,;:]', ' ', s)
    return re.sub(r'\s+', ' ', s).strip()


def num(value):
    """Число из ячейки: 12, «12», «1 234», «12,5». Иначе None."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(value) if isinstance(value, float) and value.is_integer() else value
    if isinstance(value, str):
        s = value.strip().replace('\xa0', '').replace(' ', '').replace(',', '.')
        if re.fullmatch(r'-?\d+(\.\d+)?', s):
            f = float(s)
            return int(f) if f.is_integer() else f
    return None


def _empty(value):
    return value is None or (isinstance(value, str) and (not value.strip() or value.strip() in DASH))


def _text(value):
    if isinstance(value, (datetime, date)):
        return value.strftime('%d.%m.%Y')
    return re.sub(r'\s+', ' ', str(value)).strip()


@dataclass
class Table:
    file: str
    sheet: str
    name_header: str
    columns: list                                  # подписи столбцов с данными
    rows: list = field(default_factory=list)       # [(название, {столбец: значение}, номер строки)]
    totals: dict = field(default_factory=dict)     # строка «Итого» из файла: {столбец: число}
    skipped: int = 0                               # строки без данных


def _grid(ws, limit=3000):
    rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, limit), values_only=True)]
    # Объединённые ячейки: значение верхней левой — во все ячейки диапазона.
    for rng in ws.merged_cells.ranges:
        top = rng.min_row - 1
        if top >= len(rows):
            continue
        value = rows[top][rng.min_col - 1] if rng.min_col - 1 < len(rows[top]) else None
        for r in range(rng.min_row - 1, min(rng.max_row, len(rows))):
            for c in range(rng.min_col - 1, rng.max_col):
                if c < len(rows[r]):
                    rows[r][c] = value
    return rows


def _is_numbering(row):
    """Строка «1 2 3 4 5» под шапкой."""
    vals = [num(v) for v in row if not _empty(v)]
    return len(vals) >= 2 and all(isinstance(v, int) for v in vals) and sorted(set(vals)) == list(range(1, len(set(vals)) + 1))


def _find_header(grid):
    for i, row in enumerate(grid[:30]):
        distinct = {_text(v) for v in row if not _empty(v)}
        if len(distinct) < 2:
            continue                                   # заголовок документа, растянутый на всю ширину
        best = None
        for j, v in enumerate(row):
            if isinstance(v, str) and len(v) <= 80 and NAME_HEADER.search(v) and num(v) is None:
                if best is None or 'наименован' in v.lower():
                    best = j
                if 'наименован' in v.lower():
                    break
        if best is not None:
            return i, best
    return None


def read_tables(path):
    """Все таблицы файла: [Table]. Лист без столбца с наименованиями пропускается."""
    path = Path(path)
    if path.suffix.lower() not in ('.xlsx', '.xlsm'):
        raise ValueError(f'{path.name}: нужен формат .xlsx (старый .xls откройте в Excel и сохраните как .xlsx)')
    wb = load_workbook(path, data_only=True)
    out = []
    try:
        for ws in wb.worksheets:
            if ws.sheet_state != 'visible':
                continue
            grid = _grid(ws)
            found = _find_header(grid)
            if not found:
                continue
            h, c = found
            width = max(len(r) for r in grid)
            grid = [r + [None] * (width - len(r)) for r in grid]
            label_rows, last = [h], h
            for r in range(h + 1, min(h + 4, len(grid))):
                row = grid[r]
                if all(_empty(v) for v in row):
                    break
                if _is_numbering(row):
                    last = r
                    continue
                name = row[c]
                if (_empty(name) or name == grid[h][c]) and not any(num(v) is not None for v in row):
                    label_rows.append(r)
                    last = r
                    continue
                break
            index = {}
            for j in range(width):
                if j == c:
                    continue
                parts = []
                for r in label_rows:
                    v = grid[r][j]
                    if not _empty(v) and _text(v) not in parts:
                        parts.append(_text(v))
                label = ' / '.join(parts)
                if label and ROW_NUMBER.match(fold(label).replace(' ', '')):
                    continue                           # «№ п/п» — нумерация, не данные
                index[j] = label
            table = Table(path.name, ws.title, _text(grid[h][c]), [])
            used = set()
            for r in range(last + 1, len(grid)):
                row = grid[r]
                name = row[c]
                if _empty(name):
                    continue
                name = _text(name)
                values = {}
                for j, label in index.items():
                    v = row[j]
                    if _empty(v):
                        continue
                    if not label:
                        label = f'Столбец {get_column_letter(j + 1)}'
                        index[j] = label
                    n = num(v)
                    values[label] = n if n is not None else _text(v)
                    used.add(j)
                if TOTAL_ROW.match(name):
                    table.totals = {k: v for k, v in values.items() if isinstance(v, (int, float))}
                    continue
                if not values:
                    table.skipped += 1
                    table.rows.append((name, {}, r + 1))
                    continue
                table.rows.append((name, values, r + 1))
            table.columns = [index[j] for j in sorted(index) if j in used and index[j]]
            if table.rows:
                out.append(table)
    finally:
        wb.close()
    return out


def _last_part(label):
    return fold(label.split(' / ')[-1])


@dataclass
class Org:
    name: str
    values: dict = field(default_factory=dict)     # столбец -> [(значение, файл, строка)]
    files: list = field(default_factory=list)
    in_list: bool = True


@dataclass
class Svod:
    name_header: str
    columns: list
    orgs: list                                      # [Org] в порядке вывода
    report: list                                    # [(тип, организация, файл, описание)]
    tables: list
    mode: str = 'sum'

    def value(self, org, column):
        """(значение, есть ли расхождение)."""
        items = org.values.get(column, [])
        if not items:
            return None, False
        numbers = [v for v, _, _ in items if isinstance(v, (int, float))]
        texts = [v for v, _, _ in items if not isinstance(v, (int, float))]
        if texts:
            uniq = list(dict.fromkeys(_text(v) for v, _, _ in items))
            return '; '.join(uniq), len(uniq) > 1
        if self.mode == 'sum':
            total = sum(numbers)
            return (round(total, 6) if isinstance(total, float) else total), False
        return numbers[-1], len(set(numbers)) > 1


def combine(tables, order=None, mode='sum'):
    """tables: [Table]. order: перечень организаций (строки) — порядок и проверка «кого нет».
    mode: 'sum' — числа одной организации из разных таблиц складываются;
          'check' — не складываются: берётся последнее значение, расхождения помечаются."""
    if not tables:
        raise ValueError('В выбранных файлах не найдено таблиц со столбцом «Наименование» (организация, подразделение).')
    report = []
    columns, by_key, by_last = [], {}, {}

    def column_for(label, table_lasts):
        k = fold(label)
        if k in by_key:
            return by_key[k]
        last = _last_part(label)
        # «на лицо» в одном файле и «Личный состав / на лицо» в другом — один столбец, если нет двусмысленности.
        if table_lasts.count(last) == 1 and len(by_last.get(last, [])) == 1:
            return by_last[last][0]
        columns.append(label)
        by_key[k] = label
        by_last.setdefault(last, []).append(label)
        return label

    orgs, keys = {}, []
    listed = [n.strip() for n in (order or []) if n and n.strip()]
    for n in listed:
        k = fold(n)
        if k and k not in orgs:
            orgs[k] = Org(n)
            keys.append(k)
    nospace = {k.replace(' ', ''): k for k in orgs}

    for t in tables:
        lasts = [_last_part(c) for c in t.columns]
        mapping = {c: column_for(c, lasts) for c in t.columns}
        for name, values, rownum in t.rows:
            k = fold(name)
            if k not in orgs and k.replace(' ', '') in nospace:
                k = nospace[k.replace(' ', '')]
            if k not in orgs:
                if listed and not values:
                    continue                           # пустая строка не из перечня (подпись, примечание)
                orgs[k] = Org(name, in_list=not listed)
                keys.append(k)
            org = orgs[k]
            if values and t.file not in org.files:
                org.files.append(t.file)
            for col, v in values.items():
                org.values.setdefault(mapping[col], []).append((v, t.file, rownum))
        if t.skipped:
            report.append(('Сведения', '', t.file, f'лист «{t.sheet}»: строк без данных — {t.skipped}'))

    svod = Svod(tables[0].name_header or 'Наименование', columns, [], report, tables, mode)
    in_list = [orgs[k] for k in keys if orgs[k].in_list]
    extra = [orgs[k] for k in keys if not orgs[k].in_list]
    svod.orgs = in_list + extra

    import difflib
    for org in svod.orgs:
        if listed and not org.in_list:
            hint = difflib.get_close_matches(fold(org.name), [fold(n) for n in listed], n=1, cutoff=0.7)
            near = next((n for n in listed if hint and fold(n) == hint[0]), '')
            report.append(('Нет в перечне', org.name, ', '.join(org.files),
                           'строка добавлена в конец свода' + (f'; возможно, это «{near}» — проверьте написание' if near else '')))
        if not org.values:
            report.append(('Нет данных', org.name, '', 'организации нет ни в одной таблице или строка пустая'))
            continue
        for col in columns:
            items = org.values.get(col, [])
            files = list(dict.fromkeys(f for _, f, _ in items))
            value, conflict = svod.value(org, col)
            if conflict:
                where = '; '.join(f'{f}: {_text(v)}' for v, f, _ in items)
                report.append(('Расхождение', org.name, ', '.join(files), f'«{col}»: {where}'))
            elif mode == 'sum' and len(items) > 1 and isinstance(value, (int, float)):
                report.append(('Сложено', org.name, ', '.join(files), f'«{col}» = ' + ' + '.join(_text(v) for v, _, _ in items)))

    # Итоговые строки исходных файлов: сходится ли «Итого» с суммой строк.
    for t in tables:
        if not t.totals:
            continue
        for col, total in t.totals.items():
            s = sum(v.get(col, 0) for _, v, _ in t.rows if isinstance(v.get(col), (int, float)))
            if abs(s - total) > 1e-6:
                report.append(('Итог не сходится', '', t.file, f'«{col}»: в файле {_text(total)}, по строкам {_text(s)}'))

    # «На лицо» не может быть больше «по списку».
    for col_present in columns:
        if 'на лицо' not in fold(col_present):
            continue
        prefix = fold(col_present)[:fold(col_present).rfind('на лицо')]
        roster = next((c for c in columns if 'по списку' in fold(c) and fold(c).startswith(prefix)), None)
        if not roster:
            continue
        for org in svod.orgs:
            a, _ = svod.value(org, col_present)
            b, _ = svod.value(org, roster)
            if isinstance(a, (int, float)) and isinstance(b, (int, float)) and a > b:
                report.append(('Ошибка в данных', org.name, ', '.join(org.files), f'«{col_present}» ({_text(a)}) больше, чем «{roster}» ({_text(b)})'))
    return svod


def numeric_columns(svod):
    out = []
    for col in svod.columns:
        vals = [svod.value(o, col)[0] for o in svod.orgs]
        vals = [v for v in vals if v is not None]
        if vals and all(isinstance(v, (int, float)) for v in vals):
            out.append(col)
    return out


def _widths(ws, minimum=8, maximum=50):
    for i, col in enumerate(ws.iter_cols(min_row=2 if ws.title == 'Свод' else 1), 1):
        width = max((len(_text(c.value)) for c in col[:300] if c.value is not None and not str(c.value).startswith('=')), default=0)
        ws.column_dimensions[get_column_letter(i)].width = min(maximum, max(minimum, width + 2))


def save(svod, output, title='Свод данных'):
    """Новая книга: «Свод» (с итогом), «Проверка», «Исходные данные»."""
    output = Path(output)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Свод'
    headers = ['№ п/п', svod.name_header] + svod.columns + ['Источник']
    last_col = get_column_letter(len(headers))
    ws['A1'] = f'{title} на {date.today().strftime("%d.%m.%Y")}'
    ws['A1'].font = Font(bold=True, size=13)
    ws.merge_cells(f'A1:{last_col}1')
    ws.append(headers)
    for cell in ws[2]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = BLUE
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BOX
    ws.row_dimensions[2].height = 32
    errors = {r[1] for r in svod.report if r[0] == 'Ошибка в данных'}
    for i, org in enumerate(svod.orgs, 1):
        row = [i, org.name]
        for col in svod.columns:
            row.append(svod.value(org, col)[0])
        source = ', '.join(org.files) if org.files else 'нет данных'
        if not org.in_list:
            source += ' (нет в перечне)'
        row.append(source)
        ws.append(row)
        r = ws.max_row
        for cell in ws[r]:
            cell.border = BOX
            if isinstance(cell.value, str) and cell.value.startswith('='):
                cell.data_type = 's'                   # текст из источника, не формула
        fill = YELLOW if not org.values else ORANGE if not org.in_list else None
        if fill:
            for cell in ws[r]:
                cell.fill = fill
        if org.name in errors:
            ws.cell(r, 2).font = Font(color='C00000', bold=True)
        for j, col in enumerate(svod.columns, 3):
            value, conflict = svod.value(org, col)
            if conflict:
                cell = ws.cell(r, j)
                cell.fill = YELLOW
                cell.comment = Comment('Расхождение: ' + '; '.join(f'{f}: {_text(v)}' for v, f, _ in org.values[col]), 'Искра')
    first, last = 3, ws.max_row
    totals = ['', 'ИТОГО']
    sums = set(numeric_columns(svod))
    for j, col in enumerate(svod.columns, 3):
        letter = get_column_letter(j)
        totals.append(f'=SUM({letter}{first}:{letter}{last})' if col in sums and last >= first else '')
    totals.append('')
    ws.append(totals)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.fill = GRAY
        cell.border = BOX
    _widths(ws)
    ws.column_dimensions['A'].width = 7
    ws.freeze_panes = 'C3'
    ws.print_title_rows = '2:2'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    rep = wb.create_sheet('Проверка')
    rep.append(['Что', 'Организация', 'Файл', 'Подробно'])
    for cell in rep[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = BLUE
    order = {'Ошибка в данных': 0, 'Расхождение': 1, 'Итог не сходится': 2, 'Нет в перечне': 3, 'Нет данных': 4, 'Сложено': 5, 'Сведения': 6}
    for item in sorted(svod.report, key=lambda r: order.get(r[0], 9)):
        rep.append(list(item))
        if item[0] in ('Ошибка в данных', 'Расхождение', 'Итог не сходится'):
            rep.cell(rep.max_row, 1).font = Font(color='C00000', bold=True)
    if not svod.report:
        rep.append(['Замечаний нет', '', '', ''])
    _widths(rep, maximum=90)

    src = wb.create_sheet('Исходные данные')
    src.append(['Файл', 'Лист', 'Строка', svod.name_header] + svod.columns)
    for cell in src[1]:
        cell.font = Font(bold=True)
        cell.fill = GRAY
    for t in svod.tables:
        lasts = [_last_part(c) for c in t.columns]
        mapping = {}
        for c in t.columns:
            mapping[c] = next((x for x in svod.columns if fold(x) == fold(c)), None) or \
                next((x for x in svod.columns if _last_part(x) == _last_part(c) and lasts.count(_last_part(c)) == 1), c)
        for name, values, rownum in t.rows:
            if not values:
                continue
            mapped = {mapping[k]: v for k, v in values.items()}
            src.append([t.file, t.sheet, rownum, name] + [mapped.get(c) for c in svod.columns])
            for cell in src[src.max_row]:
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    cell.data_type = 's'
    src.freeze_panes = 'E2'
    _widths(src)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation.fullCalcOnLoad = True
    wb.save(output)
    return output


def fill_form(svod, template, output):
    """Заполнить готовую форму (бланк свода): по наименованию организации в строку, по подписи столбца в ячейку.
    Ячейки с формулами не трогаются. Возвращает (заполнено ячеек, [организации формы без данных], [данные не из формы])."""
    template, output = Path(template), Path(output)
    if template.resolve() == output.resolve():
        raise ValueError('Сохраните заполненную форму под другим именем — бланк должен остаться пустым.')
    wb = load_workbook(template)
    filled, missing = 0, []
    used = set()
    try:
        ws = wb.worksheets[0]
        grid = _grid(ws)
        found = _find_header(grid)
        if not found:
            raise ValueError('В форме не найден столбец «Наименование» (организация, подразделение).')
        h, c = found
        # Подписи столбцов формы — так же, как в таблицах (шапка может быть в 2–3 строки).
        form_cols = {}
        rows_used = [h] + [r for r in range(h + 1, min(h + 4, len(grid)))
                           if grid[r] and (_empty(grid[r][c]) or grid[r][c] == grid[h][c]) and not _is_numbering(grid[r])
                           and not any(num(v) is not None for v in grid[r])]
        width = max(len(r) for r in grid)
        for j in range(width):
            if j == c:
                continue
            parts = []
            for r in rows_used:
                v = grid[r][j] if j < len(grid[r]) else None
                if not _empty(v) and _text(v) not in parts:
                    parts.append(_text(v))
            if parts:
                form_cols[j] = ' / '.join(parts)
        target = {}
        for j, label in form_cols.items():
            col = next((x for x in svod.columns if fold(x) == fold(label)), None)
            if col is None:
                same = [x for x in svod.columns if _last_part(x) == _last_part(label)]
                col = same[0] if len(same) == 1 else None
            if col:
                target[j] = col
        if not target:
            raise ValueError('Подписи столбцов формы не совпадают со столбцами таблиц.')
        by_key = {fold(o.name): o for o in svod.orgs}
        nospace = {k.replace(' ', ''): o for k, o in by_key.items()}
        start = max(rows_used) + 2                    # строка листа (с 1) сразу под шапкой
        for r in range(start, ws.max_row + 1):
            name = ws.cell(r, c + 1).value
            if _empty(name) or TOTAL_ROW.match(_text(name)) or _is_numbering([ws.cell(r, j + 1).value for j in range(width)]):
                continue
            org = by_key.get(fold(name)) or nospace.get(fold(name).replace(' ', ''))
            if org is None or not org.values:
                missing.append(_text(name))
                continue
            used.add(fold(org.name))
            for j, col in target.items():
                cell = ws.cell(r, j + 1)
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    continue
                value, _ = svod.value(org, col)
                if value is not None:
                    try:
                        cell.value = value
                    except AttributeError:
                        continue                       # часть объединённой ячейки
                    filled += 1
        wb.calculation.fullCalcOnLoad = True
        output.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output)
    finally:
        wb.close()
    extra = [o.name for o in svod.orgs if o.values and fold(o.name) not in used]
    return filled, missing, extra


def svod_files(paths, output, order=None, mode='sum', form=None):
    """Полный цикл для окна и помощника. Возвращает словарь с итогами."""
    tables, failed = [], []
    for p in paths:
        try:
            found = read_tables(p)
            if not found:
                failed.append((Path(p).name, 'нет таблицы со столбцом «Наименование»'))
            tables.extend(found)
        except Exception as exc:
            failed.append((Path(p).name, str(exc)[:200]))
    svod = combine(tables, order, mode)
    for name, why in failed:
        svod.report.insert(0, ('Файл не прочитан', '', name, why))
    out = save(svod, output)
    result = {'output': str(out), 'orgs': len(svod.orgs), 'columns': svod.columns, 'tables': len(tables),
              'failed': failed, 'report': svod.report, 'svod': svod}
    if form:
        form_out = Path(output).with_name(Path(output).stem + ' (по форме).xlsx')
        filled, missing, extra = fill_form(svod, form, form_out)
        result.update(form_output=str(form_out), form_filled=filled, form_missing=missing, form_extra=extra)
    return result


def describe(result):
    """Короткий отчёт по-русски для окна и чата."""
    svod = result['svod']
    counts = {}
    for r in svod.report:
        counts[r[0]] = counts.get(r[0], 0) + 1
    lines = [f'Таблиц: {result["tables"]}, организаций в своде: {result["orgs"]}, столбцов: {len(result["columns"])}.',
             'Столбцы: ' + ', '.join(result['columns'][:12]) + ('…' if len(result['columns']) > 12 else '')]
    for kind in ('Файл не прочитан', 'Ошибка в данных', 'Расхождение', 'Итог не сходится', 'Нет в перечне', 'Нет данных', 'Сложено'):
        if counts.get(kind):
            lines.append(f'{kind}: {counts[kind]}')
    if 'form_output' in result:
        lines.append(f'Форма заполнена: {result["form_filled"]} ячеек.')
        if result['form_missing']:
            lines.append('В форме без данных: ' + ', '.join(result['form_missing'][:10]))
        if result['form_extra']:
            lines.append('Есть в таблицах, но нет в форме: ' + ', '.join(result['form_extra'][:10]))
    lines.append('Подробности — на листе «Проверка».')
    return '\n'.join(lines)
