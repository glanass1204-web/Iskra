from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import Workbook,load_workbook
from iskra_consolidate import consolidate

def write(path,head,rows):
    w=Workbook();s=w.active;s.append(head)
    for r in rows:s.append(r)
    w.save(path)

with TemporaryDirectory() as t:
    p=Path(t); template=p/'template.xlsx'; a=p/'a.xlsx'; b=p/'b.xlsx'; out=p/'out.xlsx'
    write(template,['Организация','Количество'],[])
    write(a,['Количество','Организация','Примечание'],[[5,'А','x'],[7,'Б','y']])
    write(b,['Организация','Количество'],[['А',3],['В',9]])
    result=consolidate(template,[a,b],out,'Организация')
    assert result['rows']==4 and result['issues']>=2,result
    w=load_workbook(out,read_only=True,data_only=True)
    assert w['Сводная']['A2'].value=='А' and w['Сводная']['B2'].value==5
    assert any('Дубликат ключа' in str(r) for r in w['Проверка'].values)
    w.close()
    try:consolidate(template,[a],a)
    except ValueError:pass
    else:raise AssertionError('overwrite allowed')
    print('PASS',result)
