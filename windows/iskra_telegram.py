"""Merge several similar telegrams into one: repeated points are kept once even when worded differently,
introductory phrases are dropped, conflicting dates/numbers are marked for checking, and the result is laid out
in one standard (or the user's own) template. Nothing is invented: every sentence comes from a source telegram."""
import re
from dataclasses import dataclass, field

FIELD = re.compile(r'^\s*(Кому|От кого|Отправитель|Адресат|Дата|Номер|№|Исх\.?\s*№?|Тема|О чем|О чём|Срочность|Гриф|Подпись)\s*:\s*(.*)$',
                   re.IGNORECASE)
CANON = {'кому': 'Кому', 'адресат': 'Кому', 'от кого': 'От кого', 'отправитель': 'От кого', 'дата': 'Дата',
         'номер': 'Номер', '№': 'Номер', 'исх': 'Номер', 'тема': 'Тема', 'о чем': 'Тема', 'срочность': 'Срочность',
         'гриф': 'Гриф', 'подпись': 'Подпись'}
ORDER = ['Гриф', 'Срочность', 'Кому', 'От кого', 'Дата', 'Номер', 'Тема']
TITLE = re.compile(r'^\s*(ТЕЛЕГРАММА|ТЕЛЕФОНОГРАММА|ШИФРТЕЛЕГРАММА|РАСПОРЯЖЕНИЕ)\s*$', re.IGNORECASE)
TEXT_LABEL = re.compile(r'^\s*Текст\s*:\s*', re.IGNORECASE)
# Конец предложения — но не после сокращений адресов и званий («г. Москва, ул. Ленина, д. 5», «тел. 8…»).
SPLIT = re.compile(r'(?<=[.!?])(?<![А-ЯЁ]\.)(?<!\bг\.)(?<!\bст\.)(?<!\bп\.)(?<!№)(?<!\bул\.)(?<!\bд\.)(?<!тел\.)'
                   r'(?<!\bим\.)(?<!\bкв\.)(?<!корп\.)(?<!\bпр\.)(?<!обл\.)(?<!тов\.)(?<!\bт\.е\.)(?<!\bр-н\.)(?<!\bпос\.)'
                   r'\s+(?=[А-ЯЁA-Z0-9«"])')
NAME = re.compile(r'([А-ЯЁ]\.\s?[А-ЯЁ]\.\s?[А-ЯЁ][а-яё]+|[А-ЯЁ][а-яё]+\s+[А-ЯЁ]\.\s?[А-ЯЁ]\.)\s*$')
LIST_MARK = re.compile(r'^\s*(\d{1,2}[.)]|[-–—•])\s+')
MONTHS = {'января': 1, 'февраля': 2, 'марта': 3, 'апреля': 4, 'мая': 5, 'июня': 6, 'июля': 7, 'августа': 8,
          'сентября': 9, 'октября': 10, 'ноября': 11, 'декабря': 12}

# Вводные фразы: в объединённой телеграмме не нужны.
INTRO_SENTENCE = re.compile(r'^(здравия желаю|добрый (день|вечер)|уважаем\w*[^.]*[,!]|(настоящим )?(доводим|довожу|сообщаем|сообщаю|информируем|информирую) до (вашего )?сведения\.?$)', re.IGNORECASE)
INTRO_PREFIX = re.compile(r'^(настоящим\s+)?(доводим|довожу|сообщаем|сообщаю|информируем|информирую)\s+(до\s+(вашего\s+)?сведения,?\s*)?(что\s+)?', re.IGNORECASE)
# Основание («Во исполнение приказа…») — в начало, концовка («О выполнении доложить») — один раз в конец.
BASIS = re.compile(r'^(во исполнение|в соответствии с|на основании|согласно|в целях|в связи с)\b', re.IGNORECASE)
CLOSING = re.compile(r'^((прошу|просим) (принять к (исполнению|сведению)|руководствоваться)|о (получении|выполнении|принятых мерах|результатах)\b.*доложить|исполнение (доложить|подтвердить)|контроль\b.*(оставляю|возлагаю))', re.IGNORECASE)

INSTRUCTION = re.compile(r'(обеспечить|представить|провести|направить|доложить|организовать|прибыть|выделить|подготовить|'
                         r'исполнить|запретить|усилить|проверить|уточнить|назначить|отправить|выставить|сдать|прекратить|'
                         r'привести|принять меры|установить|выполнить|довести)', re.IGNORECASE)

STOP = set('и в во не на с со по к ко о об от до за из у же а но или для при что как это все быть был были будет '
           'его ее их а также прошу просим необходимо срочно'.split())


@dataclass
class Telegram:
    name: str
    title: str = 'ТЕЛЕГРАММА'
    fields: dict = field(default_factory=dict)
    body: list = field(default_factory=list)          # абзацы текста
    signature: list = field(default_factory=list)     # строки подписи


def _canon(label):
    key = re.sub(r'[\s.№]+$', '', label.strip().lower().replace('ё', 'е'))
    key = 'исх' if key.startswith('исх') else key
    return CANON.get(key, label.strip().capitalize())


def parse(name, text):
    t = Telegram(name)
    lines = [ln.rstrip() for ln in text.replace('\r\n', '\n').split('\n')]
    in_signature = False
    for line in lines:
        if TITLE.match(line):
            t.title = line.strip().upper()
            continue
        m = FIELD.match(line)
        if m:
            key = _canon(m.group(1))
            if key == 'Подпись':
                in_signature = True
                if m.group(2).strip():
                    t.signature.append(m.group(2).strip())
                continue
            if m.group(2).strip():
                t.fields.setdefault(key, []).append(m.group(2).strip())
            continue
        if in_signature:
            if line.strip().startswith('---'):
                break
            if line.strip():
                t.signature.append(line.strip())
            continue
        line = TEXT_LABEL.sub('', line).strip()
        if line:
            t.body.append(line)
    # Подпись без слова «Подпись:» — последние строки вида «Начальник штаба / полковник И.И. Иванов».
    if not t.signature and t.body and NAME.search(t.body[-1]) and len(t.body[-1]) <= 80:
        t.signature.append(t.body.pop())
        while t.body and len(t.body[-1]) <= 60 and not t.body[-1].endswith(('.', ';', ':')):
            t.signature.insert(0, t.body.pop())
    return t


def _words(sentence):
    out = set()
    # Числа и даты не считаем общими словами: одинаковый срок у разных организаций не делает пункты одинаковыми.
    for w in re.findall(r'[а-яёa-z]+', sentence.lower().replace('ё', 'е')):
        if w in STOP or len(w) < 3:
            continue
        out.add(w[:5] if len(w) > 5 else w)
    return out


def _facts(sentence):
    return set(re.findall(r'\d{1,2}\.\d{1,2}\.\d{2,4}|\d+(?:[.,]\d+)?', sentence))


def _lead(sentence):
    """Кому адресован пункт: первое слово, если это не глагол поручения («Госпиталю представить…» → «госпиталю»)."""
    m = re.match(r'\s*([А-ЯЁ][а-яё]+(?:\s+[№\d][\w/-]*)?)', sentence)
    if not m or INSTRUCTION.match(m.group(1)) or m.group(1).lower() in ('срочно', 'прошу', 'необходимо', 'всем', 'немедленно'):
        return ''
    return m.group(1).lower()


def similar(a, b):
    """Same point worded differently: share most meaningful words. Разным адресатам — разные пункты."""
    la, lb = _lead(a), _lead(b)
    if la and lb and la != lb:
        return False
    A, B = _words(a), _words(b)
    if not A or not B:
        return False
    inter = len(A & B)
    return inter / len(A | B) >= 0.6 or (min(len(A), len(B)) >= 4 and inter / min(len(A), len(B)) >= 0.85)


def numeric_dates(text):
    """«25 октября 2026 года» → «25.10.2026» (a date without a year is left as is)."""
    def repl(m):
        return f'{int(m.group(1)):02d}.{MONTHS[m.group(2).lower()]:02d}.{m.group(3)}'
    return re.sub(r'\b(\d{1,2})\s+(' + '|'.join(MONTHS) + r')\s+(\d{4})(?:\s*(?:года|г\.?)(?=[\s,.;:)]|$))?', repl, text, flags=re.IGNORECASE)


def tidy(sentence):
    s = re.sub(r'\s+', ' ', sentence).strip()
    s = re.sub(r'\s+([,.;:!?])', r'\1', s)
    s = re.sub(r'([,;])(?=[А-Яа-яЁё])', r'\1 ', s)
    if s and s[0].islower():
        s = s[0].upper() + s[1:]
    if s and s[-1] not in '.!?:;':
        s += '.'
    return s


def _sentences(paragraph):
    paragraph = LIST_MARK.sub('', paragraph)
    return [s for s in (x.strip() for x in SPLIT.split(paragraph)) if len(s) > 1]


def _signature_key(lines):
    return ' '.join(lines).lower().replace('ё', 'е').replace(' ', '')


@dataclass
class Merged:
    title: str
    header: dict
    basis: list
    points: list          # [(sentence, [source names])]
    closing: list
    signature: list
    report: list
    sources: list

    def text(self):
        lines = [self.title, '']
        for f in ORDER:
            if f in self.header:
                lines.append(f'{f}: {self.header[f]}')
        lines += ['', 'Текст:']
        lines += self.basis
        lines += [f'{i}. {s}' for i, (s, _) in enumerate(self.points, 1)]
        lines += self.closing
        lines += ['', 'Подпись:'] + (self.signature or ['[ДОЛЖНОСТЬ]', '[звание] [И.О. Фамилия]'])
        lines += ['', '--- ИСТОЧНИКИ (НЕ ЧАСТЬ ТЕЛЕГРАММЫ) ---']
        lines += [f'{i}. {n}' for i, n in enumerate(self.sources, 1)]
        return '\n'.join(lines)


def merge(sources, main_only=True, dates=True, fix=None):
    """sources: [(name, text)], at least two. fix: optional callable(text) -> (text, n_fixed) for typical errors."""
    if len(sources) < 2:
        raise ValueError('Выберите хотя бы две телеграммы.')
    tgs = [parse(n, t) for n, t in sources]
    report = []

    # Шапка
    header = {}
    for f in ORDER:
        values = []
        for t in tgs:
            for v in t.fields.get(f, []):
                if v.lower() not in [x.lower() for x in values]:
                    values.append(v)
        if f in ('Дата', 'Номер'):
            header[f] = f'[{f.upper()}]'
            if len(values) > 1:
                report.append(f'{f}: в исходных разные ({", ".join(values)}) — укажите новые.')
        elif not values:
            continue
        elif f in ('Кому', 'Тема'):
            header[f] = '; '.join(values)
        elif len(values) == 1:
            header[f] = values[0]
        else:
            header[f] = '[УТОЧНИТЬ: ' + ' / '.join(values) + ']'
            report.append(f'«{f}» различается: {" / ".join(values)}.')

    # Текст
    basis, points, closing = [], [], []
    dropped_intro, merged_same, conflicts = [], 0, []
    for t in tgs:
        for paragraph in t.body:
            for raw in _sentences(paragraph):
                s = raw
                if main_only:
                    if INTRO_SENTENCE.match(s):
                        dropped_intro.append(s)
                        continue
                    stripped = INTRO_PREFIX.sub('', s)
                    if stripped != s and len(stripped) > 10:
                        s = stripped
                if dates:
                    s = numeric_dates(s)
                s = tidy(s)
                if CLOSING.match(s):
                    bucket = closing
                elif BASIS.match(s) and not INSTRUCTION.search(s):
                    bucket = basis
                else:
                    bucket = points
                for i, (existing, names) in enumerate(bucket):
                    if similar(existing, s):
                        merged_same += 1
                        fa, fb = _facts(existing), _facts(s)
                        if fb - fa and fa - fb:
                            conflicts.append(f'«{existing}» — в «{t.name}» другие сроки/числа: {", ".join(sorted(fb - fa))}.')
                            new = existing + f' [УТОЧНИТЬ: в «{t.name}»: {", ".join(sorted(fb - fa))}]'
                        elif fb - fa:
                            # Во второй телеграмме есть срок/число, которого не было, — не теряем его.
                            if len(s) >= len(existing) * 0.6:
                                new = s
                            else:
                                new = existing + f' [в «{t.name}»: {", ".join(sorted(fb - fa))}]'
                            conflicts.append(f'«{existing[:80]}» — из «{t.name}» добавлено: {", ".join(sorted(fb - fa))}.')
                        elif len(s) > len(existing) and fb >= fa:
                            new = s  # более полная формулировка того же пункта
                        else:
                            new = existing
                        bucket[i] = (new, names + [t.name])
                        break
                else:
                    bucket.append((s, [t.name]))
    basis = [b for b, _ in basis]
    closing = [c for c, _ in closing][:2]

    # Подпись
    signatures = []
    for t in tgs:
        if t.signature and _signature_key(t.signature) not in [_signature_key(s) for s in signatures]:
            signatures.append(t.signature)
    if len(signatures) == 1:
        signature = signatures[0]
    elif signatures:
        signature = ['[УТОЧНИТЬ ПОДПИСЬ: ' + ' / '.join(' '.join(s) for s in signatures) + ']']
        report.append('Подписи в телеграммах разные — выберите одну.')
    else:
        signature = []
        report.append('Подпись не найдена — заполните блок подписи.')

    merged = Merged(next(iter({t.title for t in tgs})) if len({t.title for t in tgs}) == 1 else 'ТЕЛЕГРАММА',
                    header, basis, points, closing, signature, report, [t.name for t in tgs])
    if fix is not None:
        fixed_n = 0
        for i, (s, n) in enumerate(merged.points):
            s2, k = fix(s); fixed_n += k; merged.points[i] = (s2, n)
        for lst in (merged.basis, merged.closing):
            for i, s in enumerate(lst):
                s2, k = fix(s); fixed_n += k; lst[i] = s2
        if fixed_n:
            report.append(f'Исправлено типовых ошибок: {fixed_n}.')
    summary = f'Объединено телеграмм: {len(tgs)} → пунктов: {len(points)}'
    if merged_same:
        summary += f', убрано повторов: {merged_same}'
    if dropped_intro:
        summary += f', вводных фраз: {len(dropped_intro)}'
    report.insert(0, summary + '.')
    report += conflicts
    if dropped_intro:
        report.append('Убраны вводные фразы: ' + ' | '.join(dropped_intro))
    return merged


def merge_telegrams(sources, **options):
    """Compatibility wrapper: (merged text, report lines)."""
    m = merge(sources, **options)
    return m.text(), m.report


# ---------- Оформление в Word ----------

def split_editor_text(text):
    """Editor text → (title, header, body paragraphs, signature lines)."""
    text = text.split('--- ИСТОЧНИКИ (НЕ ЧАСТЬ ТЕЛЕГРАММЫ) ---')[0]
    title, header, body, signature = 'ТЕЛЕГРАММА', {}, [], []
    mode = 'header'
    for line in text.replace('\r\n', '\n').split('\n'):
        s = line.strip()
        if TITLE.match(s):
            title = s.upper(); continue
        if re.match(r'^Подпись\s*:', s, re.IGNORECASE):
            mode = 'signature'
            rest = s.split(':', 1)[1].strip()
            if rest:
                signature.append(rest)
            continue
        if mode == 'signature':
            if s:
                signature.append(s)
            continue
        if re.match(r'^Текст\s*:', s, re.IGNORECASE):
            mode = 'body'
            rest = s.split(':', 1)[1].strip()
            if rest:
                body.append(rest)
            continue
        m = FIELD.match(s)
        if mode == 'header' and m:
            header[_canon(m.group(1))] = m.group(2).strip()
            continue
        if s:
            mode = 'body' if mode == 'header' and header else mode
            body.append(s)
    return title, header, body, signature


def is_raport(text):
    head = [l.strip() for l in text.splitlines()[:15]]
    return 'РАПОРТ' in head and 'ТЕЛЕГРАММА' not in head


def save_raport_docx(text, path):
    """Рапорт в Word как принято: шапка (кому, от кого, дата) справа, «РАПОРТ» по центру, текст по ширине с отступом
    первой строки, внизу должность и строка «звание — подпись — фамилия» (фамилия у правого края)."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Cm, Pt
    lines = text.splitlines()
    i = next(n for n, l in enumerate(lines) if l.strip() == 'РАПОРТ')
    head, rest = lines[:i], lines[i + 1:]
    while rest and not rest[0].strip():
        rest = rest[1:]
    while rest and not rest[-1].strip():
        rest = rest[:-1]
    # подпись — последний блок после пустой строки (должность, «звание … фамилия»)
    cut = max((n for n, l in enumerate(rest) if not l.strip()), default=len(rest))
    body, sign = rest[:cut], [l for l in rest[cut:] if l.strip()]
    doc = Document()
    for sec in doc.sections:
        sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(2), Cm(1), Cm(2), Cm(2)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(14)
    normal.paragraph_format.space_before, normal.paragraph_format.space_after = Pt(0), Pt(0)
    width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    for l in head:
        if l.strip():
            doc.add_paragraph(l.strip()).alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run('РАПОРТ').italic = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    for l in body:
        if l.strip():
            p = doc.add_paragraph(l.strip())
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(1.25)
    doc.add_paragraph()
    for l in sign:
        parts = [x.strip() for x in re.split(r'\t+', l) if x.strip()]
        p = doc.add_paragraph()
        if len(parts) >= 2:                                 # «майор → (подпись) → Перов В. В.»
            p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
            p.add_run(parts[0] + '\t' + parts[-1])
        else:
            p.add_run(l.strip())
    doc.save(path)
    return path


def save_docx(text, path, template=None):
    """Lays the telegram out in Word: the user's template with {Кому}, {Текст}, {Подпись}… or the standard one."""
    from docx import Document
    text = text.split('--- ИСТОЧНИКИ')[0].rstrip()
    if is_raport(text):
        return save_raport_docx(text, path)
    if template:
        # Шаблон с метками и текст с «Кому: / Текст: / Подпись:» — подставляем по меткам; иначе текст уже в виде
        # шаблона (так его пишет нейросеть) — переносим строки с оформлением шаблона.
        if is_labeled(text) and has_placeholders(template_text(template)):
            return _fill_template(template, path, *split_editor_text(text))
        return save_like_template(text, path, template)
    title, header, body, signature = split_editor_text(text)
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Cm, Pt
    doc = Document()
    for s in doc.sections:  # ГОСТ Р 7.0.97-2016
        s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(2), Cm(1), Cm(2), Cm(2)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(14)
    pf = normal.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(0), 1.0
    width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin

    def para(text='', bold=False, align=None, indent=None):
        p = doc.add_paragraph()
        if text:
            r = p.add_run(text); r.bold = bold
        if align is not None:
            p.alignment = align
        if indent is not None:
            p.paragraph_format.first_line_indent = indent
        return p

    for f in ('Гриф', 'Срочность'):
        if f in header:
            para(header[f], bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    para(title, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para()
    for f in ('Кому', 'От кого', 'Дата', 'Номер', 'Тема'):
        if f in header:
            p = para()
            p.add_run(f + ': ').bold = True
            p.add_run(header[f])
    para()
    for line in body:
        para(line, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=Cm(1.25))
    para()
    for i, line in enumerate(signature):
        p = para()
        m = NAME.search(line)
        if i == len(signature) - 1 and m and m.start() > 0:
            # «полковник                     И.И. Иванов»: звание слева, фамилия у правого поля.
            p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
            p.add_run(line[:m.start()].strip() + '\t' + m.group(1))
        else:
            p.add_run(line)
    doc.save(path)
    return path


def _fill_template(template, path, title, header, body, signature):
    from copy import deepcopy
    from docx import Document
    doc = Document(template)
    values = {'кому': header.get('Кому', ''), 'от кого': header.get('От кого', ''), 'дата': header.get('Дата', ''),
              'номер': header.get('Номер', ''), 'тема': header.get('Тема', ''), 'гриф': header.get('Гриф', ''),
              'срочность': header.get('Срочность', ''), 'заголовок': title}
    blocks = {'текст': body, 'подпись': signature}

    def paragraphs():
        yield from list(doc.paragraphs)
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    yield from list(cell.paragraphs)

    for p in paragraphs():
        full = p.text
        block = re.fullmatch(r'\s*\{(текст|подпись)\}\s*', full, re.IGNORECASE)
        if block:
            # Абзац-заполнитель размножается: у каждой строки его отступы, шрифт и выравнивание.
            lines = blocks[block.group(1).lower()] or ['']
            for line in lines:
                new = deepcopy(p._p)
                p._p.addprevious(new)
                from docx.text.paragraph import Paragraph
                np_ = Paragraph(new, p._parent)
                runs = np_.runs
                runs[0].text = line
                for r in runs[1:]:
                    r.text = ''
            p._p.getparent().remove(p._p)
            continue
        if '{' in full:
            new_text = re.sub(r'\{([^{}]+)\}', lambda m: values.get(m.group(1).strip().lower(), m.group(0)), full)
            if new_text != full and p.runs:
                p.runs[0].text = new_text
                for r in p.runs[1:]:
                    r.text = ''
    doc.save(path)
    return path


TEMPLATE_HELP = ('Свой шаблон: документ Word, оформленный как нужно — подойдёт и обычная удачная телеграмма отдела. '
                 'Искра и нейросеть составляют текст по нему (та же шапка, порядок, обороты, подпись), а в Word переносят его '
                 'оформление. Ещё точнее — если вместо данных написать метки: '
                 '{Кому}, {От кого}, {Дата}, {Номер}, {Тема}, а отдельными абзацами — {Текст} и {Подпись}. '
                 'Оформление абзаца {Текст} (отступ первой строки, выравнивание, шрифт) получат все пункты.')


# ---------- свой шаблон: не только при сохранении в Word, но и для самого текста ----------

PLACEHOLDER = re.compile(r'\{(кому|от кого|дата|номер|тема|гриф|срочность|заголовок|текст|подпись)\}', re.IGNORECASE)


def _body_items(doc):
    """Абзацы и таблицы документа по порядку: ('p', Paragraph) / ('t', Table)."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit('}', 1)[-1]
        if tag == 'p':
            yield 'p', Paragraph(child, doc)
        elif tag == 'tbl':
            yield 't', Table(child, doc)


def template_text(template):
    """Текст шаблона строками, как он выглядит (таблицы — строка на строку таблицы, ячейки через пробелы)."""
    from docx import Document
    doc = Document(template)
    lines = []
    for kind, item in _body_items(doc):
        if kind == 'p':
            lines.append(item.text)
        else:
            for row in item.rows:
                cells = list(dict.fromkeys(c.text.strip() for c in row.cells if c.text.strip()))
                lines.append('    '.join(cells))
    while lines and not lines[-1].strip():
        lines.pop()
    return '\n'.join(lines)


def has_placeholders(text):
    return bool(PLACEHOLDER.search(text or ''))


def render_template(template, fields):
    """Шаблон с метками → текст для редактора: метки заменены значениями, пустые — [УТОЧНИТЬ: …].
    fields: кому, от кого, дата, номер, тема, текст (строки), подпись (строки)."""
    text = template_text(template)

    def value(m):
        key = m.group(1).lower()
        v = fields.get(key)
        if isinstance(v, (list, tuple)):
            v = '\n'.join(x for x in v if x)
        return v or f'[УТОЧНИТЬ: {key}]'
    return PLACEHOLDER.sub(value, text)


def _similar(a, b):
    import difflib
    a, b = re.sub(r'\s+', ' ', a.strip().lower()), re.sub(r'\s+', ' ', b.strip().lower())
    if not a or not b:
        return 1.0 if a == b else 0.0
    return difflib.SequenceMatcher(None, a, b).ratio()


def save_like_template(text, path, template):
    """Текст в Word с оформлением шаблона: каждая строка получает оформление самого похожего абзаца шаблона
    (шапка — как шапка, пункты — как текст, подпись — как подпись). Поля, колонтитулы и стили шаблона сохраняются."""
    from copy import deepcopy
    from docx import Document
    from docx.text.paragraph import Paragraph
    doc = Document(template)
    items = list(_body_items(doc))
    paras = [p for kind, p in items if kind == 'p']
    if not paras:
        return _fill_template(template, path, *split_editor_text(text))
    filled = [p for p in paras if p.text.strip()]
    body_like = max(filled, key=lambda p: len(p.text)) if filled else paras[0]
    empty_like = next((p for p in paras if not p.text.strip()), body_like)
    lines = text.rstrip().split('\n')
    chosen, pos = [], 0
    for line in lines:
        if not line.strip():
            chosen.append(empty_like); continue
        best, score = None, 0.0
        for j in range(pos, len(paras)):
            s = _similar(line, paras[j].text)
            if s > score:
                best, score = j, s
        if best is not None and score >= 0.55:
            chosen.append(paras[best]); pos = best + 1
        elif len(line) > 60:
            chosen.append(body_like)
        else:
            nxt = next((paras[j] for j in range(pos, len(paras)) if paras[j].text.strip()), body_like)
            chosen.append(nxt)
    anchor = items[0][1]._p if items[0][0] == 'p' else items[0][1]._tbl
    for line, src in zip(lines, chosen):
        new = deepcopy(src._p)
        anchor.addprevious(new)
        p = Paragraph(new, src._parent)
        runs = p.runs
        if runs:
            runs[0].text = line
            for r in runs[1:]:
                r.text = ''
        elif line:
            p.add_run(line)
    for kind, item in items:                          # старое содержание шаблона убираем, оформление осталось
        el = item._p if kind == 'p' else item._tbl
        if el.getparent() is not None and not el.tag.endswith('sectPr'):
            el.getparent().remove(el)
    doc.save(path)
    return path


def is_labeled(text):
    """Текст в виде «Кому: … Текст: … Подпись:» (стандартный вид Искры)."""
    return bool(re.search(r'^\s*Кому\s*:', text, re.MULTILINE)) and bool(re.search(r'^\s*(Текст|Подпись)\s*:?\s*$', text, re.MULTILINE))
