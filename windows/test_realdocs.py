"""Документы, прошедшие через разные программы (LibreOffice: .docx → .doc/.odt → .docx — стили Standard,
Text body, WW8Num, как у пользователя). Таблица ПУ: перенос, сутки, шрифт шаблона, зелёный только у изменений.
Нужен LibreOffice (soffice); на компьютере без него тест пропускается."""
import shutil
import subprocess
import tempfile
import unittest
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

import iskra_highlight as H

HEADS = ['Наименование ПУ', 'Подчиненные в/ч', 'Удаление от линии фронта', 'Старший на ПУ', 'Телефоны',
         'Количество ОС/ВВСТ', 'Положение', 'Характер ПУ, количество запасных районов', 'Время нахождения в районе']
WAS = [['', '', '77 км', '', '', '', '', '', ''], ['', '', '105 км', '', '', '', '', '', ''],
       ['КП 5447 обр СпН', '', '56 км', 'НШ гв п/п-к Залупа В.А.', '12188 4444\n18197 4400', '27/2',
        'с-в окраина ЗАЛУПИНСК', 'Заброшенное помещение,\n1 запасной район', '4 сут. с 19.10.2026'],
       ['КП 7001', '', '40 км', 'НШ м-р Иванов И.И.', '11111 2222', '30/5', 'ЛИМАН', 'палатки', '2 сут. с 21.10.2026']]
GREEN = {(2, 0, -1), (2, 2, -1), (2, 5, -1), (2, 6, 1)}


def build(path, rows, green=(), how='font', font='Times New Roman', size=10):
    d = Document()
    d.styles['Normal'].font.name, d.styles['Normal'].font.size = font, Pt(size)
    d.add_paragraph().add_run('СВЕДЕНИЯ о пунктах управления').bold = True
    t = d.add_table(rows=2 + len(rows) + 3, cols=9)
    t.style = 'Table Grid'
    for k, h in enumerate(HEADS):
        t.cell(0, k).paragraphs[0].add_run(h).bold = True
    sec = t.cell(1, 0).merge(t.cell(1, 8))
    sec.paragraphs[0].add_run('СЕВЕРОДОНЕЦКОЕ НАПРАВЛЕНИЕ').bold = True
    for r_i, values in enumerate(rows, 2):
        for k, v in enumerate(values):
            c = t.cell(r_i, k)
            for i, line in enumerate(v.split('\n')):
                run = (c.paragraphs[0] if i == 0 else c.add_paragraph()).add_run(line)
                if (r_i - 2, k, i) in green or (r_i - 2, k, -1) in green:
                    if how == 'font':
                        run.font.color.rgb = RGBColor(0x00, 0xB0, 0x50)
                    else:
                        run.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN
            if r_i < 4:
                shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), 'C2D69B' if k % 2 else '9BBB59')
                c._tc.get_or_add_tcPr().append(shd)
    d.add_paragraph('Начальник штаба')
    d.save(path)


@unittest.skipUnless(shutil.which('soffice'), 'нет LibreOffice')
class ThroughOtherPrograms(unittest.TestCase):
    def roundtrip(self, src, fmt):
        mid, back = self.d / f'mid_{fmt}', self.d / f'back_{fmt}'
        subprocess.run(['soffice', '--headless', '--convert-to', fmt, '--outdir', str(mid), str(src)], capture_output=True, timeout=180)
        subprocess.run(['soffice', '--headless', '--convert-to', 'docx:MS Word 2007 XML', '--outdir', str(back),
                        str(mid / f'{src.stem}.{fmt}')], capture_output=True, timeout=180)
        out = self.d / f'{src.stem}_{fmt}.docx'
        (back / f'{src.stem}.docx').rename(out)
        return out

    def test_pu_table_after_libreoffice(self):
        self.d = Path(tempfile.mkdtemp())
        now = [r[:] for r in WAS]
        now[2][2], now[2][5], now[2][6] = '43 км', '27/4', 'с-в окраина ЗАЛУПИНСК\nПЕЧОРИНСК'
        build(self.d / 'tpl.docx', WAS)
        build(self.d / 'base.docx', WAS, font='Calibri', size=9)
        build(self.d / 'src.docx', now, GREEN, font='Courier New', size=11)
        build(self.d / 'srcm.docx', now, GREEN, how='marker', font='Arial', size=8)
        tpls = [self.d / 'tpl.docx', self.roundtrip(self.d / 'tpl.docx', 'doc'), self.roundtrip(self.d / 'tpl.docx', 'odt')]
        srcs = [self.d / 'src.docx', self.d / 'srcm.docx', self.roundtrip(self.d / 'src.docx', 'doc')]
        bases = [None, self.d / 'base.docx', self.roundtrip(self.d / 'base.docx', 'doc')]
        for tpl in tpls:
            for src in srcs:
                for base in bases:
                    with self.subTest(tpl=tpl.name, src=src.name, base=base and base.name):
                        out = self.d / 'out.docx'
                        rep = H.merge(tpl, [src], out, mark=H.FONT, base=base, today=date(2026, 10, 24))
                        t = Document(str(out)).tables[0]
                        self.assertEqual([t.cell(4, k).text for k in (2, 5, 6, 8)],
                                         ['43 км', '27/4', 'с-в окраина ЗАЛУПИНСК\nПЕЧОРИНСК', '1 сут. с 24.10.2026'])
                        self.assertEqual([t.cell(5, k).text for k in (2, 8)], ['40 км', '4 сут. с 21.10.2026'])

                        def green(r, k):
                            return [any(x.font.color.rgb is not None and str(x.font.color.rgb) == '00B050' for x in p.runs)
                                    for p in t.cell(r, k).paragraphs]
                        self.assertEqual((green(4, 2), green(4, 6), green(4, 8), green(5, 2), green(5, 8)),
                                         ([True], [False, True], [True], [False], [False]))
                        fonts = set()
                        for r in t._tbl.iter(qn('w:r')):
                            if ''.join(x.text or '' for x in r.iter(qn('w:t'))).strip():
                                rpr = r.find(qn('w:rPr'))
                                fonts.add((rpr.find(qn('w:rFonts')).get(qn('w:ascii')), rpr.find(qn('w:sz')).get(qn('w:val'))))
                        self.assertEqual(fonts, {('Times New Roman', '20')})
                        self.assertEqual(rep.unplaced, [])


if __name__ == '__main__':
    unittest.main()
