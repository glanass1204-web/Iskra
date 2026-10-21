import tempfile
import unittest
from datetime import date
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from iskra_inbox import Inbox, analyze
from iskra_compose import reply, reminder, compose, svod_report, unfilled
from iskra_services import telegram_check

TODAY = date(2026, 10, 21)
TG = '''ТЕЛЕГРАММА
От кого: Штаб округа
Номер: 154/2
Дата: 20.10.2026
Текст:
Довожу до вашего сведения, что проводится проверка готовности.
1. Представить сведения о численности личного состава до 25.10.2026.
2. Обеспечить готовность техники.
Подпись: Начальник штаба полковник И.И. Иванов'''


def docx(path, text):
    d = Document()
    for line in text.split('\n'):
        d.add_paragraph(line)
    d.save(path)
    return path


class InboxTest(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp())
        self.box = self.root / 'Входящие'
        self.box.mkdir()

    def test_analyze_telegram(self):
        item = analyze(docx(self.box / 'тлг.docx', TG), TODAY)
        self.assertEqual(item['kind'], 'Телеграмма')
        self.assertEqual(item['number'], '№ 154/2')
        self.assertEqual(item['deadline'], '25.10.2026')
        self.assertEqual(item['sender'], 'Штаб округа')

    def test_new_files_only_once(self):
        (self.box / 'старый.txt').write_text('Старый документ, срок до 01.01.2026.', encoding='utf-8')
        inbox = Inbox(self.root / 'data')
        inbox.add_folder(self.box)
        self.assertEqual(inbox.scan(TODAY, settle=0), [])          # уже лежавшие файлы — не новые
        docx(self.box / 'тлг.docx', TG)
        wb = Workbook(); ws = wb.active
        ws.append(['Наименование', 'По списку', 'На лицо']); ws.append(['Госпиталь', 10, 9])
        wb.save(self.box / 'сведения.xlsx')
        (self.box / '~$тлг.docx').write_text('lock', encoding='utf-8')
        found = inbox.scan(TODAY, settle=0)
        self.assertEqual(sorted(i['name'] for i in found), ['сведения.xlsx', 'тлг.docx'])
        table = next(i for i in found if i['name'] == 'сведения.xlsx')
        self.assertTrue(table.get('svod'))
        self.assertEqual(inbox.scan(TODAY, settle=0), [])
        inbox.set_status(str(self.box / 'тлг.docx'), 'выполнен')
        again = Inbox(self.root / 'data')                            # журнал сохраняется между запусками
        self.assertEqual(again.counts(), (1, 0))

    def test_file_still_copying_waits(self):
        inbox = Inbox(self.root / 'data')
        inbox.add_folder(self.box)
        (self.box / 'a.txt').write_text('Документ', encoding='utf-8')
        self.assertEqual(inbox.scan(TODAY, settle=0.01), [])       # первый раз — только заметили
        import time; time.sleep(0.02)
        self.assertEqual(len(inbox.scan(TODAY, settle=0.01)), 1)


class ComposeTest(unittest.TestCase):
    def test_reply_by_points(self):
        text = reply(TG, {'from': 'Госпиталь', 'signature': 'Начальник госпиталя\nполковник П.П. Петров'}, 'тлг.docx')
        self.assertIn('Кому: Штаб округа', text)
        self.assertIn('На № 154/2 от 20.10.2026 докладываю:', text)
        self.assertIn('1. По пункту «Представить сведения о численности личного состава до 25.10.2026» (срок 25.10.2026): '
                      '[УТОЧНИТЬ: что сделано].', text)
        self.assertIn('2. По пункту «Обеспечить готовность техники»', text)
        self.assertNotIn('Довожу', text)
        self.assertIn('полковник П.П. Петров', text)
        self.assertEqual(len(unfilled(text)), 3)                     # 2 пункта + номер
        self.assertTrue(any('Не заполнено' in i for i in telegram_check(text)))

    def test_reminder_and_kinds(self):
        item = {'name': 'тлг.docx', 'number': '№ 154/2', 'date': '20.10.2026', 'deadline': '25.10.2026',
                               'deadline_text': 'Представить сведения о численности до 25.10.2026.', 'sender': 'Госпиталь'}
        text = reminder(item)
        self.assertIn('по телеграмме № 154/2 от 20.10.2026 установлен срок 25.10.2026', text)
        u = compose('Указание', basis='приказа № 5', points=['провести проверку', 'доложить итоги'], deadline='30.10.2026')
        self.assertIn('Во исполнение приказа № 5 ТРЕБУЮ:\n1. Провести проверку.\n2. Доложить итоги.\nО выполнении доложить до 30.10.2026.', u)
        self.assertIn('прошу в срок до [ДАТА] представить', compose('Запрос сведений'))

    def test_svod_report(self):
        from iskra_svod import Table, combine
        t = Table('a.xlsx', 'Лист1', 'Наименование', ['Личный состав / по списку', 'Личный состав / на лицо'],
                  [('Госпиталь', {'Личный состав / по списку': 120, 'Личный состав / на лицо': 115}, 2),
                   ('Полигон', {'Личный состав / по списку': 80, 'Личный состав / на лицо': 78}, 3)])
        text = svod_report(combine([t], ['Госпиталь', 'Полигон', 'Модуль 1']))
        self.assertIn('1. Госпиталь: по списку — 120, на лицо — 115;', text)
        self.assertIn('3. Модуль 1: [УТОЧНИТЬ: нет данных].', text)
        self.assertIn('Итого: по списку — 200, на лицо — 193.', text)


if __name__ == '__main__':
    unittest.main()
