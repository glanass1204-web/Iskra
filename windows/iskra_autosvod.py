"""Автосвод: таблицы подразделений падают в папку — свод за день собирается сам при каждом новом или исправленном файле.
Показывает, кто ещё не прислал, и что изменилось со вчера («Полигон: на лицо −5»).
Сводные книги — в папке сводов («Свод ДД.ММ.ГГГГ.xlsx»), снимки цифр по дням — в data/autosvod/."""
import json
import os
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from iskra_svod import combine, fold, numeric_columns, read_tables, save

DATE_FMT = '%d.%m.%Y'


def _sign(paths):
    out = []
    for p in paths:
        try:
            st = Path(p).stat()
            out.append([Path(p).name, st.st_size, int(st.st_mtime)])
        except OSError:
            pass
    return out


def _n(v):
    if isinstance(v, float):
        return f'{v:.2f}'.rstrip('0').rstrip('.').replace('.', ',')
    return str(v)


def dynamics(today_values, prev_values):
    """[(организация, столбец, было, стало, изменение)] — только числа, только изменившиеся."""
    out = []
    prev_by_key = {fold(org): (org, cols) for org, cols in prev_values.items()}
    for org, cols in today_values.items():
        _, before = prev_by_key.get(fold(org), (org, {}))
        for col, value in cols.items():
            old = before.get(col)
            if isinstance(value, (int, float)) and isinstance(old, (int, float)) and value != old:
                delta = value - old
                out.append((org, col, old, value, round(delta, 6) if isinstance(delta, float) else delta))
    return out


def describe_change(row):
    org, col, old, new, delta = row
    return f'{org}: {col.split(" / ")[-1]} {"+" if delta > 0 else "−"}{_n(abs(delta))} (было {_n(old)}, стало {_n(new)})'


class AutoSvod:
    def __init__(self, root):
        self.root = Path(root)
        self.path = self.root / 'autosvod.json'
        self.snap_dir = self.root / 'autosvod'
        try:
            self.settings = json.loads(self.path.read_text(encoding='utf-8'))
        except (OSError, ValueError):
            self.settings = {}
        self.settings.setdefault('enabled', False)
        self.last = None

    def save_settings(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        tmp = self.path.with_suffix('.tmp')
        tmp.write_text(json.dumps(self.settings, ensure_ascii=False, indent=1), encoding='utf-8')
        os.replace(tmp, self.path)

    def configure(self, folder, out, enabled=True, order=None):
        self.settings.update({'folder': str(folder), 'out': str(out), 'enabled': bool(enabled)})
        if order is not None:
            self.settings['order'] = [o for o in order if o.strip()]
        self.save_settings()

    @property
    def ready(self):
        return bool(self.settings.get('folder')) and Path(self.settings['folder']).is_dir()

    def today_files(self, today=None):
        """Таблицы за день: .xlsx в папке (и на уровень глубже), изменённые сегодня. Свои своды не берутся."""
        today = today or date.today()
        folder = Path(self.settings.get('folder', ''))
        out = Path(self.settings.get('out') or folder).resolve()
        files = []
        for p in list(folder.glob('*.xls*')) + list(folder.glob('*/*.xls*')):
            if p.name.startswith(('~$', 'Свод ')) or p.suffix.lower() not in ('.xlsx', '.xlsm'):
                continue
            if p.resolve().parent == out and p.name.startswith('Свод'):
                continue
            try:
                if datetime.fromtimestamp(p.stat().st_mtime).date() == today:
                    files.append(p)
            except OSError:
                continue
        return sorted(files, key=lambda p: p.stat().st_mtime)          # позже присланный файл — последним

    # ---------- снимки по дням ----------

    def snapshot_path(self, day):
        return self.snap_dir / f'{day.isoformat()}.json'

    def previous_snapshot(self, today):
        if not self.snap_dir.is_dir():
            return None, {}
        days = sorted(p.stem for p in self.snap_dir.glob('*.json') if p.stem < today.isoformat())
        if not days:
            return None, {}
        day = days[-1]
        try:
            return date.fromisoformat(day), json.loads((self.snap_dir / f'{day}.json').read_text(encoding='utf-8'))
        except (OSError, ValueError):
            return None, {}

    # ---------- сборка ----------

    def build(self, today=None, order=None):
        today = today or date.today()
        files = self.today_files(today)
        order = order if order is not None else self.settings.get('order', [])
        prev_day, prev = self.previous_snapshot(today)
        result = {'date': today.strftime(DATE_FMT), 'files': [p.name for p in files], 'output': '', 'missing': [],
                  'changes': [], 'prev_date': prev_day.strftime(DATE_FMT) if prev_day else '', 'orgs': 0, 'failed': []}
        tables = []
        for p in files:
            try:
                tables.extend(read_tables(p))
            except Exception as exc:
                result['failed'].append(f'{p.name}: {exc}'[:200])
        if not tables:
            result['missing'] = list(order) or sorted(prev)
            self.last = result
            return result
        # Повторно присланный файл исправляет прежний: берётся последнее значение, разница видна на листе «Проверка».
        svod = combine(tables, order, mode='check')
        nums = set(numeric_columns(svod))
        values = {o.name: {c: svod.value(o, c)[0] for c in svod.columns if c in nums and svod.value(o, c)[0] is not None}
                  for o in svod.orgs if o.values}
        result['orgs'] = len(values)
        sent = {fold(o) for o in values}
        expected = list(order) or list(prev)
        result['missing'] = [o for o in expected if fold(o) not in sent]
        result['changes'] = dynamics(values, prev)
        out_dir = Path(self.settings.get('out') or self.settings['folder'])
        out_dir.mkdir(parents=True, exist_ok=True)
        output = out_dir / f'Свод {today.strftime(DATE_FMT)}.xlsx'
        save(svod, output, title='Свод')
        self._add_sheets(output, result)
        self.snap_dir.mkdir(parents=True, exist_ok=True)
        self.snapshot_path(today).write_text(json.dumps(values, ensure_ascii=False, indent=1), encoding='utf-8')
        result['output'] = str(output)
        result['svod'] = svod
        self.settings['last_sign'] = _sign(files)
        self.settings['last_build'] = datetime.now().strftime('%d.%m.%Y %H:%M')
        self.save_settings()
        self.last = result
        return result

    def _add_sheets(self, output, result):
        wb = load_workbook(output)
        ws = wb.create_sheet('Динамика', 1)
        ws.append([f'Изменения к {result["prev_date"]}' if result['prev_date'] else 'Предыдущего свода нет — динамика появится завтра'])
        ws['A1'].font = Font(bold=True)
        ws.append(['Организация', 'Показатель', 'Было', 'Стало', 'Изменение'])
        for c in ws[2]:
            c.font = Font(bold=True)
        green, red = PatternFill('solid', fgColor='E2EFDA'), PatternFill('solid', fgColor='FCE4D6')
        for org, col, old, new, delta in result['changes']:
            ws.append([org, col, old, new, delta])
            ws.cell(ws.max_row, 5).fill = green if delta > 0 else red
        for col, width in zip('ABCDE', (32, 34, 10, 10, 12)):
            ws.column_dimensions[col].width = width
        miss = wb.create_sheet('Не прислали', 2)
        miss.append(['Организация'])
        miss['A1'].font = Font(bold=True)
        for o in result['missing']:
            miss.append([o])
        miss.column_dimensions['A'].width = 40
        wb.calculation.fullCalcOnLoad = True
        wb.save(output)
        wb.close()

    def check(self, today=None):
        """Вызов по таймеру: пересобрать, если в папке появились новые или изменённые таблицы за сегодня."""
        if not self.settings.get('enabled') or not self.ready:
            return None
        files = self.today_files(today)
        if not files or _sign(files) == self.settings.get('last_sign'):
            return None
        return self.build(today)

    def summary_lines(self, result=None):
        r = result or self.last
        if not r:
            return []
        lines = [f'Свод на {r["date"]}: прислали {len(r["files"])} файл(ов), организаций в своде: {r["orgs"]}.']
        if r['missing']:
            lines.append('Не прислали: ' + ', '.join(r['missing'][:15]) + ('…' if len(r['missing']) > 15 else ''))
        if r['changes']:
            lines.append(f'Изменения к {r["prev_date"]}: ' + '; '.join(describe_change(c) for c in r['changes'][:8]) +
                         ('…' if len(r['changes']) > 8 else ''))
        elif r['prev_date']:
            lines.append(f'Изменений к {r["prev_date"]} нет.')
        return lines
